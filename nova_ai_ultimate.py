"""
Nova AI Ultimate - Windows desktop voice/text assistant
Version: FULL WORKING SMOOTH FIX - old functions kept intact
Built by combining the user's Nova AI, Nova AI Desktop Tool, SmartAssistant, and AI chat project ideas.

Main goals:
- Voice + text chat
- Gemini-style conversational assistant with optional Gemini/OpenAI/Groq/OpenRouter API support
- Windows laptop control through safe, known actions
- No hard-coded API keys. Put keys in .env.
"""

import os
import re
import json
import sys
import time
import shutil
import difflib
import threading
import datetime
import subprocess
import webbrowser
import ctypes
import platform
import argparse
import hashlib
import secrets
import socket
import base64
import mimetypes
import math
import importlib.util
from pathlib import Path
from urllib.parse import quote_plus

try:
    import tkinter as tk
    from tkinter import ttk, messagebox, scrolledtext
except Exception:
    tk = None

try:
    import requests
except Exception:
    requests = None

try:
    from dotenv import load_dotenv
except Exception:
    def load_dotenv(*args, **kwargs):
        return False

try:
    import pyttsx3
except Exception:
    pyttsx3 = None

try:
    import speech_recognition as sr
    SR_IMPORT_ERROR = ""
except Exception:
    sr = None
    SR_IMPORT_ERROR = "SpeechRecognition is not installed"

try:
    import pyautogui
except Exception:
    pyautogui = None

try:
    import pyperclip
except Exception:
    pyperclip = None

try:
    import pygetwindow as gw
except Exception:
    gw = None

APP_NAME = "Nova AI Ultimate"
USER_NAME = os.getenv("NOVA_USER_NAME", "Adarsh")
BASE_DIR = Path(__file__).resolve().parent
ENV_FILE = BASE_DIR / ".env"
SETTINGS_FILE = BASE_DIR / "settings.json"
MEMORY_FILE = BASE_DIR / "memory.json"
CONTACTS_FILE = BASE_DIR / "contacts.json"
LOG_FILE = BASE_DIR / "nova_log.txt"
APP_INDEX_FILE = BASE_DIR / "app_index.json"
LAPTOP_PROFILE_FILE = BASE_DIR / "laptop_profile.json"
APP_PLAYBOOK_FILE = BASE_DIR / "app_playbooks.json"
VOICE_CACHE_FILE = BASE_DIR / "voice_cache.json"
MOBILE_ACCESS_FILE = BASE_DIR / "mobile_access.txt"
DESKTOP = Path.home() / "Desktop"
FILES_DIR = DESKTOP / "Nova AI Files"
DRAFTS_DIR = FILES_DIR / "Drafts"
SCREENSHOT_DIR = FILES_DIR / "Screenshots"
MEDIA_DIR = FILES_DIR / "Media Studio"
MEDIA_UPLOAD_DIR = MEDIA_DIR / "Uploads"
MEDIA_IMAGE_DIR = MEDIA_DIR / "Images"
MEDIA_VIDEO_DIR = MEDIA_DIR / "Videos"
MEDIA_EDIT_DIR = MEDIA_DIR / "Edited"
FILE_STUDIO_DIR = FILES_DIR / "File Studio"
FILE_UPLOAD_DIR = FILE_STUDIO_DIR / "Uploads"
FILE_CREATED_DIR = FILE_STUDIO_DIR / "Created"
FILE_REPORT_DIR = FILE_STUDIO_DIR / "Reports"


def configure_runtime_dirs():
    """Use Desktop by default, but fall back to the project folder if Windows blocks Desktop writes."""
    global FILES_DIR, DRAFTS_DIR, SCREENSHOT_DIR, MEDIA_DIR, MEDIA_UPLOAD_DIR, MEDIA_IMAGE_DIR, MEDIA_VIDEO_DIR, MEDIA_EDIT_DIR, FILE_STUDIO_DIR, FILE_UPLOAD_DIR, FILE_CREATED_DIR, FILE_REPORT_DIR
    def assign(root):
        global FILES_DIR, DRAFTS_DIR, SCREENSHOT_DIR, MEDIA_DIR, MEDIA_UPLOAD_DIR, MEDIA_IMAGE_DIR, MEDIA_VIDEO_DIR, MEDIA_EDIT_DIR, FILE_STUDIO_DIR, FILE_UPLOAD_DIR, FILE_CREATED_DIR, FILE_REPORT_DIR
        FILES_DIR = root
        DRAFTS_DIR = FILES_DIR / "Drafts"
        SCREENSHOT_DIR = FILES_DIR / "Screenshots"
        MEDIA_DIR = FILES_DIR / "Media Studio"
        MEDIA_UPLOAD_DIR = MEDIA_DIR / "Uploads"
        MEDIA_IMAGE_DIR = MEDIA_DIR / "Images"
        MEDIA_VIDEO_DIR = MEDIA_DIR / "Videos"
        MEDIA_EDIT_DIR = MEDIA_DIR / "Edited"
        FILE_STUDIO_DIR = FILES_DIR / "File Studio"
        FILE_UPLOAD_DIR = FILE_STUDIO_DIR / "Uploads"
        FILE_CREATED_DIR = FILE_STUDIO_DIR / "Created"
        FILE_REPORT_DIR = FILE_STUDIO_DIR / "Reports"

    root = DESKTOP / "Nova AI Files"
    assign(root)
    def ensure_and_test(root):
        assign(root)
        for folder in [FILES_DIR, DRAFTS_DIR, SCREENSHOT_DIR, MEDIA_DIR, MEDIA_UPLOAD_DIR, MEDIA_IMAGE_DIR, MEDIA_VIDEO_DIR, MEDIA_EDIT_DIR, FILE_STUDIO_DIR, FILE_UPLOAD_DIR, FILE_CREATED_DIR, FILE_REPORT_DIR]:
            folder.mkdir(parents=True, exist_ok=True)
        test = MEDIA_DIR / ".nova_dir_test.tmp"
        test.write_text("ok", encoding="utf-8")
        test.unlink(missing_ok=True)

    try:
        ensure_and_test(root)
    except Exception:
        root = BASE_DIR / "Nova AI Files"
        ensure_and_test(root)


configure_runtime_dirs()

load_dotenv(ENV_FILE)

if pyautogui:
    pyautogui.FAILSAFE = True
    pyautogui.PAUSE = 0.05

DEFAULT_SETTINGS = {
    "voice_enabled": True,
    "speak_rate": 145,
    "voice_name": "",
    "voice_language_mode": "auto",
    "hindi_voice_name": "",
    "english_voice_name": "",
    "voice_volume": 100,
    "language": "en-IN",
    "assistant_style": "Human, friendly, short, action-first, English-only replies.",
    "force_english_output": True,
    "confirm_dangerous_actions": True,
    "ai_provider_order": ["gemini", "openai", "groq", "openrouter"],
    "gemini_model": "gemini-2.5-flash",
    "openai_model": "gpt-4.1-mini",
    "groq_model": "llama-3.1-8b-instant",
    "openrouter_model": "meta-llama/llama-3.1-8b-instruct:free",
    "control_mode": True,
    "local_control_permission": False,
    "google_browser_permission": False,
    "gmail_browser_permission": False,
    "tor_browser_permission": False,
    "local_only_mode": True,
    "plan_mode_enabled": False,
    "pursue_goal_mode_enabled": False,
    "auto_scan_on_start": True,
    "startup_scan_interval_hours": 12,
    "assistant_enabled": True,
    "access_pin_hash": "",
    "access_pin_salt": "",
    "max_command_length": 1200,
    "web_rate_limit_per_minute": 80,
    "auto_security_headers": True,
    "chat_context_limit": 12,
    "chatgpt_features_enabled": True,
    "max_media_upload_mb": 80,
    "media_ai_vision_enabled": True,
    "spotify_control_enabled": True,
    "whatsapp_control_enabled": True,
    "media_studio_enabled": True,
    "file_studio_enabled": True,
    "security_tools_enabled": True,
    "autopilot_enabled": True,
    "privacy_guard_enabled": True,
    "dark_web_safety_enabled": True
}

DEFAULT_CONTACTS = {
    "ritik": "Ritik",
    "hrithik": "Hrithik",
    "hritik": "Hrithik",
    "piyush": "Piyush",
    "dishant": "Dishant",
    "adarsh": "Adarsh",
    "dosti": "Dosti"
}

APP_ALIASES = {
    "chrome": ["chrome"], "google chrome": ["chrome"],
    "brave": ["brave"], "brave browser": ["brave"], "brave app": ["brave"],
    "edge": ["msedge"], "microsoft edge": ["msedge"],
    "tor": ["tor"], "tor browser": ["tor"], "tor private browser": ["tor"],
    "word": ["winword"], "ms word": ["winword"], "microsoft word": ["winword"],
    "excel": ["excel"], "microsoft excel": ["excel"],
    "powerpoint": ["powerpnt"], "power point": ["powerpnt"],
    "notepad": ["notepad"], "calculator": ["calc"], "paint": ["mspaint"],
    "cmd": ["cmd"], "command prompt": ["cmd"], "powershell": ["powershell"],
    "file explorer": ["explorer"], "explorer": ["explorer"], "files": ["explorer"], "file": ["explorer"],
    "this pc": ["explorer"], "my computer": ["explorer"],
    "setting": ["ms-settings:"], "settings": ["ms-settings:"], "windows settings": ["ms-settings:"],
    "system setting": ["ms-settings:about"], "system settings": ["ms-settings:about"], "about pc": ["ms-settings:about"],
    "wifi": ["ms-settings:network-wifi"], "wifi setting": ["ms-settings:network-wifi"], "wifi settings": ["ms-settings:network-wifi"],
    "bluetooth": ["ms-settings:bluetooth"], "bluetooth setting": ["ms-settings:bluetooth"], "bluetooth settings": ["ms-settings:bluetooth"],
    "network settings": ["ms-settings:network"], "sound settings": ["ms-settings:sound"],
    "display settings": ["ms-settings:display"], "apps settings": ["ms-settings:appsfeatures"],
    "windows security": ["windowsdefender:"], "security": ["windowsdefender:"], "virus protection": ["windowsdefender:"],
    "firewall": ["ms-settings:windowsdefender"], "firewall settings": ["ms-settings:windowsdefender"],
    "windows update": ["ms-settings:windowsupdate"], "update settings": ["ms-settings:windowsupdate"],
    "privacy settings": ["ms-settings:privacy"], "privacy": ["ms-settings:privacy"],
    "vpn settings": ["ms-settings:network-vpn"], "vpn": ["ms-settings:network-vpn"],
    "proxy settings": ["ms-settings:network-proxy"], "proxy": ["ms-settings:network-proxy"],
    "cast settings": ["ms-settings-connectabledevices:devicediscovery"], "wireless display": ["ms-settings-connectabledevices:devicediscovery"],
    "display projection": ["ms-settings-connectabledevices:devicediscovery"], "project settings": ["ms-settings:project"],
    "printers settings": ["ms-settings:printers"], "printer settings": ["ms-settings:printers"],
    "mouse settings": ["ms-settings:mousetouchpad"], "keyboard settings": ["ms-settings:keyboard"],
    "touchpad settings": ["ms-settings:devices-touchpad"], "camera settings": ["ms-settings:camera"],
    "microphone settings": ["ms-settings:privacy-microphone"], "mic settings": ["ms-settings:privacy-microphone"],
    "storage settings": ["ms-settings:storagesense"], "battery settings": ["ms-settings:batterysaver"],
    "power settings": ["ms-settings:powersleep"], "notification settings": ["ms-settings:notifications"],
    "focus settings": ["ms-settings:quiethours"], "date time settings": ["ms-settings:dateandtime"],
    "language settings": ["ms-settings:regionlanguage"], "region settings": ["ms-settings:regionformatting"],
    "account settings": ["ms-settings:accounts"], "signin settings": ["ms-settings:signinoptions"],
    "personalization settings": ["ms-settings:personalization"], "theme settings": ["ms-settings:themes"],
    "taskbar settings": ["ms-settings:taskbar"], "startup apps": ["ms-settings:startupapps"],
    "default apps": ["ms-settings:defaultapps"], "installed apps": ["ms-settings:appsfeatures"],
    "optional features": ["ms-settings:optionalfeatures"], "troubleshoot settings": ["ms-settings:troubleshoot"],
    "recovery settings": ["ms-settings:recovery"], "activation settings": ["ms-settings:activation"],
    "clock": ["ms-clock:"], "alarms": ["ms-clock:"], "control panel": ["control"],
    "whatsapp": ["whatsapp:"], "whatsapp app": ["whatsapp:"],
    "spotify": ["spotify:"], "spotify app": ["spotify:"], "youtube": ["https://www.youtube.com"], "youtube app": ["https://www.youtube.com"],
    "gmail": ["https://mail.google.com"], "google": ["https://www.google.com"],
    "play store": ["ms-windows-store://search/?query=Google%20Play%20Games"],
    "microsoft store": ["ms-windows-store:"], "store": ["ms-windows-store:"],
    "vs code": ["code"], "visual studio code": ["code"], "visual studio": ["devenv"],
    "copilot": ["microsoft-copilot:"], "myasus": ["MyASUS"], "my asus": ["MyASUS"], "myasus tools": ["MyASUS"], "my asus tools": ["MyASUS"],
    "asus": ["MyASUS"], "asus app": ["MyASUS"], "camera": ["microsoft.windows.camera:"], "photos": ["ms-photos:"],
    "calendar": ["outlookcal:"], "mail": ["outlookmail:"], "xbox": ["xbox:"]
}

DEFAULT_APP_PLAYBOOKS = {
    "spotify": {
        "type": "desktop app",
        "can_do": [
            "Open Spotify Desktop",
            "Search for a named song/artist/playlist",
            "Try one safe Play click on the selected result",
            "Use media keys for next, previous, pause, and resume",
        ],
        "commands": [
            "open spotify",
            "play jhol song in spotify",
            "play song in spotify",
            "next song",
            "pause spotify",
        ],
        "limits": [
            "Spotify has no reliable local public API without login/developer setup",
            "Ads, login screens, layout changes, or device selection can block automation",
            "Nova must not claim a song is playing unless it can verify it",
        ],
        "execution_rule": "If a song name is given, search then click one likely result. If no song name is given, open a playlist/current context and click Play only once.",
    },
    "whatsapp": {
        "type": "desktop app",
        "can_do": [
            "Open WhatsApp Desktop",
            "Search/open a contact best-effort",
            "Try message/call navigation best-effort",
        ],
        "commands": [
            "open whatsapp",
            "search Dishant in WhatsApp",
            "call Dishant on WhatsApp",
            "message Ritik hello in WhatsApp",
        ],
        "limits": [
            "WhatsApp Desktop does not expose a guaranteed public call/message automation API",
            "Login, permissions, layout changes, or contact ambiguity can block automation",
            "Nova must never open Google for WhatsApp unless the user clearly asks for web/browser",
        ],
        "execution_rule": "Always try local WhatsApp Desktop first; report exactly what was attempted.",
    },
    "youtube": {
        "type": "website/PWA",
        "can_do": [
            "Open YouTube app/PWA if installed",
            "Open YouTube search results",
            "Open the first found video when Python web access can read results",
        ],
        "commands": ["open youtube", "search youtube python basics", "play lo-fi music on youtube"],
        "limits": ["Browser autoplay may block playback", "Nova cannot guarantee a video started unless it can inspect the player"],
        "execution_rule": "Use YouTube only when the user says YouTube or asks web video playback.",
    },
    "google": {
        "type": "website",
        "can_do": ["Open Google search", "Open Google services when permission is ON", "Use online AI for answers when API is reachable"],
        "commands": ["search google laptop battery tips", "open google drive", "open google photos"],
        "limits": ["Search result pages are opened in browser; Nova does not silently click unknown results"],
        "execution_rule": "Use Google search only for explicit search/google commands, not failed app-open commands.",
    },
    "gmail": {
        "type": "website",
        "can_do": ["Open Gmail", "Open Gmail search", "Open a compose draft with recipient/subject/body"],
        "commands": ["open gmail", "search mail from college", "compose email to name@example.com subject hello body hi"],
        "limits": ["Sending, deleting, archiving, or marking emails requires user review/final click"],
        "execution_rule": "Never send or delete email automatically.",
    },
    "windows_settings": {
        "type": "Windows settings",
        "can_do": ["Open specific settings pages", "Toggle some settings by supported safe APIs", "Ask for UAC/admin when needed"],
        "commands": ["open bluetooth settings", "turn off wifi", "open display settings", "volume down"],
        "limits": ["Some Windows toggles are blocked by design or need manual confirmation/UAC"],
        "execution_rule": "Verify or say best-effort; never say done when Windows only opened a settings page.",
    },
    "files": {
        "type": "local files",
        "can_do": ["Create documents/code files", "Analyze selected files", "Open named files from common folders", "Export chat"],
        "commands": ["make python file with date time code", "analyze selected PDF", "open research paper word file"],
        "limits": ["Nova scans common folders and its own workspace, not every private file on the whole disk by default"],
        "execution_rule": "Use selected/uploaded file first, then known common folders; avoid guessing wrong files.",
    },
    "media_studio": {
        "type": "local photo/video tools",
        "can_do": ["Analyze images/videos", "Create local prompt art", "Crop/edit/merge selected media", "Use FFmpeg for video"],
        "commands": ["analyze this photo", "cinematic portrait keep same face", "crop 9:16", "merge selected photos"],
        "limits": ["Local edits are Pillow/FFmpeg drafts; true semantic image/video generation needs a connected AI model/API"],
        "execution_rule": "Work only on selected files for GUI media tasks; never mix old latest files with selected files.",
    },
}

DANGEROUS_WORDS = ["delete", "remove", "uninstall", "shutdown", "restart", "log off", "format", "erase"]
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif", ".tif", ".tiff"}
VIDEO_EXTS = {".mp4", ".mov", ".mkv", ".avi", ".webm", ".m4v", ".wmv"}
TEXT_FILE_EXTS = {
    ".txt", ".md", ".py", ".js", ".ts", ".html", ".css", ".json", ".csv",
    ".log", ".bat", ".cmd", ".ps1", ".xml", ".yaml", ".yml", ".ini", ".env"
}
DOC_FILE_EXTS = {".docx", ".xlsx", ".pptx", ".pdf", ".rtf"}
GENERAL_FILE_EXTS = TEXT_FILE_EXTS | DOC_FILE_EXTS | IMAGE_EXTS | VIDEO_EXTS

# -------------------- Storage --------------------
def load_json(path: Path, default):
    try:
        if path.exists():
            return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        pass
    path.write_text(json.dumps(default, indent=4), encoding="utf-8")
    return default


def save_json(path: Path, data):
    path.write_text(json.dumps(data, indent=4), encoding="utf-8")


settings = load_json(SETTINGS_FILE, DEFAULT_SETTINGS)
_settings_changed = False
for _k, _v in DEFAULT_SETTINGS.items():
    if _k not in settings:
        settings[_k] = _v
        _settings_changed = True
if _settings_changed:
    save_json(SETTINGS_FILE, settings)
contacts = load_json(CONTACTS_FILE, DEFAULT_CONTACTS)
memory = load_json(MEMORY_FILE, {"notes": [], "history": []})


def save_app_playbooks():
    existing = load_json(APP_PLAYBOOK_FILE, DEFAULT_APP_PLAYBOOKS)
    changed = False
    for key, value in DEFAULT_APP_PLAYBOOKS.items():
        if key not in existing:
            existing[key] = value
            changed = True
    if changed or not APP_PLAYBOOK_FILE.exists():
        save_json(APP_PLAYBOOK_FILE, existing)
    return existing


def app_playbooks():
    return save_app_playbooks()


def app_playbook_summary():
    books = app_playbooks()
    parts = []
    for name in ["spotify", "whatsapp", "youtube", "google", "gmail", "windows_settings", "files", "media_studio"]:
        book = books.get(name, {})
        can_do = "; ".join(book.get("can_do", [])[:3])
        limits = "; ".join(book.get("limits", [])[:2])
        if can_do:
            parts.append(f"{name}: can {can_do}. Limits: {limits}.")
    return " App/website playbooks: " + " | ".join(parts)


def playbook_for_command(c):
    c = normalize_words(c)
    if "spotify" in c:
        return "spotify"
    if "whatsapp" in c:
        return "whatsapp"
    if "youtube" in c:
        return "youtube"
    if "gmail" in c or "email" in c or "mail" in c:
        return "gmail"
    if "google" in c or "search" in c:
        return "google"
    if "setting" in c or "settings" in c or any(x in c for x in ["wifi", "bluetooth", "volume", "brightness", "firewall", "defender"]):
        return "windows_settings"
    if "file" in c or "document" in c or "folder" in c:
        return "files"
    if any(x in c for x in ["photo", "image", "video", "media", "crop", "merge"]):
        return "media_studio"
    return ""


def app_playbook_command(raw):
    c = normalize_words(raw)
    key = playbook_for_command(c)
    intent_words = ["how", "what can", "what do", "kaise", "kya", "guide", "help", "use", "work", "status", "capability", "capabilities", "learn", "about"]
    if not any(word in c for word in intent_words):
        return False, ""
    if not key:
        last_app = str(assistant_context().get("last_app", ""))
        key = playbook_for_command(last_app)
    if not key:
        return False, ""
    book = app_playbooks().get(key, {})
    if not book:
        return False, ""
    lines = [
        f"{key.replace('_', ' ').title()} playbook",
        "What Nova can do:",
    ]
    lines.extend(f"- {item}" for item in book.get("can_do", []))
    lines.append("Good commands:")
    lines.extend(f"- {item}" for item in book.get("commands", []))
    lines.append("Limits:")
    lines.extend(f"- {item}" for item in book.get("limits", []))
    lines.append("Execution rule: " + book.get("execution_rule", "Use the safest local action first and report honestly."))
    return True, "\n".join(lines)


def hash_pin(pin: str, salt: str) -> str:
    return hashlib.sha256((salt + str(pin)).encode("utf-8")).hexdigest()


def access_lock_enabled() -> bool:
    return bool(settings.get("access_pin_hash") and settings.get("access_pin_salt"))


def set_access_pin(pin: str):
    pin = re.sub(r"\D", "", str(pin or ""))
    if len(pin) < 4:
        return False, "PIN must be at least 4 digits."
    salt = secrets.token_hex(16)
    settings["access_pin_salt"] = salt
    settings["access_pin_hash"] = hash_pin(pin, salt)
    save_json(SETTINGS_FILE, settings)
    return True, "Local PIN lock is enabled. Nova will ask for this PIN in the GUI."


def verify_access_pin(pin: str) -> bool:
    if not access_lock_enabled():
        return True
    return secrets.compare_digest(hash_pin(pin, settings.get("access_pin_salt", "")), settings.get("access_pin_hash", ""))


def set_assistant_power(enabled: bool) -> str:
    """Single source of truth for the GUI ON/OFF button and voice commands."""
    settings["assistant_enabled"] = bool(enabled)
    if enabled:
        settings["control_mode"] = True
        settings["local_control_permission"] = True
    save_json(SETTINGS_FILE, settings)
    return "Assistant is ON. Clear laptop commands can run again." if enabled else "Assistant is OFF. I will not run laptop-control actions until you turn it on."


def issue_access_token() -> str:
    token = secrets.token_urlsafe(32)
    ACCESS_TOKENS.add(token)
    return token


def generate_pin(length=6) -> str:
    return "".join(secrets.choice("0123456789") for _ in range(length))


def get_lan_ip() -> str:
    """Best-effort LAN IP for mobile access."""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
            if ip and not ip.startswith("127."):
                return ip
    except Exception:
        pass
    try:
        host = socket.gethostname()
        for ip in socket.gethostbyname_ex(host)[2]:
            if ip and not ip.startswith("127."):
                return ip
    except Exception:
        pass
    return "127.0.0.1"

# -------------------- Logging + Speech --------------------
_engine = None
_speech_process = None
_speech_lock = threading.Lock()
ACCESS_TOKENS = set()
_NETWORK_CACHE = {"time": 0.0, "value": None}
_VOICE_CACHE = {"time": 0.0, "details": []}
SERVER_MOBILE_MODE = False
SERVER_MOBILE_URL = ""
SERVER_LOCAL_URL = ""


def stop_speech():
    """Stop any Windows/pyttsx3 speech currently running."""
    global _speech_process
    with _speech_lock:
        proc = _speech_process
        _speech_process = None
    if proc and proc.poll() is None:
        try:
            proc.terminate()
            try:
                proc.wait(timeout=1.2)
            except Exception:
                proc.kill()
        except Exception:
            pass
    return True


def init_voice():
    global _engine
    if not pyttsx3:
        return None
    if _engine is None:
        try:
            _engine = pyttsx3.init()
            _engine.setProperty("rate", int(settings.get("speak_rate", 172)))
            _engine.setProperty("volume", max(0.0, min(1.0, int(settings.get("voice_volume", 100)) / 100)))
        except Exception:
            _engine = None
    return _engine


def log(text):
    line = f"[{datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {text}"
    try:
        print(line)
    except Exception:
        pass
    try:
        LOG_FILE.open("a", encoding="utf-8").write(line + "\n")
    except Exception:
        pass


def safe_print(*args):
    try:
        print(*args)
    except Exception as e:
        try:
            log("print failed: " + str(e))
        except Exception:
            pass


def _powershell_exe():
    return shutil.which("powershell.exe") or shutil.which("powershell") or "powershell"


def windows_voice_details(force: bool = False):
    if platform.system().lower() != "windows":
        return []
    now = time.time()
    if not force and _VOICE_CACHE.get("details") and now - float(_VOICE_CACHE.get("time", 0)) < 300:
        return list(_VOICE_CACHE["details"])
    if not force and VOICE_CACHE_FILE.exists():
        try:
            cached = load_json(VOICE_CACHE_FILE, {})
            details = cached.get("details", [])
            if details and now - float(cached.get("time", 0)) < 86400:
                _VOICE_CACHE.update({"time": now, "details": details})
                return list(details)
        except Exception:
            pass
    ps = (
        "Add-Type -AssemblyName System.Speech; "
        "$s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
        "$s.GetInstalledVoices() | ForEach-Object { "
        "[pscustomobject]@{Name=$_.VoiceInfo.Name; Culture=$_.VoiceInfo.Culture.Name; Gender=$_.VoiceInfo.Gender.ToString()} "
        "} | ConvertTo-Json -Compress"
    )
    try:
        result = subprocess.run(
            [_powershell_exe(), "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            timeout=12,
        )
        if result.returncode != 0 or not result.stdout.strip():
            return list(_VOICE_CACHE.get("details") or [])
        data = json.loads(result.stdout)
        if isinstance(data, dict):
            data = [data]
        details = [
            {
                "name": str(item.get("Name", "")),
                "culture": str(item.get("Culture", "")),
                "gender": str(item.get("Gender", "")),
            }
            for item in data if item.get("Name")
        ]
        _VOICE_CACHE.update({"time": now, "details": details})
        try:
            save_json(VOICE_CACHE_FILE, {"time": now, "details": details})
        except Exception:
            pass
        return list(details)
    except Exception as e:
        log(f"Voice details error: {e}")
        if VOICE_CACHE_FILE.exists():
            try:
                details = load_json(VOICE_CACHE_FILE, {}).get("details", [])
                if details:
                    return list(details)
            except Exception:
                pass
        return list(_VOICE_CACHE.get("details") or [])


def text_looks_hindi(text: str) -> bool:
    sample = str(text or "").lower()
    if re.search(r"[\u0900-\u097f]", sample):
        return True
    hindi_words = [
        "namaste", "aap", "tum", "main", "mujhe", "karo", "kar diya", "kholo", "band",
        "chalu", "batao", "kaise", "kya", "hai", "nahi", "haan", "theek", "samajh",
        "laptop", "settings", "permission", "awaaz", "aawaz", "saaf", "hindi",
    ]
    return sum(1 for word in hindi_words if re.search(r"\b" + re.escape(word) + r"\b", sample)) >= 3


def available_hindi_voice():
    details = windows_voice_details()
    candidates = [
        v["name"] for v in details
        if v.get("culture", "").lower().startswith("hi")
        or "hindi" in v.get("name", "").lower()
        or any(x in v.get("name", "").lower() for x in ["kalpana", "hemant", "heera", "ravi"])
    ]
    return candidates[0] if candidates else ""


def available_english_voice():
    details = windows_voice_details()
    preferred = ["zira", "aria", "jenny", "david", "guy"]
    for key in preferred:
        for v in details:
            name = v.get("name", "")
            culture = v.get("culture", "").lower()
            if key in name.lower() and culture.startswith("en"):
                return name
    for v in details:
        if v.get("culture", "").lower().startswith("en"):
            return v.get("name", "")
    names = windows_voice_names()
    return names[0] if names else ""


def prepare_speech_text(text: str) -> str:
    """Make spoken output clearer. Code stays in chat, voice gets a clean summary."""
    value = str(text or "")
    if "```" in value:
        value = re.sub(r"```[\s\S]*?```", " Code is written in the chat. ", value)
    value = re.sub(r"https?://\S+", " link is in the chat ", value)
    value = re.sub(r"\s+", " ", value).strip()
    # If Windows has no Hindi TTS voice, do not let an English voice garble Devanagari.
    if re.search(r"[\u0900-\u097f]", value) and not available_hindi_voice():
        return "Hindi text is written in chat. Hindi Windows voice is not installed yet, so I will not read it unclearly. Say: setup Hindi voice."
    return value


def choose_voice_for_text(text: str) -> str:
    explicit = str(settings.get("voice_name", "")).strip()
    mode = str(settings.get("voice_language_mode", "auto")).lower()
    if mode != "auto" and explicit:
        return explicit
    if text_looks_hindi(text):
        selected = str(settings.get("hindi_voice_name", "")).strip() or available_hindi_voice()
        if selected:
            return selected
    return explicit or str(settings.get("english_voice_name", "")).strip() or available_english_voice()


def _windows_sapi_speak(text: str, wait: bool = False):
    """Reliable Windows SAPI speech. Uses a temp UTF-8 text file to avoid quote bugs."""
    global _speech_process
    text = prepare_speech_text(text)
    text_file = BASE_DIR / ".nova_speech_text.txt"
    try:
        text_file.write_text(str(text), encoding="utf-8")
    except Exception:
        pass
    selected_voice = choose_voice_for_text(text).replace("'", "''").strip()
    voice_line = f"try {{ $s.SelectVoice('{selected_voice}') }} catch {{ }}; " if selected_voice else ""
    base_rate = int(settings.get("speak_rate", 145))
    if text_looks_hindi(text):
        base_rate = min(base_rate, 135)
    rate = max(-10, min(10, int((base_rate - 170) / 15)))
    volume = max(0, min(100, int(settings.get("voice_volume", 100))))
    path = str(text_file).replace("'", "''")
    ps = (
        "$ErrorActionPreference='Stop'; "
        f"$t = Get-Content -Raw -Encoding UTF8 '{path}'; "
        "Add-Type -AssemblyName System.Speech; "
        "$s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
        "$s.SetOutputToDefaultAudioDevice(); "
        f"{voice_line}"
        f"$s.Rate = {rate}; "
        f"$s.Volume = {volume}; "
        "$s.Speak($t); "
        "$s.Dispose();"
    )
    creationflags = 0
    startupinfo = None
    if platform.system().lower() == "windows":
        try:
            creationflags = subprocess.CREATE_NO_WINDOW
        except Exception:
            creationflags = 0
        try:
            startupinfo = subprocess.STARTUPINFO()
            startupinfo.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        except Exception:
            startupinfo = None
    proc = subprocess.Popen(
        [_powershell_exe(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", ps],
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=creationflags,
        startupinfo=startupinfo,
    )
    with _speech_lock:
        _speech_process = proc
    if wait:
        proc.wait(timeout=max(10, min(90, len(str(text)) // 8 + 8)))
    return proc


def speak(text, wait: bool = False):
    """Speak through Windows desktop audio. wait=True is used by the GUI read bar."""
    text = str(text).strip()
    if not text:
        return True
    log(f"Nova: {text}")
    if not settings.get("voice_enabled", True):
        return False
    if platform.system().lower() == "windows":
        try:
            stop_speech()
            _windows_sapi_speak(text, wait=wait)
            return True
        except Exception as e:
            log(f"Windows voice error: {e}")
    engine = init_voice()
    if engine:
        try:
            engine.say(text)
            engine.runAndWait()
            return True
        except Exception as e:
            log(f"Voice error: {e}")
    return False


def audio_status():
    """Small diagnostics payload for the web GUI and assistant doctor."""
    return {
        "voice_enabled": bool(settings.get("voice_enabled", True)),
        "voice_volume": settings.get("voice_volume", 100),
        "voice_name": settings.get("voice_name", ""),
        "voice_language_mode": settings.get("voice_language_mode", "auto"),
        "hindi_voice_name": settings.get("hindi_voice_name", ""),
        "windows": platform.system().lower() == "windows",
        "powershell": bool(_powershell_exe()),
        "pyttsx3": bool(pyttsx3),
        "windows_voices": windows_voice_names() if platform.system().lower() == "windows" else [],
        "voice_details": windows_voice_details() if platform.system().lower() == "windows" else [],
        "hindi_voice_available": bool(available_hindi_voice()),
    }


def windows_voice_names():
    return [v["name"] for v in windows_voice_details()]


def set_voice_by_name(name):
    voices = windows_voice_names()
    if not voices:
        return False, "No Windows voices found."
    wanted = (name or "").strip().lower()
    selected = ""
    if wanted in ["female", "girl", "woman", "ladki"]:
        selected = next((v for v in voices if any(x in v.lower() for x in ["zira", "heera", "female"])), "")
    elif wanted in ["male", "boy", "man", "ladka"]:
        selected = next((v for v in voices if any(x in v.lower() for x in ["david", "ravi", "mark", "male"])), "")
    else:
        selected = next((v for v in voices if wanted and wanted in v.lower()), "")
    if not selected:
        return False, "I could not find that voice. Available voices: " + ", ".join(voices)
    settings["voice_name"] = selected
    save_json(SETTINGS_FILE, settings)
    speak(f"Voice changed to {selected}.")
    return True, f"Voice changed to {selected}."


def apply_clear_voice_settings():
    """Make current Windows voice slower and easier to understand."""
    changed = False
    if int(settings.get("speak_rate", 172)) > 145:
        settings["speak_rate"] = 145
        changed = True
    if int(settings.get("voice_volume", 100)) < 95:
        settings["voice_volume"] = 100
        changed = True
    if settings.get("voice_language_mode") != "auto":
        settings["voice_language_mode"] = "auto"
        changed = True
    english = available_english_voice()
    if english and not settings.get("english_voice_name"):
        settings["english_voice_name"] = english
        changed = True
    hindi = available_hindi_voice()
    if hindi and settings.get("hindi_voice_name") != hindi:
        settings["hindi_voice_name"] = hindi
        changed = True
    if changed:
        save_json(SETTINGS_FILE, settings)
    return hindi


def voice_status_text():
    details = windows_voice_details()
    hindi = available_hindi_voice()
    lines = [
        "Voice Status",
        f"- Speech rate: {settings.get('speak_rate', 145)} (clear mode)",
        f"- Voice mode: {settings.get('voice_language_mode', 'auto')}",
        f"- Selected voice: {settings.get('voice_name') or 'auto'}",
        f"- Hindi voice: {hindi or 'not installed'}",
        "- Installed voices: " + (", ".join(f"{v['name']} ({v.get('culture') or 'unknown'})" for v in details) if details else "none detected"),
    ]
    if not hindi:
        lines.append("- Hindi clarity needs Windows Hindi text-to-speech voice. Say: setup Hindi voice.")
    return "\n".join(lines)


def fix_voice_clarity():
    hindi = apply_clear_voice_settings()
    if hindi:
        speak("Hello Adarsh. Voice clarity is now slower and clearer.")
        return True, f"Voice clarity fixed. Hindi voice detected: {hindi}. Speech is slower and auto language mode is ON."
    speak("Voice clarity improved. Hindi Windows voice is not installed yet.")
    return True, "Voice clarity improved: speech is slower, volume is full, and auto voice mode is ON. Hindi voice is not installed, so Hindi text may still sound wrong. Say: setup Hindi voice."


def setup_hindi_voice():
    apply_clear_voice_settings()
    hindi = available_hindi_voice()
    if hindi:
        settings["hindi_voice_name"] = hindi
        settings["voice_language_mode"] = "auto"
        save_json(SETTINGS_FILE, settings)
        speak("Hindi voice is ready.")
        return True, f"Hindi voice is already installed and selected for Hindi text: {hindi}."

    script = """
$caps = @(
  'Language.Basic~~~hi-IN~0.0.1.0',
  'Language.TextToSpeech~~~hi-IN~0.0.1.0',
  'Language.Speech~~~hi-IN~0.0.1.0'
)
$installed = @()
foreach ($cap in $caps) {
    $state = Get-WindowsCapability -Online -Name $cap -ErrorAction SilentlyContinue
    if ($state -and $state.State -ne 'Installed') {
        Add-WindowsCapability -Online -Name $cap -ErrorAction Stop | Out-Null
    }
    $after = Get-WindowsCapability -Online -Name $cap -ErrorAction SilentlyContinue
    $installed += ($cap + '=' + $after.State)
}
Start-Process 'ms-settings:speech'
$NovaOk = $true
$NovaMessage = 'Hindi language/voice capability install attempted. Status: ' + ($installed -join ', ') + '. Restart Nova after Windows finishes installing voices.'
"""
    ok, msg = run_elevated_powershell(script, timeout=600)
    if ok:
        return True, msg
    open_uri("ms-settings:speech")
    return True, msg + " Speech settings are open. Add Hindi voice manually, then restart Nova."

# -------------------- Text Normalization --------------------
def clean_command(command: str) -> str:
    c = (command or "").lower().strip()
    c = re.sub(r"[^\w\s:/.-]", " ", c)
    if c in ["assistant on", "assistant off", "assistant switch", "assistant power", "nova on", "nova off"]:
        return c
    for wake in ["nova ai", "hey nova", "ok nova", "nova", "assistant"]:
        if c.startswith(wake + " "):
            c = c[len(wake):].strip()
        elif c == wake:
            return ""
    return re.sub(r"\s+", " ", c).strip()


def normalize_words(c: str) -> str:
    c = clean_command(c)
    c = re.sub(r"\bopen\s*\.\s*", "open ", c)
    replacements = {
        "please": "", "pls": "", "plz": "", "jara": "", "zara": "",
        "mujhe": "", "mere liye": "", "mere laptop me": "", "laptop me": "",
        "kar sakte ho": "", "kar sakta hai": "", "kya tum": "", "tum": "",
        "kholo": "open", "khol": "open", "kholna": "open", "chalu karo": "open", "start karo": "open", "launch karo": "open",
        "band karo": "close", "band kar do": "close", "close kar do": "close", "exit karo": "close",
        "samne lao": "focus", "front me lao": "focus", "aage lao": "focus", "screen par lao": "focus",
        "whats app": "whatsapp", "what's app": "whatsapp", "watsapp": "whatsapp",
        "power point": "powerpoint", "micro soft": "microsoft", "visual studio code": "vs code",
        "gaana": "song", "gana": "song", "gaane": "song", "gane": "song", "music chalao": "play song", "song chalao": "play song", "chalao": "play",
        "wi fi": "wifi", "wi-fi": "wifi", "blue tooth": "bluetooth", "screen shot": "screenshot",
        "type karo": "type", "likho": "type", "likhna": "type", "likhkar": "type", "likh do": "type", "paste karo": "paste", "banaa do": "make", "bana do": "make", "create karo": "create",
        "search karo": "search", "dhundo": "search", "dhoondo": "search", "find karo": "find",
        "call lagao": "call", "call karo": "call", "message bhejo": "message", "bhejo": "send",
        "karo": "", "kar do": "", "chahiye": "", "chahie": "", "do mujhe": "",
        "khul jaaye": "open", "khula": "open", "sev": "save", "save kiya": "saved",
        "setting app": "settings", "setting": "settings", "system setting": "system settings",
        "wifi on": "turn on wifi", "wifi off": "turn off wifi",
        "jo bhi bolo": "do what i say", "jo bolu": "do what i say",
        "wahi karo": "do what i say", "pura control": "full control",
        "pura assistant": "working assistant", "kaam karne vala assistant": "working assistant",
        "kaam karne wala assistant": "working assistant", "thik karo": "repair",
        "theek karo": "repair", "fix karo": "fix",
        "isse": "this", "isko": "this", "ye": "this", "yeh": "this",
        "usse": "it", "usko": "it", "vo": "it", "woh": "it", "vahi": "it",
    }
    for a, b in replacements.items():
        c = re.sub(r"(?<!\w)" + re.escape(a) + r"(?!\w)", b, c)
    c = c.replace("settingss", "settings")
    c = re.sub(r"^(.+?)\s+(open|start|launch)$", r"\2 \1", c)
    return re.sub(r"\s+", " ", c).strip()


def remove_words(text: str, words) -> str:
    out = text
    for w in words:
        out = re.sub(rf"\b{re.escape(w)}\b", " ", out, flags=re.I)
    return re.sub(r"\s+", " ", out).strip()


def clean_entity_name(text: str) -> str:
    text = re.sub(r"[^\w\s+]", " ", text or "")
    return re.sub(r"\s+", " ", text).strip()


def resolve_contact_name(text: str) -> str:
    cleaned = clean_entity_name(text)
    key = cleaned.lower()
    saved = contacts.get(key)
    if isinstance(saved, dict):
        return str(saved.get("name") or cleaned).strip()
    if saved:
        return str(saved).strip()
    return cleaned


def assistant_context():
    return memory.setdefault("context", {})


def remember_command_context(c: str, reply: str = ""):
    ctx = assistant_context()
    ctx["last_command"] = c
    ctx["last_reply"] = str(reply or "")[:240]
    ctx["updated_at"] = datetime.datetime.now().isoformat()
    if any(x in c for x in ["whatsapp", "message ", "call "]):
        ctx["last_app"] = "whatsapp"
        possible_contact = remove_words(c, ["open", "search", "find", "call", "message", "whatsapp", "on", "in", "to", "ko", "par", "send"])
        possible_contact = resolve_contact_name(possible_contact)
        if possible_contact and possible_contact not in ["this", "it"]:
            ctx["last_contact"] = possible_contact
    if "spotify" in c:
        ctx["last_app"] = "spotify"
        ctx["last_platform"] = "spotify"
        query = remove_words(c, ["play", "song", "songs", "music", "on", "in", "spotify", "app", "that", "this"])
        if query:
            ctx["last_media_query"] = query
    elif "youtube" in c:
        ctx["last_app"] = "youtube"
        ctx["last_platform"] = "youtube"
        query = remove_words(c, ["play", "search", "find", "song", "songs", "music", "on", "in", "youtube", "app", "that", "this"])
        if query:
            ctx["last_media_query"] = query
    elif c.startswith(("open ", "start ", "launch ")):
        app = clean_entity_name(remove_words(c, ["open", "start", "launch", "app", "application", "tool", "tools", "this", "it"]))
        if app:
            ctx["last_app"] = app
    if any(x in c for x in ["file", "document", "word", "excel", "powerpoint", "note"]):
        ctx["last_file_command"] = c
    save_json(MEMORY_FILE, memory)


def resolve_followup_command(c: str) -> str:
    ctx = assistant_context()
    last_app = ctx.get("last_app", "")
    last_platform = ctx.get("last_platform", "")
    last_media_query = ctx.get("last_media_query", "")
    last_contact = ctx.get("last_contact", "")
    if c in ["open it", "open this", "open that", "start it", "launch it"] and last_app:
        return "open " + last_app
    if c in ["close it", "close this", "close that", "close current", "band this"] or c.startswith("close it "):
        return "close window"
    if c in ["play it", "play this", "play that"] and last_platform:
        return "play " + (last_media_query or "song") + " in " + last_platform
    if c in ["call him", "call her", "call them", "call this", "call it"] and last_contact:
        return "call " + last_contact + " on whatsapp"
    if c in ["message him", "message her", "message them", "message this"] and last_contact:
        return "message " + last_contact + " in whatsapp"
    return c


def safe_filename(name: str, ext: str) -> Path:
    name = re.sub(r"[^a-zA-Z0-9 _.-]", "", name).strip() or "Nova_File"
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return FILES_DIR / f"{name}_{stamp}.{ext}"

# -------------------- Voice Listening --------------------
def listen_once(timeout=8, phrase_time_limit=8) -> str:
    if not sr:
        speak("Speech recognition package is not installed.")
        return ""
    recognizer = sr.Recognizer()
    try:
        with sr.Microphone() as source:
            log("Listening...")
            recognizer.adjust_for_ambient_noise(source, duration=0.5)
            audio = recognizer.listen(source, timeout=timeout, phrase_time_limit=phrase_time_limit)
        text = recognizer.recognize_google(audio, language=settings.get("language", "en-IN"))
        text = normalize_words(text)
        log(f"You: {text}")
        return text
    except sr.WaitTimeoutError:
        return ""
    except sr.UnknownValueError:
        return ""
    except Exception as e:
        log(f"Listen error: {e}")
        return ""

# -------------------- Human-like AI Brain --------------------
def api_key(name):
    return os.getenv(name, "").strip()


def setting_value(key, default=""):
    env_name = key.upper()
    return os.getenv(env_name, "").strip() or str(settings.get(key, default)).strip()


def available_providers():
    checks = {
        "gemini": "GEMINI_API_KEY",
        "openai": "OPENAI_API_KEY",
        "groq": "GROQ_API_KEY",
        "openrouter": "OPENROUTER_API_KEY",
    }
    return [name for name, key in checks.items() if api_key(key)]


def api_status():
    providers = available_providers()
    network = network_status()
    python_http = bool(network.get("python_http", False))
    models = {
        "gemini": setting_value("gemini_model", "gemini-2.5-flash"),
        "openai": setting_value("openai_model", "gpt-4.1-mini"),
        "groq": setting_value("groq_model", "llama-3.1-8b-instant"),
        "openrouter": setting_value("openrouter_model", "meta-llama/llama-3.1-8b-instruct:free"),
    }
    ordered = [p for p in settings.get("ai_provider_order", []) if p in providers]
    ordered += [p for p in providers if p not in ordered]
    return {
        "online": bool(ordered) and requests is not None and python_http,
        "python_http": python_http,
        "network_message": network.get("message", ""),
        "providers": providers,
        "order": ordered,
        "models": models,
        "requests_installed": requests is not None,
        "laptop_connected": bool(settings.get("laptop_connected", False)) or LAPTOP_PROFILE_FILE.exists(),
        "laptop_profile": str(LAPTOP_PROFILE_FILE),
        "control_mode": bool(settings.get("control_mode", True)),
        "local_control_permission": bool(settings.get("local_control_permission", False)),
        "google_browser_permission": bool(settings.get("google_browser_permission", False)),
        "gmail_browser_permission": bool(settings.get("gmail_browser_permission", False)),
        "tor_browser_permission": bool(settings.get("tor_browser_permission", False)),
        "local_only_mode": bool(settings.get("local_only_mode", True)),
        "assistant_enabled": bool(settings.get("assistant_enabled", True)),
        "access_lock_enabled": access_lock_enabled(),
    }


def setup_needs():
    status = api_status()
    needs = []
    if not status["laptop_connected"]:
        needs.append({
            "id": "connect_laptop",
            "title": "Connect laptop profile",
            "message": "Nova needs permission to scan local apps, folders, drives, and settings on this laptop only.",
            "confirmCommand": "confirm repair assistant",
        })
    if not status["local_control_permission"]:
        needs.append({
            "id": "local_control_permission",
            "title": "Allow local control",
            "message": "Nova needs local mouse, keyboard, clipboard, and window-control permission for clear laptop commands.",
            "confirmCommand": "confirm repair assistant",
        })
    if not pyautogui or not pyperclip or not gw or not pil_available():
        missing = []
        if not pyautogui:
            missing.append("pyautogui")
        if not pyperclip:
            missing.append("pyperclip")
        if not gw:
            missing.append("pygetwindow")
        if not pil_available():
            missing.append("Pillow")
        needs.append({
            "id": "python_control_packages",
            "title": "Install control packages",
            "message": "Nova is missing local Python control packages: " + ", ".join(missing) + ". Run INSTALL_AND_RUN.bat or pip install requirements.txt.",
            "confirmCommand": "",
        })
    if not status["online"]:
        needs.append({
            "id": "online_brain_key",
            "title": "Optional online thinking",
            "message": status.get("network_message") or "For smarter Gemini/ChatGPT-style thinking, add an API key in .env. Laptop control still works offline.",
            "confirmCommand": "",
        })
    return needs


FEATURE_TOGGLES = [
    {
        "key": "assistant_enabled",
        "label": "Assistant Core",
        "description": "Allows Nova to answer and execute commands. Turning it off pauses most assistant actions.",
    },
    {
        "key": "control_mode",
        "label": "Laptop Control",
        "description": "Allows keyboard, mouse, window, and local app control for clear commands.",
    },
    {
        "key": "voice_enabled",
        "label": "Voice Output",
        "description": "Allows Nova to speak answers using Windows voice.",
    },
    {
        "key": "spotify_control_enabled",
        "label": "Spotify Control",
        "description": "Allows Nova to open Spotify, search songs, and press play/media keys.",
    },
    {
        "key": "whatsapp_control_enabled",
        "label": "WhatsApp Control",
        "description": "Allows Nova to open WhatsApp Desktop and best-effort contact search/call automation.",
    },
    {
        "key": "chatgpt_features_enabled",
        "label": "ChatGPT Tools",
        "description": "Allows copy/read/like/fork/regenerate/export chat tools and ChatGPT-style chat helpers.",
    },
    {
        "key": "media_studio_enabled",
        "label": "Media Studio",
        "description": "Allows photo/video upload, analyze, create, crop, edit, and merge tools.",
    },
    {
        "key": "file_studio_enabled",
        "label": "File Studio",
        "description": "Allows file upload, file creation, and file analysis tools.",
    },
    {
        "key": "security_tools_enabled",
        "label": "Nova Shield",
        "description": "Allows defensive Windows Security, Defender, firewall, and protection status commands.",
    },
    {
        "key": "autopilot_enabled",
        "label": "Autopilot",
        "description": "Allows Nova to split one command into safe ordered steps and run them automatically.",
    },
    {
        "key": "plan_mode_enabled",
        "label": "Plan Mode",
        "description": "Makes Nova answer with a plan first instead of executing the command.",
    },
    {
        "key": "pursue_goal_mode_enabled",
        "label": "Pursue Goal",
        "description": "Allows Nova to use safe autopilot behavior for clear multi-step goals.",
    },
    {
        "key": "auto_scan_on_start",
        "label": "Startup Learning Scan",
        "description": "Refreshes Nova's local app, setting, folder, and capability profile when Nova starts.",
    },
    {
        "key": "privacy_guard_enabled",
        "label": "Privacy Guard",
        "description": "Redacts owner details, keys, emails, phone numbers, and local paths before online AI calls.",
    },
    {
        "key": "dark_web_safety_enabled",
        "label": "Dark Web Safety",
        "description": "Allows defensive dark-web safety explanations and leak-risk guidance, never illegal links or evasion.",
    },
    {
        "key": "google_browser_permission",
        "label": "Google Browser",
        "description": "Allows Nova to open Google services in your already logged-in browser session.",
    },
    {
        "key": "gmail_browser_permission",
        "label": "Gmail Browser",
        "description": "Allows Nova to open/search/compose Gmail drafts in your browser session. Sending still needs your final review.",
    },
    {
        "key": "tor_browser_permission",
        "label": "Tor Browser",
        "description": "Allows Nova to open Tor Browser for explicit legal privacy browsing/search commands only.",
    },
    {
        "key": "local_only_mode",
        "label": "Local Only",
        "description": "Keeps Nova focused on this laptop and avoids deployment/external storage behavior.",
    },
    {
        "key": "confirm_dangerous_actions",
        "label": "Danger Confirm",
        "description": "Requires confirmation before risky actions like uninstall, delete, shutdown, restart, or security changes.",
    },
]


def feature_settings_payload():
    return {
        "items": [
            {
                "key": item["key"],
                "label": item["label"],
                "description": item["description"],
                "enabled": bool(settings.get(item["key"], DEFAULT_SETTINGS.get(item["key"], False))),
            }
            for item in FEATURE_TOGGLES
        ]
    }


def set_feature_setting(key, enabled):
    allowed = {item["key"] for item in FEATURE_TOGGLES}
    if key not in allowed:
        return False, "Unknown feature setting."
    enabled = bool(enabled)
    if key == "assistant_enabled":
        reply = set_assistant_power(enabled)
        return True, reply
    settings[key] = enabled
    if enabled and key == "plan_mode_enabled":
        settings["pursue_goal_mode_enabled"] = False
    if enabled and key == "pursue_goal_mode_enabled":
        settings["plan_mode_enabled"] = False
    if key == "control_mode":
        settings["local_control_permission"] = enabled
    settings["feature_settings_updated_at"] = datetime.datetime.now().isoformat()
    save_json(SETTINGS_FILE, settings)
    state = "ON" if enabled else "OFF"
    label = next((item["label"] for item in FEATURE_TOGGLES if item["key"] == key), key)
    return True, f"{label} is {state}."


def windows_internet_connected():
    if platform.system().lower() != "windows":
        return False
    try:
        flags = ctypes.c_ulong()
        return bool(ctypes.windll.wininet.InternetGetConnectedState(ctypes.byref(flags), 0))
    except Exception:
        return False


def network_status(force: bool = False):
    now = time.time()
    if not force and _NETWORK_CACHE.get("value") and now - float(_NETWORK_CACHE.get("time", 0)) < 45:
        return dict(_NETWORK_CACHE["value"])
    if not requests:
        value = {"online": False, "python_http": False, "message": "Python requests package is missing."}
        _NETWORK_CACHE.update({"time": now, "value": value})
        return dict(value)
    checks = [
        "https://www.google.com/generate_204",
        "https://www.gstatic.com/generate_204",
    ]
    errors = []
    for url in checks:
        try:
            r = requests.get(url, timeout=5)
            if r.status_code in [200, 204]:
                value = {"online": True, "python_http": True, "message": "Network is reachable."}
                _NETWORK_CACHE.update({"time": now, "value": value})
                return dict(value)
        except Exception as e:
            errors.append(str(e))
    if windows_internet_connected():
        short = "; ".join(errors[-2:])
        value = {
            "online": True,
            "python_http": False,
            "message": "Windows internet is connected, but Python/API HTTP is blocked or filtered. Browser may work; API calls may need firewall/security allow. Detail: " + short,
        }
        _NETWORK_CACHE.update({"time": now, "value": value})
        return dict(value)
    value = {"online": False, "python_http": False, "message": "Network check failed: " + "; ".join(errors[-2:])}
    _NETWORK_CACHE.update({"time": now, "value": value})
    return dict(value)


def remember_history(role, content):
    memory.setdefault("history", [])
    memory["history"].append({"role": role, "content": content, "time": datetime.datetime.now().isoformat()})
    memory["history"] = memory["history"][-200:]
    save_json(MEMORY_FILE, memory)


def record_interaction(user_text, assistant_text):
    memory.setdefault("history", [])
    memory["history"].append({
        "role": "user",
        "content": user_text,
        "time": datetime.datetime.now().isoformat(),
    })
    memory["history"].append({
        "role": "assistant",
        "content": assistant_text,
        "time": datetime.datetime.now().isoformat(),
    })
    memory["history"] = memory["history"][-200:]
    save_json(MEMORY_FILE, memory)


def recent_history_pairs(limit=12):
    hist = memory.get("history", [])[-limit * 2:]
    pairs = []
    last_user = None
    for item in hist:
        role = item.get("role")
        content = item.get("content", "")
        if role == "user":
            last_user = content
        elif role == "assistant" and last_user:
            pairs.append({"user": last_user, "assistant": content, "time": item.get("time", "")})
            last_user = None
    return pairs[-limit:]


def redact_private_text(value):
    text = str(value or "")
    if not settings.get("privacy_guard_enabled", True):
        return text
    replacements = [
        (re.compile(re.escape(str(Path.home())), re.I), "[HOME]"),
        (re.compile(r"C:\\Users\\[^\\\s]+", re.I), "[USER_PATH]"),
        (re.compile(r"[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}"), "[EMAIL]"),
        (re.compile(r"\b(?:\+?91[-\s]?)?[6-9]\d{9}\b"), "[PHONE]"),
        (re.compile(r"\b((?:api[_-]?key|token|secret|password|pass|otp|pin))\s*[:=]\s*[\w.\-+/=]{6,}", re.I), r"\1=[REDACTED]"),
        (re.compile(r"AIza[0-9A-Za-z_\-]{20,}"), "[GEMINI_KEY]"),
        (re.compile(r"sk-[A-Za-z0-9_\-]{20,}"), "[API_KEY]"),
    ]
    if USER_NAME:
        replacements.append((re.compile(re.escape(USER_NAME), re.I), "[OWNER]"))
    for pattern, repl in replacements:
        text = pattern.sub(repl, text)
    return text


def privacy_status_text():
    return (
        "Privacy Guard is " + ("ON" if settings.get("privacy_guard_enabled", True) else "OFF") + ".\n"
        "- Local command history stays in memory.json on this laptop.\n"
        "- Online AI calls get redacted text when Privacy Guard is ON.\n"
        "- I will not send laptop-owner details, local paths, API keys, emails, phone numbers, PINs, passwords, or OTPs intentionally.\n"
        "- Tor Browser permission is " + ("ON" if settings.get("tor_browser_permission", False) else "OFF") + ". It is used only when you clearly ask for Tor/private browsing.\n"
        "- I also will not use dark-web sources or hidden services for illegal searching."
    )


def save_feedback(kind, user_text="", assistant_text=""):
    memory.setdefault("feedback", [])
    memory["feedback"].append({
        "kind": kind,
        "user": user_text,
        "assistant": assistant_text,
        "time": datetime.datetime.now().isoformat(),
    })
    memory["feedback"] = memory["feedback"][-100:]
    save_json(MEMORY_FILE, memory)


def chatgpt_feature_help():
    return (
        "ChatGPT-style tools are ON.\n\n"
        "- Ask normal questions and coding questions with proper line-by-line code blocks.\n"
        "- Use answer buttons for copy, read/stop, like, dislike, fork, and regenerate.\n"
        "- Say or type: new chat, export chat, summarize chat, copy last answer, read last answer, regenerate last answer.\n"
        "- Laptop actions still work separately: open apps, settings, files, Spotify, WhatsApp best-effort, screenshots, volume, and typing."
    )


def last_history_pair():
    pairs = recent_history_pairs(1)
    return pairs[-1] if pairs else {}


def start_new_chat_session():
    memory["current_session_id"] = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    save_json(MEMORY_FILE, memory)
    return "New chat started. Old chat history is still saved in memory and History."


def export_chat_history():
    hist = memory.get("history", [])[-200:]
    if not hist:
        return False, "No chat history to export yet."
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = DRAFTS_DIR / f"Nova_ChatGPT_Style_History_{stamp}.md"
    lines = [f"# {APP_NAME} Chat History", "", f"Exported: {datetime.datetime.now()}", ""]
    for item in hist:
        role = str(item.get("role", "unknown")).title()
        when = item.get("time", "")
        content = str(item.get("content", "")).strip()
        if not content:
            continue
        lines.extend([f"## {role}", f"`{when}`" if when else "", "", content, ""])
    path.write_text("\n".join(lines), encoding="utf-8")
    try:
        os.startfile(path)
    except Exception:
        pass
    return True, f"Exported chat history to {path.name} in Desktop/Nova AI Files/Drafts."


def summarize_chat_history():
    pairs = recent_history_pairs(int(settings.get("chat_context_limit", 12)))
    if not pairs:
        return "No chat history yet. Start a conversation first."
    bullets = []
    for pair in pairs[-8:]:
        user = pair.get("user", "").strip().replace("\n", " ")
        assistant = pair.get("assistant", "").strip().replace("\n", " ")
        bullets.append(f"- You: {user[:120]}\n  Nova: {assistant[:160]}")
    return "Recent chat summary:\n" + "\n".join(bullets)


def copy_last_answer():
    pair = last_history_pair()
    answer = pair.get("assistant", "")
    if not answer:
        return False, "No assistant answer to copy yet."
    if not pyperclip:
        return False, "Clipboard package is missing, so I cannot copy it yet."
    pyperclip.copy(answer)
    return True, "Copied the last assistant answer to clipboard."


def read_last_answer():
    pair = last_history_pair()
    answer = pair.get("assistant", "")
    if not answer:
        return False, "No assistant answer to read yet."
    ok = speak(answer)
    return ok, "Reading the last assistant answer." if ok else "Voice output is not working. Try test voice or run INSTALL_AUDIO_FIX.cmd."


def regenerate_last_answer():
    pair = last_history_pair()
    question = pair.get("user", "")
    if not question:
        return "No previous user message found to regenerate."
    return ai_reply(question)


def fork_last_answer():
    pair = last_history_pair()
    if not pair:
        return False, "No answer to fork yet."
    memory.setdefault("forks", [])
    memory["forks"].append({
        "user": pair.get("user", ""),
        "assistant": pair.get("assistant", ""),
        "time": datetime.datetime.now().isoformat(),
    })
    memory["forks"] = memory["forks"][-50:]
    save_json(MEMORY_FILE, memory)
    return True, "Fork saved. You can continue from that answer in the chat box."


def system_prompt():
    app_count = 0
    laptop_connected = bool(settings.get("laptop_connected", False))
    try:
        if APP_INDEX_FILE.exists():
            app_count = int(json.loads(APP_INDEX_FILE.read_text(encoding="utf-8")).get("apps_count", 0))
            laptop_connected = laptop_connected or LAPTOP_PROFILE_FILE.exists()
    except Exception:
        app_count = 0
    recent = recent_history_pairs(5)
    recent_text = ""
    if recent:
        recent_text = " Recent chat memory: " + " | ".join(
            f"User: {redact_private_text(p.get('user',''))[:120]} Assistant: {redact_private_text(p.get('assistant',''))[:160]}" for p in recent
        )
    playbook_text = app_playbook_summary()
    owner_label = USER_NAME if not settings.get("privacy_guard_enabled", True) else "the laptop owner"
    return (
        f"You are {APP_NAME}, a human-like Windows laptop assistant for {owner_label}. "
        "Always reply in clear English only, even when the user speaks Hindi or Hinglish. "
        "You may understand Hindi/Hinglish commands, but your final answer and voice output must stay English. "
        f"You remember local chat/search history and the learned laptop profile. Laptop connected: {laptop_connected}. Learned apps count: {app_count}. "
        f"Laptop Control Mode is {'on' if settings.get('control_mode', True) else 'off'}; when it is on, action commands should be short, direct, and honest. "
        f"Local-only mode is {'on' if settings.get('local_only_mode', True) else 'off'}; do not claim anything is deployed, uploaded, or stored outside this laptop. "
        f"Privacy Guard is {'on' if settings.get('privacy_guard_enabled', True) else 'off'}; do not ask for or reveal private owner details, secrets, passwords, OTPs, API keys, emails, phone numbers, or local paths. "
        "Think before answering, but never expose hidden chain-of-thought. "
        "If the user asks a normal question, answer the question directly. "
        "Before suggesting laptop actions, use the local laptop profile and app/website playbooks. "
        "For commands, think in this order: identify target app/site/setting, check local capability, execute the safest specific action, verify when possible, and say the exact limitation if blocked. "
        "If the user asks for code, write clean code in fenced code blocks with proper line breaks and indentation. "
        "Never put code in one mixed paragraph. "
        "For code requests, use this structure: one short title, one fenced code block, then short usage notes if needed. "
        "For media tasks, be honest: local image analyze/edit/generate can run in Media Studio; real semantic vision needs Gemini; MP4 video crop/merge needs FFmpeg. "
        "If the user asks a laptop action, do not pretend you performed it; local code performs actions before you are called. "
        "Never answer with random future-tense lines like 'I will open' unless you are explaining a limitation. "
        "You can help with laptop tasks, apps, files, WhatsApp guidance, and normal questions. "
        "Never claim you performed a desktop action unless the local action handler actually did it. "
        "Be honest in the application UI: if an action is best-effort automation, say exactly what was tried and what the user may need to click manually. "
        "For WhatsApp calls, never promise that the call has started; say that the chat/call button was tried because WhatsApp Desktop can block automation. "
        "You cannot bypass Windows passwords, UAC, app login, app permissions, or missing apps. Say this clearly when needed. "
        "For dangerous tasks like uninstall/delete/shutdown, remind that user confirmation is required."
        "Tor Browser may be opened only for explicit legal privacy browsing/search commands. "
        "For dark web topics, provide defensive education only: risks, scams, malware, leaks, privacy basics, and security steps. Do not provide illegal marketplaces, hidden-service links, credential access, exploit steps, evasion, or instructions for wrongdoing."
        + playbook_text
        + recent_text
    )


def call_gemini(message):
    if not requests or not api_key("GEMINI_API_KEY"):
        return None
    message = redact_private_text(message)
    model = setting_value("gemini_model", "gemini-2.5-flash")
    url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
    payload = {
        "system_instruction": {"parts": [{"text": system_prompt()}]},
        "contents": [{"role": "user", "parts": [{"text": message}]}],
        "generationConfig": {"temperature": 0.65, "maxOutputTokens": 600},
    }
    r = requests.post(
        url,
        headers={"x-goog-api-key": api_key("GEMINI_API_KEY"), "Content-Type": "application/json"},
        json=payload,
        timeout=35,
    )
    data = r.json()
    if r.status_code >= 400:
        raise RuntimeError(data.get("error", {}).get("message", "Gemini API error"))
    return data["candidates"][0]["content"]["parts"][0]["text"].strip()


def call_openai(message):
    if not requests or not api_key("OPENAI_API_KEY"):
        return None
    message = redact_private_text(message)
    r = requests.post(
        "https://api.openai.com/v1/responses",
        headers={"Authorization": f"Bearer {api_key('OPENAI_API_KEY')}", "Content-Type": "application/json"},
        json={"model": setting_value("openai_model", "gpt-4.1-mini"), "input": [{"role": "system", "content": system_prompt()}, {"role": "user", "content": message}], "max_output_tokens": 600},
        timeout=25,
    )
    data = r.json()
    if r.status_code >= 400:
        raise RuntimeError(data.get("error", {}).get("message", "OpenAI API error"))
    return data.get("output_text", "").strip()


def call_groq(message):
    if not requests or not api_key("GROQ_API_KEY"):
        return None
    message = redact_private_text(message)
    r = requests.post(
        "https://api.groq.com/openai/v1/chat/completions",
        headers={"Authorization": f"Bearer {api_key('GROQ_API_KEY')}", "Content-Type": "application/json"},
        json={"model": setting_value("groq_model", "llama-3.1-8b-instant"), "messages": [{"role": "system", "content": system_prompt()}, {"role": "user", "content": message}], "temperature": 0.65, "max_tokens": 600},
        timeout=25,
    )
    data = r.json()
    if r.status_code >= 400:
        raise RuntimeError(data.get("error", {}).get("message", "Groq API error"))
    return data["choices"][0]["message"]["content"].strip()


def call_openrouter(message):
    if not requests or not api_key("OPENROUTER_API_KEY"):
        return None
    message = redact_private_text(message)
    r = requests.post(
        "https://openrouter.ai/api/v1/chat/completions",
        headers={"Authorization": f"Bearer {api_key('OPENROUTER_API_KEY')}", "Content-Type": "application/json", "HTTP-Referer": "http://localhost", "X-Title": APP_NAME},
        json={"model": setting_value("openrouter_model", "meta-llama/llama-3.1-8b-instruct:free"), "messages": [{"role": "system", "content": system_prompt()}, {"role": "user", "content": message}], "temperature": 0.65, "max_tokens": 600},
        timeout=25,
    )
    data = r.json()
    if r.status_code >= 400:
        raise RuntimeError(data.get("error", {}).get("message", "OpenRouter API error"))
    return data["choices"][0]["message"]["content"].strip()


def ai_reply(message):
    remember_history("user", message)
    providers = {
        "gemini": call_gemini,
        "openai": call_openai,
        "groq": call_groq,
        "openrouter": call_openrouter,
    }
    errors = []
    status = api_status()
    if not status["requests_installed"]:
        return offline_reply(message, ["requests package is not installed"])
    if not status["order"]:
        return offline_reply(message, ["No API key found in .env. Add GEMINI_API_KEY, OPENAI_API_KEY, GROQ_API_KEY, or OPENROUTER_API_KEY."])
    if not status.get("python_http", False):
        return offline_reply(message, [status.get("network_message") or "Python/API internet is blocked or filtered."])
    for p in status["order"]:
        try:
            reply = providers[p](message)
            if reply:
                reply = improve_code_blocks(reply, message)
                remember_history("assistant", reply)
                return reply
        except Exception as e:
            errors.append(f"{p}: {e}")
            log(f"AI provider failed: {p}: {e}")
    return offline_reply(message, errors)


def offline_reply(c, errors=None):
    c = normalize_words(c)
    if re.search(r"\b(hello|hi|hey)\b", c):
        return f"Hello {USER_NAME}, I am ready. Tell me what you want me to do."
    if "your name" in c or "who are you" in c:
        return "I am Nova AI Ultimate, your local Windows assistant."
    if "what can you do" in c or "help" == c:
        return "I can open and close apps, search Google or YouTube, type text, make files, take screenshots, control volume, manage windows, edit/analyze photos and videos through Media Studio, and chat using Gemini or other APIs."
    if "api" in c or "key" in c:
        return "Put your API keys in the .env file. Use GEMINI_API_KEY, OPENAI_API_KEY, GROQ_API_KEY, or OPENROUTER_API_KEY."
    if errors:
        short = "; ".join(str(e) for e in errors[-2:])
        return f"Online brain failed, so I answered offline. Problem: {short}"
    return "I understood you, but this exact command is not added yet. Try open chrome, play song on YouTube, make Word file, screenshot, type hello, or ask me a question."


def force_english_reply(reply):
    if reply == "EXIT" or not isinstance(reply, str) or not settings.get("force_english_output", True):
        return reply
    text = reply
    phrase_replacements = [
        (
            "Generic Spotify play samjha: available playlist open karke Play ko sirf ek baar press karne ki koshish kar raha hoon. Ab extra Play/Pause key nahi bhejunga, taaki song start hote hi stop na ho.",
            "Generic Spotify play understood: I am opening an available playlist and trying to press Play only once. I will not send the extra Play/Pause key, so the song should not stop immediately after starting.",
        ),
        ("karne ki koshish kar raha hoon", "trying to do it"),
        ("kar raha hoon", "trying"),
        ("kar diya", "done"),
        ("set kar diya", "set"),
        ("open kar diya", "opened"),
        ("khol raha hoon", "opening"),
        ("khul gaya hai", "opened"),
        ("nahi mila", "was not found"),
        ("nahi mili", "was not found"),
        ("nahi hua", "did not work"),
        ("nahi ho paya", "could not be done"),
        ("nahi karunga", "will not do it"),
        ("nahi bolunga", "will not say it"),
        ("nahi bhejunga", "will not send"),
        ("ke liye", "for"),
        ("hote hi", "immediately after"),
        ("sirf ek baar", "only once"),
        ("start hote hi stop na ho", "should not stop immediately after starting"),
        ("stop na ho", "should not stop"),
    ]
    for src, dst in phrase_replacements:
        text = text.replace(src, dst)
    word_replacements = {
        "samjha": "understood",
        "karke": "and",
        "press": "press",
        "agar": "if",
        "Agar": "If",
        "ya": "or",
        "aur": "and",
        "bhi": "also",
        "hai": "is",
        "mein": "in",
        "me": "in",
        "par": "on",
        "ko": "",
        "Ab": "Now",
        "taaki": "so that",
        "Pehle": "First",
        "karo": "",
        "bolo": "say",
    }
    for src, dst in word_replacements.items():
        text = re.sub(rf"\b{re.escape(src)}\b", dst, text)
    text = improve_code_blocks(text)
    parts = re.split(r"(```[\s\S]*?```)", text)
    cleaned = []
    for idx, part in enumerate(parts):
        if idx % 2 == 1:
            cleaned.append(part)
        else:
            part = re.sub(r"[ \t]{2,}", " ", part)
            part = re.sub(r"\n{3,}", "\n\n", part)
            cleaned.append(part)
    return "".join(cleaned).strip()


def calculator_code_reply():
    return (
        "Basic Calculator in Python\n\n"
        "```python\n"
        "def calculator():\n"
        "    print(\"Calculator Program\")\n"
        "    print(\"1. Addition\")\n"
        "    print(\"2. Subtraction\")\n"
        "    print(\"3. Multiplication\")\n"
        "    print(\"4. Division\")\n"
        "\n"
        "    choice = input(\"Enter your choice (1/2/3/4): \")\n"
        "\n"
        "    if choice in (\"1\", \"2\", \"3\", \"4\"):\n"
        "        num1 = float(input(\"Enter first number: \"))\n"
        "        num2 = float(input(\"Enter second number: \"))\n"
        "\n"
        "        if choice == \"1\":\n"
        "            print(num1, \"+\", num2, \"=\", num1 + num2)\n"
        "        elif choice == \"2\":\n"
        "            print(num1, \"-\", num2, \"=\", num1 - num2)\n"
        "        elif choice == \"3\":\n"
        "            print(num1, \"*\", num2, \"=\", num1 * num2)\n"
        "        elif choice == \"4\":\n"
        "            if num2 != 0:\n"
        "                print(num1, \"/\", num2, \"=\", num1 / num2)\n"
        "            else:\n"
        "                print(\"Error! Division by zero is not allowed.\")\n"
        "    else:\n"
        "        print(\"Invalid choice. Please choose a valid option.\")\n"
        "\n"
        "\n"
        "calculator()\n"
        "```\n\n"
        "Run it with: `python calculator.py`"
    )


def _protect_python_strings(code):
    strings = []

    def repl(match):
        strings.append(match.group(0))
        return f"__NOVA_STRING_{len(strings) - 1}__"

    pattern = r"(?s)(\"\"\".*?\"\"\"|'''.*?'''|\"(?:\\.|[^\"\\])*\"|'(?:\\.|[^'\\])*')"
    return re.sub(pattern, repl, code), strings


def _restore_python_strings(code, strings):
    for idx, value in enumerate(strings):
        code = code.replace(f"__NOVA_STRING_{idx}__", value)
    return code


def expand_one_line_python(code):
    """Best-effort repair for AI replies that put a Python program on one line."""
    original = code.strip()
    if "\n" in original or len(original) < 60:
        return code
    if not re.search(r"\b(def|if|elif|else|for|while|try|except|class|print|return|input)\b", original):
        return code

    safe, strings = _protect_python_strings(original)
    safe = re.sub(r"\s+", " ", safe).strip()

    safe = re.sub(r"\s+(?=(?:def|class|if|elif|else:|for|while|try:|except|return)\b)", "\n", safe)
    safe = re.sub(r"\s+(?=print\()", "\n", safe)
    safe = re.sub(r"\s+(?=[A-Za-z_]\w*\s=(?!=))", "\n", safe)
    safe = re.sub(r"\s+([A-Za-z_]\w*\(\))\s*$", r"\n\1", safe)
    safe = safe.replace(": print(", ":\nprint(")
    safe = safe.replace(": input(", ":\ninput(")
    safe = safe.replace(": return ", ":\nreturn ")
    safe = safe.replace(": if ", ":\nif ")

    raw_lines = [line.strip() for line in safe.splitlines() if line.strip()]
    lines = []
    indent = 0
    prev = ""
    for line in raw_lines:
        if line.startswith(("elif ", "else:", "except ", "finally:")):
            indent = max(0, indent - 1)
        if prev.endswith(":") and not line.startswith(("elif ", "else:", "except ", "finally:")):
            indent += 1
        lines.append("    " * indent + line)
        prev = line
    repaired = "\n".join(lines)
    repaired = _restore_python_strings(repaired, strings)

    try:
        compile(repaired, "<nova-code-repair>", "exec")
        return repaired
    except Exception:
        return repaired if len(repaired.splitlines()) >= 4 else code


def improve_code_blocks(reply, original_message=""):
    if not isinstance(reply, str) or "```" not in reply:
        return reply

    def repl(match):
        lang = (match.group(1) or "").strip().lower()
        code = match.group(2).strip()
        likely_python = lang in ["python", "py"] or (
            not lang and re.search(r"\b(def|print|input|elif|calculator\(\))\b", code)
        )
        if likely_python:
            if "def calculator" in code and "division" in code.lower() and "choice" in code:
                code = calculator_code_reply().split("```python\n", 1)[1].split("\n```", 1)[0]
            else:
                code = expand_one_line_python(code)
            lang = "python"
        fence_lang = lang or ""
        return "```" + fence_lang + "\n" + code.strip() + "\n```"

    text = re.sub(r"```([a-zA-Z0-9_+-]*)\s*\n?([\s\S]*?)```", repl, reply)

    # If an AI forgot fences but clearly returned a one-line Python calculator, repair it.
    if "```" not in text and "calculator" in normalize_words(original_message) and "def calculator" in text:
        code = expand_one_line_python(text)
        if code != text:
            return "Basic Calculator in Python\n\n```python\n" + code + "\n```"
    return text


# -------------------- Windows Actions --------------------
def run_start(target: str) -> bool:
    try:
        # Use cmd/start as a list so spaces in URI/app names do not get split.
        subprocess.Popen(["cmd", "/c", "start", "", target])
        return True
    except Exception as e:
        log(f"start error: {e}")
        return False


def find_known_executable(name):
    name = name.lower().strip()
    roots = [
        Path(os.environ.get("LOCALAPPDATA", "")),
        Path(os.environ.get("PROGRAMFILES", "")),
        Path(os.environ.get("PROGRAMFILES(X86)", "")),
    ]
    known = {
        "brave": [
            "BraveSoftware/Brave-Browser/Application/brave.exe",
            "BraveSoftware/Brave-Browser-Nightly/Application/brave.exe",
        ],
        "tor": [
            "Tor Browser/Browser/firefox.exe",
            "TorBrowser/Browser/firefox.exe",
        ],
        "tor browser": [
            "Tor Browser/Browser/firefox.exe",
            "TorBrowser/Browser/firefox.exe",
        ],
        "whatsapp": [
            "WhatsApp/WhatsApp.exe",
            "Packages/5319275A.WhatsAppDesktop_cv1g1gvanyjgm/LocalCache/Roaming/WhatsApp/WhatsApp.exe",
        ],
    }
    for rel in known.get(name, []):
        for root in roots:
            p = root / rel
            if p.exists():
                return p
    if name in ["tor", "tor browser"]:
        for p in [
            DESKTOP / "Tor Browser" / "Browser" / "firefox.exe",
            Path.home() / "Downloads" / "Tor Browser" / "Browser" / "firefox.exe",
            BASE_DIR / "tools" / "Tor Browser" / "Browser" / "firefox.exe",
        ]:
            if p.exists():
                return p
    return None


def executable_from_profile(name):
    name_l = clean_entity_name(name).lower()
    if not LAPTOP_PROFILE_FILE.exists() or not name_l:
        return None
    try:
        profile = json.loads(LAPTOP_PROFILE_FILE.read_text(encoding="utf-8"))
        programs = profile.get("installed_programs", [])
        matches = []
        for item in programs:
            label = str(item.get("name", "")).lower()
            if label == name_l or name_l in label or label in name_l:
                matches.append(item)
        if not matches and len(name_l) >= 5:
            names = [str(item.get("name", "")).lower() for item in programs]
            close = difflib.get_close_matches(name_l, names, n=1, cutoff=0.72)
            if close:
                matches = [item for item in programs if str(item.get("name", "")).lower() == close[0]]
        for item in matches:
            icon = str(item.get("DisplayIcon", "")).strip().strip('"')
            icon = re.sub(r",\d+$", "", icon).strip().strip('"')
            if icon.lower().endswith(".exe") and Path(icon).exists():
                return Path(icon)
            folder = str(item.get("InstallLocation", "")).strip().strip('"')
            if folder and Path(folder).exists():
                exe_files = sorted(Path(folder).glob("*.exe"), key=lambda p: (name_l not in p.stem.lower(), len(p.name)))
                if exe_files:
                    return exe_files[0]
    except Exception as e:
        log(f"profile executable lookup error: {e}")
    return None


def shortcut_dirs():
    dirs = [
        ("Start Menu", Path(os.environ.get("APPDATA", "")) / "Microsoft/Windows/Start Menu/Programs"),
        ("Start Menu", Path(os.environ.get("PROGRAMDATA", "")) / "Microsoft/Windows/Start Menu/Programs"),
        ("Desktop", DESKTOP),
        ("Public Desktop", Path(os.environ.get("PUBLIC", "")) / "Desktop"),
    ]
    onedrive = os.environ.get("OneDrive") or os.environ.get("ONEDRIVE")
    if onedrive:
        dirs.append(("OneDrive Desktop", Path(onedrive) / "Desktop"))
    return dirs


def start_menu_dirs():
    return [path for source, path in shortcut_dirs() if source == "Start Menu"]


def scan_start_menu_apps():
    items = []
    for source, d in shortcut_dirs():
        if d.exists():
            for pattern in ["*.lnk", "*.url"]:
                for p in d.rglob(pattern):
                    # Keep both raw shortcut name and parent folder words for better matching.
                    label = (p.stem + " " + p.parent.stem).lower().strip()
                    items.append((p.stem.lower(), p))
                    if label != p.stem.lower():
                        items.append((label, p))
    return items


def scan_windows_start_apps(limit=1500):
    """Scan Windows Start apps, including Store apps and installed PWAs."""
    if platform.system().lower() != "windows":
        return []
    try:
        result = subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-Command",
                "Get-StartApps | Select-Object Name,AppID | ConvertTo-Json -Compress",
            ],
            capture_output=True,
            text=True,
            timeout=18,
        )
        if result.returncode != 0:
            log("Get-StartApps failed: " + ((result.stderr or result.stdout or "").strip()[:300]))
            return []
        raw = (result.stdout or "").strip()
        if not raw:
            return []
        data = json.loads(raw)
        if isinstance(data, dict):
            data = [data]
        apps = []
        seen = set()
        for item in data:
            name = str(item.get("Name", "")).strip()
            app_id = str(item.get("AppID", "")).strip()
            key = (name.lower(), app_id.lower())
            if not name or not app_id or key in seen:
                continue
            seen.add(key)
            apps.append({"name": name, "app_id": app_id, "source": "start_app"})
        return sorted(apps, key=lambda x: x["name"].lower())[:limit]
    except Exception as e:
        log(f"Get-StartApps scan error: {e}")
        return []


def learn_laptop():
    return startup_system_scan(force=True)


def installed_programs(limit=400):
    programs = []
    if platform.system().lower() != "windows":
        return programs
    try:
        import winreg
        roots = [
            (winreg.HKEY_CURRENT_USER, r"Software\Microsoft\Windows\CurrentVersion\Uninstall"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\Microsoft\Windows\CurrentVersion\Uninstall"),
            (winreg.HKEY_LOCAL_MACHINE, r"Software\WOW6432Node\Microsoft\Windows\CurrentVersion\Uninstall"),
        ]
        seen = set()
        for hive, subkey in roots:
            try:
                with winreg.OpenKey(hive, subkey) as key:
                    for i in range(winreg.QueryInfoKey(key)[0]):
                        try:
                            child_name = winreg.EnumKey(key, i)
                            with winreg.OpenKey(key, child_name) as child:
                                name = winreg.QueryValueEx(child, "DisplayName")[0]
                                if not name or name.lower() in seen:
                                    continue
                                seen.add(name.lower())
                                item = {"name": name}
                                for field in ["DisplayVersion", "Publisher", "InstallLocation", "DisplayIcon"]:
                                    try:
                                        item[field] = winreg.QueryValueEx(child, field)[0]
                                    except Exception:
                                        pass
                                programs.append(item)
                        except Exception:
                            continue
            except Exception:
                continue
    except Exception as e:
        log(f"installed programs scan error: {e}")
    return sorted(programs, key=lambda x: x.get("name", "").lower())[:limit]


def laptop_drives():
    drives = []
    if platform.system().lower() == "windows":
        for letter in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
            root = Path(f"{letter}:\\")
            if root.exists():
                try:
                    total, used, free = shutil.disk_usage(root)
                    drives.append({
                        "drive": str(root),
                        "total_gb": round(total / (1024 ** 3), 1),
                        "free_gb": round(free / (1024 ** 3), 1),
                    })
                except Exception:
                    drives.append({"drive": str(root)})
    return drives


def common_folders():
    candidates = {
        "desktop": DESKTOP,
        "documents": Path.home() / "Documents",
        "downloads": Path.home() / "Downloads",
        "pictures": Path.home() / "Pictures",
        "videos": Path.home() / "Videos",
        "music": Path.home() / "Music",
        "nova_files": FILES_DIR,
    }
    onedrive = os.environ.get("OneDrive") or os.environ.get("ONEDRIVE")
    if onedrive:
        candidates["onedrive"] = Path(onedrive)
        candidates["onedrive_desktop"] = Path(onedrive) / "Desktop"
        candidates["onedrive_documents"] = Path(onedrive) / "Documents"
    return {name: str(path) for name, path in candidates.items() if path.exists()}


def permission_status():
    return {
        "keyboard_mouse_automation": bool(pyautogui),
        "clipboard": bool(pyperclip),
        "window_focus": bool(gw),
        "offline_voice_recognition_package": bool(sr),
        "online_requests_package": bool(requests),
        "windows_voice": platform.system().lower() == "windows",
        "browser_mic_note": "Allow microphone in the browser address bar when the GUI asks.",
        "limits": [
            "Cannot bypass Windows password, UAC prompts, app login, or blocked app permissions.",
            "WhatsApp calls are best-effort because WhatsApp Desktop has no guaranteed public call command.",
        ],
    }


def connect_laptop():
    apps = []
    seen = set()
    for label, path in scan_start_menu_apps():
        key = str(path).lower()
        if key in seen:
            continue
        seen.add(key)
        apps.append({"name": path.stem, "path": str(path), "source": "shortcut"})
    start_apps = scan_windows_start_apps()
    profile = {
        "computer": platform.node(),
        "windows": platform.platform(),
        "connected_at": datetime.datetime.now().isoformat(),
        "apps_count": len(apps),
        "shortcuts": apps[:1500],
        "start_apps_count": len(start_apps),
        "start_apps": start_apps[:1500],
        "installed_programs": installed_programs(),
        "drives": laptop_drives(),
        "folders": common_folders(),
        "known_settings": sorted([k for k in APP_ALIASES if "setting" in k or k in ["wifi", "bluetooth", "clock", "camera", "photos"]]),
        "permissions": permission_status(),
        "network": network_status(),
    }
    save_json(LAPTOP_PROFILE_FILE, profile)
    save_json(APP_INDEX_FILE, {
        "computer": profile["computer"],
        "windows": profile["windows"],
        "learned_at": profile["connected_at"],
        "apps_count": profile["apps_count"] + profile["start_apps_count"],
        "apps": profile["shortcuts"],
        "start_apps": profile["start_apps"],
        "known_settings": profile["known_settings"],
    })
    settings["laptop_connected"] = True
    settings["last_laptop_connect"] = profile["connected_at"]
    save_json(SETTINGS_FILE, settings)
    return True, (
        f"Laptop connected locally. I indexed {profile['apps_count']} shortcuts, "
        f"{profile['start_apps_count']} Windows Start apps/PWAs, "
        f"{len(profile['installed_programs'])} installed programs, {len(profile['drives'])} drives, "
        "and common folders/settings. I will use this laptop profile before any web fallback."
    )


def _iso_age_hours(value):
    if not value:
        return None
    try:
        then = datetime.datetime.fromisoformat(str(value))
        return (datetime.datetime.now() - then).total_seconds() / 3600
    except Exception:
        return None


def startup_system_scan(force=False):
    """Refresh Nova's local capability map without reading private file contents."""
    save_app_playbooks()
    if not settings.get("auto_scan_on_start", True) and not force:
        return True, "Startup learning scan is OFF. Existing app playbooks are loaded; say 'scan laptop' to refresh manually."

    age = _iso_age_hours(settings.get("last_startup_scan_at"))
    try:
        interval = max(1, int(settings.get("startup_scan_interval_hours", 12)))
    except Exception:
        interval = 12
    stale = (
        force
        or age is None
        or age >= interval
        or not LAPTOP_PROFILE_FILE.exists()
        or not APP_INDEX_FILE.exists()
        or not APP_PLAYBOOK_FILE.exists()
    )
    if not stale:
        return True, (
            "Startup learning scan is fresh. I already have the local app/settings profile and app playbooks loaded. "
            f"Last scan was about {age:.1f} hours ago."
        )

    ok, msg = connect_laptop()
    profile = {}
    try:
        profile = json.loads(LAPTOP_PROFILE_FILE.read_text(encoding="utf-8")) if LAPTOP_PROFILE_FILE.exists() else {}
    except Exception:
        profile = {}
    now = datetime.datetime.now().isoformat()
    settings["last_startup_scan_at"] = now
    settings["last_startup_scan_summary"] = {
        "time": now,
        "shortcuts_count": profile.get("apps_count", 0),
        "start_apps_count": profile.get("start_apps_count", 0),
        "total_apps_count": int(profile.get("apps_count", 0) or 0) + int(profile.get("start_apps_count", 0) or 0),
        "programs_count": len(profile.get("installed_programs", [])),
        "drives_count": len(profile.get("drives", [])),
        "folders_count": len(profile.get("folders", {})),
        "playbooks_count": len(app_playbooks()),
    }
    save_json(SETTINGS_FILE, settings)
    if not ok:
        return False, msg
    return True, (
        msg
        + " Startup learning scan completed. I learned how to route common apps/websites through local playbooks first: "
        "Spotify, WhatsApp, YouTube, Google/Gmail, Windows Settings, files, and Media Studio. "
        "I scanned app/settings capability metadata only; I did not read all private file contents."
    )


def assistant_diagnostics(repair=False):
    if repair:
        settings["control_mode"] = True
        settings["local_control_permission"] = True
        settings["google_browser_permission"] = True
        settings["gmail_browser_permission"] = True
        settings["voice_enabled"] = True
        settings["local_only_mode"] = True
        settings["assistant_enabled"] = True
        save_json(SETTINGS_FILE, settings)
        if not LAPTOP_PROFILE_FILE.exists():
            connect_laptop()

    status = api_status()
    network = network_status()
    profile = {}
    if LAPTOP_PROFILE_FILE.exists():
        try:
            profile = json.loads(LAPTOP_PROFILE_FILE.read_text(encoding="utf-8"))
        except Exception:
            profile = {}

    app_checks = {}
    for app in ["brave", "whatsapp", "spotify", "chrome", "edge", "vs code"]:
        app_checks[app] = bool(find_start_menu_app(app) or find_known_executable(app) or executable_from_profile(app))

    missing = []
    if not status["online"]:
        missing.append("online AI API key missing or requests package unavailable")
    if not network.get("online"):
        missing.append(network.get("message", "network check failed"))
    if not pyautogui:
        missing.append("pyautogui missing, so mouse/keyboard control is limited")
    if not pyperclip:
        missing.append("pyperclip missing, so paste/copy helpers are limited")
    if not gw:
        missing.append("pygetwindow missing, so app focus/control is limited")
    if not windows_voice_names():
        missing.append("Windows voice list unavailable")
    if not (settings.get("laptop_connected") or LAPTOP_PROFILE_FILE.exists()):
        missing.append("laptop profile not connected")

    lines = [
        "Assistant Doctor",
        f"- Laptop profile: {'connected' if profile else 'not connected'} ({profile.get('apps_count', 0)} shortcuts, {profile.get('start_apps_count', 0)} Start apps/PWAs, {len(profile.get('installed_programs', []))} programs)",
        f"- Startup learning scan: {'ON' if settings.get('auto_scan_on_start', True) else 'OFF'}; last scan: {settings.get('last_startup_scan_at', 'not yet')}",
        f"- Control mode: {'ON' if settings.get('control_mode', True) else 'OFF'}",
        f"- Assistant switch: {'ON' if settings.get('assistant_enabled', True) else 'OFF'}",
        f"- Local PIN lock: {'ON' if access_lock_enabled() else 'not set'}",
        f"- Local-only mode: {'ON' if settings.get('local_only_mode', True) else 'OFF'}",
        f"- Mouse/keyboard: {'ready' if pyautogui else 'missing package'}",
        f"- Clipboard: {'ready' if pyperclip else 'missing package'}",
        f"- Window focus: {'ready' if gw else 'missing package'}",
        f"- Windows voice: {'ready' if windows_voice_names() else 'not detected'}",
        f"- Browser mic: needs Edge/Chrome permission from address bar; browser speech can still fail even when internet works",
        f"- Network: {network.get('message', 'unknown')}",
        f"- Online brain: {'ready: ' + ', '.join(status['order']) if status['online'] else 'offline, add API key in .env'}",
        "- Apps: " + ", ".join(f"{k}={'yes' if v else 'no'}" for k, v in app_checks.items()),
        "- Security: say 'security status', 'open firewall', 'open vpn settings', or 'open privacy settings'.",
    ]
    needs = setup_needs()
    if needs:
        lines.append("- Popups needed: " + "; ".join(item["title"] for item in needs[:4]))
    if missing:
        lines.append("- Needs attention: " + "; ".join(missing[:6]))
    else:
        lines.append("- Result: ready for normal laptop control.")
    if repair:
        lines.append("- Repair applied: control permissions ON, voice ON, local-only mode ON, laptop profile checked.")
    return "\n".join(lines)


def full_function_audit(repair=False):
    """Non-destructive health audit for the main Nova feature groups."""
    if repair:
        prepare_assistant_runtime()
        apply_clear_voice_settings()

    results = []

    def add(name, ok, detail=""):
        results.append((name, bool(ok), detail))

    def write_check(folder):
        try:
            folder.mkdir(parents=True, exist_ok=True)
            p = folder / ".nova_write_test.tmp"
            p.write_text("ok", encoding="utf-8")
            p.unlink(missing_ok=True)
            return True, "writable"
        except Exception as e:
            return False, str(e)

    status = api_status()
    media = media_tool_status()
    network = network_status()
    voice_count = 0
    try:
        voice_count = len(windows_voice_names())
    except Exception:
        voice_count = 0

    add("Python packages", all(importlib.util.find_spec(m) for m in ["dotenv", "requests", "pyttsx3", "speech_recognition", "pyaudio", "pyautogui", "pyperclip", "pygetwindow", "docx", "openpyxl", "pptx", "PIL"]), "core packages installed")
    add("Workspace files", Path(__file__).exists() and SETTINGS_FILE.exists() and MEMORY_FILE.exists(), str(BASE_DIR))
    ok, detail = write_check(FILES_DIR)
    add("Nova files folder", ok, detail)
    ok, detail = write_check(MEDIA_DIR)
    add("Media Studio folder", ok, detail)
    add("Laptop profile", bool(settings.get("laptop_connected") or LAPTOP_PROFILE_FILE.exists()), str(LAPTOP_PROFILE_FILE))
    add("App index", APP_INDEX_FILE.exists(), str(APP_INDEX_FILE))
    add("Startup learning scan", bool(settings.get("auto_scan_on_start", True)), "last scan: " + str(settings.get("last_startup_scan_at", "not yet")))
    add("App playbooks", APP_PLAYBOOK_FILE.exists(), str(APP_PLAYBOOK_FILE))
    add("Control mode", bool(settings.get("control_mode", True) and settings.get("local_control_permission", False)), "mouse/keyboard permission profile saved")
    add("Assistant ON/OFF", "assistant_enabled" in settings, "switch state: " + ("ON" if settings.get("assistant_enabled", True) else "OFF"))
    add("Safety confirmations", str(handle_command("shutdown")).startswith("For safety I need"), "dangerous commands require confirm")
    add("ChatGPT tools", "ChatGPT-style tools are ON" in chatgpt_feature_help(), "new/export/fork/read/regenerate commands available")
    add("Command router", "I could not find definitelymissingapp" in handle_command("open definitelymissingapp"), "unknown apps do not auto-search Google")
    add("Media image tools", bool(media.get("pillow")), "Pillow ready" if media.get("pillow") else "Pillow missing")
    add("Media video tools", bool(media.get("ffmpeg")), "FFmpeg ready" if media.get("ffmpeg") else "FFmpeg missing; run INSTALL_MEDIA_TOOLS.cmd for MP4 crop/merge")
    add("Media video analysis", bool(media.get("ffprobe")), "FFprobe ready" if media.get("ffprobe") else "FFprobe missing; video metadata is limited")
    add("Voice output", bool(settings.get("voice_enabled", True) and voice_count), f"{voice_count} Windows voices detected")
    add("Voice input", bool(sr), "SpeechRecognition installed" if sr else SR_IMPORT_ERROR)
    add("Online brain config", bool(status["online"]), "providers: " + (", ".join(status["order"]) if status["order"] else "none"))
    add("Network checker", bool(network.get("online")), network.get("message", "unknown"))
    add("Python/API HTTP", bool(network.get("python_http", network.get("online"))), "API traffic reachable" if network.get("python_http") else "Python HTTP may be blocked even if browser internet works; run ALLOW_NOVA_NETWORK.cmd")
    add("Mobile mode support", Path(BASE_DIR / "START_MOBILE_MODE.cmd").exists(), "LAN IP helper: " + get_lan_ip())
    add("Media installer", Path(BASE_DIR / "INSTALL_MEDIA_TOOLS.cmd").exists(), "optional FFmpeg/Pillow installer present")
    add("Network permission helper", Path(BASE_DIR / "ALLOW_NOVA_NETWORK.cmd").exists(), "optional Windows Firewall helper present")

    passed = sum(1 for _, ok, _ in results if ok)
    total = len(results)
    lines = [
        "Nova Full Function Check",
        f"- Passed: {passed}/{total}",
    ]
    for name, ok, detail in results:
        mark = "OK" if ok else "NEEDS ATTENTION"
        lines.append(f"- {name}: {mark}. {detail}")

    if passed == total:
        lines.append("Result: all checked non-destructive systems look ready.")
    else:
        lines.append("Result: Nova can still run, but the NEEDS ATTENTION items above explain what is limited. I will not give fake success for those limited functions.")
    if repair:
        lines.append("Repair applied: runtime folders checked, assistant/control/voice settings normalized.")
    return "\n".join(lines)


def prepare_assistant_runtime():
    changed = False
    for key in ["control_mode", "local_control_permission", "local_only_mode"]:
        if not settings.get(key, False):
            settings[key] = True
            changed = True
    if "voice_enabled" not in settings:
        settings["voice_enabled"] = True
        changed = True
    if int(settings.get("speak_rate", 172)) > 145:
        settings["speak_rate"] = 145
        changed = True
    if settings.get("voice_language_mode") != "auto":
        settings["voice_language_mode"] = "auto"
        changed = True
    if "english_voice_name" not in settings:
        settings["english_voice_name"] = ""
        changed = True
    if "hindi_voice_name" not in settings:
        settings["hindi_voice_name"] = ""
        changed = True
    if "Hindi-English" in str(settings.get("assistant_style", "")):
        settings["assistant_style"] = DEFAULT_SETTINGS["assistant_style"]
        changed = True
    if changed:
        save_json(SETTINGS_FILE, settings)
    try:
        ok, msg = startup_system_scan(force=not LAPTOP_PROFILE_FILE.exists())
        log(msg)
    except Exception as e:
        log(f"startup learning scan failed: {e}")


def open_shortcut(path: Path) -> bool:
    try:
        os.startfile(str(path))
        return True
    except Exception as e:
        log(f"shortcut open error: {e}")
        try:
            subprocess.Popen(["cmd", "/c", "start", "", str(path)])
            return True
        except Exception as e2:
            log(f"shortcut fallback open error: {e2}")
        return False


def find_start_menu_app(name):
    apps = scan_start_menu_apps()
    name_l = name.lower().strip()
    # Exact/contains first; fuzzy only for longer names to avoid wrong apps like VLC for "files".
    for n, p in apps:
        if n == name_l:
            return p
    for n, p in apps:
        if name_l in n or n in name_l:
            return p
    if len(name_l) >= 5:
        names = [n for n, _ in apps]
        matches = difflib.get_close_matches(name_l, names, n=1, cutoff=0.72)
        if matches:
            for n, p in apps:
                if n == matches[0]:
                    return p
    return None


def find_windows_start_app(name):
    name_l = clean_entity_name(name).lower()
    if not name_l:
        return None
    candidates = []
    if APP_INDEX_FILE.exists():
        try:
            candidates.extend(json.loads(APP_INDEX_FILE.read_text(encoding="utf-8")).get("start_apps", []))
        except Exception:
            pass
    if LAPTOP_PROFILE_FILE.exists():
        try:
            candidates.extend(json.loads(LAPTOP_PROFILE_FILE.read_text(encoding="utf-8")).get("start_apps", []))
        except Exception:
            pass
    if not candidates:
        candidates = scan_windows_start_apps()
    seen = set()
    unique = []
    for item in candidates:
        app_id = str(item.get("app_id", "")).strip()
        label = str(item.get("name", "")).strip()
        key = (label.lower(), app_id.lower())
        if label and app_id and key not in seen:
            seen.add(key)
            unique.append({"name": label, "app_id": app_id})
    for item in unique:
        label_l = item["name"].lower()
        if label_l == name_l:
            return item
    for item in unique:
        label_l = item["name"].lower()
        if name_l in label_l or label_l in name_l:
            return item
    if len(name_l) >= 5:
        names = [item["name"].lower() for item in unique]
        matches = difflib.get_close_matches(name_l, names, n=1, cutoff=0.75)
        if matches:
            for item in unique:
                if item["name"].lower() == matches[0]:
                    return item
    return None


def open_windows_start_app(item):
    if not item:
        return False
    app_id = str(item.get("app_id", "")).strip()
    if not app_id:
        return False
    target = "shell:AppsFolder\\" + app_id
    try:
        subprocess.Popen(["explorer.exe", target])
        return True
    except Exception as e:
        log(f"Start app open error: {e}")
        return run_start(target)


def looks_like_local_app_request(c):
    if not settings.get("control_mode", True):
        return False
    if any(x in c for x in ["what", "why", "how", "who", "when", "where", "kaise", "kyu", "kya", "?"]):
        return False
    name = clean_entity_name(remove_words(c, ["app", "application", "tool", "tools"]))
    if not name or len(name.split()) > 4:
        return False
    return bool(name in APP_ALIASES or find_start_menu_app(name) or find_windows_start_app(name) or find_known_executable(name) or executable_from_profile(name))


def open_app(name):
    raw = normalize_words(name)
    name = clean_entity_name(remove_words(raw, ["open", "start", "launch", "app", "application", "tool", "tools"]))
    # common speech mistakes
    if name in ["", "it"]:
        return False, "Tell me which app to open."
    if name in ["google", "browser"]:
        shortcut = find_start_menu_app("google") or find_start_menu_app("chrome") or find_start_menu_app("brave") or find_start_menu_app("edge")
        if shortcut and open_shortcut(shortcut):
            return True, f"Opening {shortcut.stem}."
        return False, "I could not find a browser app locally. Say 'search Google for browser' or install a browser."

    # Laptop-first: if Start Menu has the requested app/PWA, open it before any web URL.
    shortcut = find_start_menu_app(name)
    if shortcut and open_shortcut(shortcut):
        return True, f"Opening {shortcut.stem}."
    start_app = find_windows_start_app(name)
    if start_app and open_windows_start_app(start_app):
        return True, f"Opening {start_app.get('name', name)}."
    exe = find_known_executable(name)
    if exe:
        try:
            subprocess.Popen([str(exe)])
            return True, f"Opening {name}."
        except Exception as e:
            log(f"known executable open error: {e}")
            if run_start(str(exe)):
                return True, f"Tried to open {name} with Windows shell fallback."
    profile_exe = executable_from_profile(name)
    if profile_exe:
        try:
            subprocess.Popen([str(profile_exe)])
            return True, f"Opening {profile_exe.stem} from learned laptop profile."
        except Exception as e:
            log(f"profile executable open error: {e}")
            if run_start(str(profile_exe)):
                return True, f"Tried to open {profile_exe.stem} from learned laptop profile with Windows shell fallback."

    # Alias/setting URI after local app check.
    targets = APP_ALIASES.get(name, [])
    for target in targets:
        if target.startswith("http"):
            app_shortcut = find_start_menu_app(name.replace(" app", ""))
            if app_shortcut and open_shortcut(app_shortcut):
                return True, f"Opening {app_shortcut.stem}."
            return False, f"I could not find {name} installed on this laptop. Say 'open {name} in browser' or 'install {name}' if you want."
        if target.lower() in ["calc"]:
            if run_start("calculator:") or run_start("calc.exe"):
                return True, "Opening calculator."
        if target.lower() == "myasus":
            shortcut = find_start_menu_app("myasus") or find_start_menu_app("my asus") or find_start_menu_app("asus")
            if shortcut and open_shortcut(shortcut):
                return True, f"Opening {shortcut.stem}."
            return False, "MyASUS was not found locally. Say 'install MyASUS' if you want me to open install options."
        if run_start(target):
            return True, f"Tried to open {name}. If it did not appear, the app may be missing or blocked by Windows."

    # Do NOT automatically Google for open-app commands. It caused wrong behavior.
    # Search only when user explicitly says search/google.
    return False, f"I could not find {name} as an installed app. Say 'learn laptop' once, or say 'search Google for {name}' if you want web search."

def close_app(name):
    name = remove_words(normalize_words(name), ["close", "stop", "exit", "quit", "band", "app", "application", "karo", "to"]).strip()
    if not name:
        return window_control("close window")
    process_map = {
        "google": "chrome.exe", "chrome": "chrome.exe", "edge": "msedge.exe", "spotify": "Spotify.exe", "whatsapp": "WhatsApp.exe",
        "word": "WINWORD.EXE", "ms word": "WINWORD.EXE", "excel": "EXCEL.EXE", "powerpoint": "POWERPNT.EXE",
        "notepad": "notepad.exe", "calculator": "CalculatorApp.exe", "paint": "mspaint.exe", "vs code": "Code.exe",
        "visual studio code": "Code.exe", "files": "explorer.exe", "file": "explorer.exe", "explorer": "explorer.exe"
    }
    exe = process_map.get(name, name if name.lower().endswith(".exe") else name + ".exe")
    try:
        result = subprocess.run(["taskkill", "/IM", exe, "/F"], capture_output=True, text=True, timeout=8)
        if result.returncode == 0:
            return True, f"Closed {name}."
        return False, f"I could not close {name}. It may not be running."
    except Exception as e:
        return False, f"Could not close {name}: {e}"


def focus_app(c):
    name = clean_entity_name(remove_words(normalize_words(c), ["focus", "bring", "front", "show", "open", "app", "application", "this", "it"]))
    if not name:
        return False, "Tell me which app to bring to the front."
    if focus_window(name):
        return True, f"Brought {name} to the front."
    return open_app("open " + name)


def search_web(query, engine="google"):
    c = normalize_words(query)
    query = remove_words(c, ["search", "find", "google", "youtube", "on", "for"])
    if not query:
        return False, "Tell me what to search."
    if engine == "youtube":
        webbrowser.open("https://www.youtube.com/results?search_query=" + quote_plus(query))
        return True, f"Searching YouTube for {query}."
    webbrowser.open("https://www.google.com/search?q=" + quote_plus(query))
    return True, f"Searching Google for {query}."


def tor_browser_status_text():
    exe = find_known_executable("tor") or find_start_menu_app("tor browser") or find_start_menu_app("tor")
    return (
        "Tor Browser status\n"
        f"- Permission: {'ON' if settings.get('tor_browser_permission', False) else 'OFF'}\n"
        f"- Installed/found: {'yes' if exe else 'no'}\n"
        f"- Path: {exe if exe else 'not found locally'}\n"
        "- Use: open tor, search privacy tips in tor, or open tor browser.\n"
        "- Safety: Nova will use Tor only for legal privacy browsing and defensive research."
    )


def tor_request_is_unsafe(c):
    blocked = [
        "buy drugs", "sell drugs", "weapon", "gun", "explosive", "stolen", "carding", "credit card",
        "credential", "password dump", "leaked password", "otp", "bank login", "hack account",
        "malware", "ransomware", "botnet", "phishing", "exploit kit", "illegal market", "marketplace",
        "hire hacker", "dox", "doxx", "hitman", "child abuse", "csam",
    ]
    return any(x in c for x in blocked)


def open_tor_browser(url=""):
    if not settings.get("tor_browser_permission", False):
        return True, "Tor Browser permission is OFF in Nova Settings. Turn it ON and allow the popup, or say: allow tor browser permission."
    exe = find_known_executable("tor") or find_known_executable("tor browser")
    if exe:
        try:
            args = [str(exe)]
            if url:
                args.append(url)
            subprocess.Popen(args)
            return True, "Opening Tor Browser." if not url else "Opening the requested legal/private search in Tor Browser."
        except Exception as e:
            return False, f"Tor Browser was found but could not open: {e}"
    shortcut = find_start_menu_app("tor browser") or find_start_menu_app("tor")
    if shortcut and open_shortcut(shortcut):
        return True, "Opening Tor Browser."
    return False, "Tor Browser was not found on this laptop. Install Tor Browser first, then say: learn laptop."


def tor_browser_command(raw):
    c = normalize_words(raw)
    if c in ["tor status", "tor browser status", "privacy browser status"]:
        return True, tor_browser_status_text()
    if tor_request_is_unsafe(c):
        return True, "I cannot use Tor for illegal markets, stolen data, hacking, malware, credential access, or harm. I can help with legal privacy, security education, and defensive leak-risk checks."
    if ".onion" in c:
        return True, "I will not open unknown .onion links automatically. I can open Tor Browser for you, but hidden-service links must be verified and opened manually by you."
    if any(x in c for x in ["search", "find", "look up", "research"]):
        query = remove_words(c, ["search", "find", "look", "up", "research", "in", "on", "with", "using", "tor", "browser", "private"])
        query = query.strip()
        if not query:
            return open_tor_browser()
        return open_tor_browser("https://duckduckgo.com/?q=" + quote_plus(query))
    if any(x in c for x in ["open", "start", "launch"]) or c in ["tor", "tor browser"]:
        return open_tor_browser()
    return True, tor_browser_status_text()


def google_command(command):
    if not settings.get("google_browser_permission", False):
        return True, "Google Browser permission is OFF in Nova Settings. Turn it ON and allow the popup before I open Google services."
    c = normalize_words(command)
    if "gmail" in c or re.search(r"\b(email|mail)\b", c):
        return gmail_command(command)
    service_urls = {
        "google account": "https://myaccount.google.com",
        "my google account": "https://myaccount.google.com",
        "google drive": "https://drive.google.com",
        "drive": "https://drive.google.com",
        "google docs": "https://docs.google.com",
        "docs": "https://docs.google.com",
        "google sheets": "https://sheets.google.com",
        "sheets": "https://sheets.google.com",
        "google calendar": "https://calendar.google.com",
        "calendar": "https://calendar.google.com",
        "google photos": "https://photos.google.com",
        "photos": "https://photos.google.com",
        "google": "https://www.google.com",
    }
    for name, url in service_urls.items():
        if name in c:
            webbrowser.open(url)
            return True, f"Opening {name} in your browser session."
    if c.startswith("google "):
        query = remove_words(c, ["google", "search", "for"])
        if query:
            return search_web("search google " + query, "google")
    return False, ""


def gmail_command(command):
    if not settings.get("gmail_browser_permission", False):
        return True, "Gmail Browser permission is OFF in Nova Settings. Turn it ON and allow the popup before I open Gmail."
    raw = str(command or "").strip()
    c = normalize_words(raw)
    if any(x in c for x in ["delete", "archive", "mark read", "mark unread", "send now"]):
        webbrowser.open("https://mail.google.com")
        return True, "Gmail opened. For safety, I will not delete/archive/send inside Gmail automatically; review and click the final action yourself."
    if any(x in c for x in ["compose", "write email", "new email", "send email", "mail to", "email to"]):
        to_match = re.search(r"[\w.+-]+@[\w.-]+\.[a-zA-Z]{2,}", raw)
        to = to_match.group(0) if to_match else ""
        subject = ""
        body = ""
        subject_match = re.search(r"\bsubject\b\s+(.+?)(?:\s+\bbody\b|\s+\bmessage\b|$)", raw, flags=re.I)
        if subject_match:
            subject = subject_match.group(1).strip()
        body_match = re.search(r"\b(?:body|message)\b\s+(.+)$", raw, flags=re.I)
        if body_match:
            body = body_match.group(1).strip()
        url = "https://mail.google.com/mail/?view=cm&fs=1"
        if to:
            url += "&to=" + quote_plus(to)
        if subject:
            url += "&su=" + quote_plus(subject)
        if body:
            url += "&body=" + quote_plus(body)
        webbrowser.open(url)
        return True, "Gmail compose opened in your logged-in browser. Review it, then press Send manually."
    if "search" in c or "find" in c:
        query = remove_words(c, ["search", "find", "gmail", "email", "mail", "in", "on", "for"])
        if query:
            webbrowser.open("https://mail.google.com/mail/u/0/#search/" + quote_plus(query))
            return True, f"Searching Gmail for {query} in your browser session."
    webbrowser.open("https://mail.google.com")
    return True, "Opening Gmail in your browser session."


def focus_window(title_word):
    if not gw:
        return False
    try:
        for w in gw.getAllWindows():
            if title_word.lower() in (w.title or "").lower():
                if w.isMinimized:
                    w.restore()
                w.activate()
                return True
    except Exception as e:
            log(f"focus error: {e}")
    return False


def click_whatsapp_call_button(video=False):
    if not pyautogui or not gw:
        return False
    try:
        for w in gw.getAllWindows():
            if "whatsapp" in (w.title or "").lower():
                if w.isMinimized:
                    w.restore()
                w.activate()
                time.sleep(0.35)
                # WhatsApp Desktop keeps call icons near the chat header's right side.
                # This coordinate fallback is used after selecting a chat because
                # WhatsApp does not expose a reliable public "start call" command.
                x_offset = 96 if video else 146
                x = max(w.left + 30, w.left + w.width - x_offset)
                y = w.top + 58
                pyautogui.click(x, y)
                return True
    except Exception as e:
        log(f"WhatsApp call click error: {e}")
    return False


def send_media_key(action="playpause"):
    """Send Windows global media keys. Returns true if the key was sent.
    Windows/Spotify do not expose a reliable no-login local playback API here,
    so replies must stay honest and say the key was pressed, not "song is playing".
    """
    key_map = {
        "play": "playpause",
        "pause": "playpause",
        "playpause": "playpause",
        "next": "nexttrack",
        "previous": "prevtrack",
        "prev": "prevtrack",
        "stop": "stop",
    }
    key = key_map.get(action, "playpause")
    if pyautogui:
        try:
            pyautogui.press(key)
            return True
        except Exception as e:
            log(f"media key pyautogui failed: {e}")
    if platform.system().lower() == "windows":
        vk_map = {"playpause": 0xB3, "nexttrack": 0xB0, "prevtrack": 0xB1, "stop": 0xB2}
        try:
            vk = vk_map.get(key, 0xB3)
            ctypes.windll.user32.keybd_event(vk, 0, 0, 0)
            ctypes.windll.user32.keybd_event(vk, 0, 2, 0)
            return True
        except Exception as e:
            log(f"media key ctypes failed: {e}")
    return False


SPOTIFY_DEFAULT_PLAYLIST_URI = "spotify:playlist:37i9dQZF1DXcBWIGoYBM5M"


def focus_spotify_app():
    if focus_window("spotify"):
        time.sleep(0.25)
        return True
    run_start("spotify:")
    time.sleep(2.0)
    return focus_window("spotify")


def spotify_window_rect():
    if not gw:
        return None
    try:
        for w in gw.getAllWindows():
            if "spotify" in (w.title or "").lower():
                if w.isMinimized:
                    w.restore()
                w.activate()
                return w
    except Exception as e:
        log(f"Spotify window read failed: {e}")
    return None


def spotify_click_play_candidates(song_result=False):
    """Click one likely Spotify play position.
    Multiple clicks can toggle playback back to pause, so this intentionally
    stops after the first successful click.
    """
    if not pyautogui:
        return False
    clicked = False
    w = spotify_window_rect()
    if not w:
        return False
    try:
        candidates = []
        if song_result:
            candidates.extend([
                (w.left + int(w.width * 0.26), w.top + int(w.height * 0.35)),
                (w.left + int(w.width * 0.32), w.top + int(w.height * 0.42)),
                (w.left + int(w.width * 0.45), w.top + int(w.height * 0.40)),
            ])
        candidates.extend([
            (w.left + int(w.width * 0.50), w.top + w.height - 70),
            (w.left + int(w.width * 0.08), w.top + w.height - 72),
            (w.left + int(w.width * 0.16), w.top + w.height - 72),
        ])
        for x, y in candidates:
            pyautogui.click(x, y)
            clicked = True
            time.sleep(0.55)
            return True
        return clicked
    except Exception as e:
        log(f"Spotify click play failed: {e}")
        return clicked


def delayed_spotify_play_attempt(query="", generic=False):
    """Best-effort UI attempt after opening Spotify/search. It deliberately does not claim success."""
    try:
        time.sleep(2.4 if query or generic else 1.4)
        focus_spotify_app()
        if pyautogui:
            if query:
                # Do not depend only on spotify:search URI. Ctrl+L focuses Spotify search
                # in most desktop builds, then we type the query and open the first result.
                pyautogui.hotkey("ctrl", "l")
                time.sleep(0.25)
                paste_text(query)
                time.sleep(1.0)
                pyautogui.press("enter")
                time.sleep(0.9)
                spotify_click_play_candidates(song_result=True)
            elif generic:
                # A generic "play song" should try to start any available music, not ask for a name.
                # The playlist URI gives Spotify a playable context when nothing is selected.
                # Click the playlist/result area, not the bottom play/pause toggle.
                spotify_click_play_candidates(song_result=True)
        elif generic:
            # No UI automation available. This can toggle play/pause, so use it only
            # when there is no safer click path.
            send_media_key("play")
    except Exception as e:
        log(f"Spotify play attempt failed: {e}")


def youtube_first_video_url(query):
    if not requests:
        return ""
    try:
        url = "https://www.youtube.com/results?search_query=" + quote_plus(query)
        headers = {"User-Agent": "Mozilla/5.0 NovaAssistant/1.0"}
        r = requests.get(url, headers=headers, timeout=8)
        ids = re.findall(r'"videoId":"([A-Za-z0-9_-]{11})"', r.text)
        seen = []
        for vid in ids:
            if vid not in seen:
                seen.append(vid)
        if seen:
            return "https://www.youtube.com/watch?v=" + seen[0]
    except Exception as e:
        log(f"YouTube first result lookup failed: {e}")
    return ""


def media_control_command(c):
    c = normalize_words(c)
    if "spotify" in c and not settings.get("spotify_control_enabled", True):
        return True, "Spotify Control is OFF in Nova Settings."
    if any(x in c for x in ["next song", "next track", "next music", "spotify next", "youtube next"]):
        ok = send_media_key("next")
        return True, "Pressed Windows media Next key." if ok else "I could not send the media Next key."
    if any(x in c for x in ["previous song", "previous track", "prev song", "back song", "spotify previous"]):
        ok = send_media_key("previous")
        return True, "Pressed Windows media Previous key." if ok else "I could not send the media Previous key."
    if any(x in c for x in ["pause song", "pause music", "pause spotify", "pause youtube"]):
        ok = send_media_key("pause")
        return True, "Pressed Windows media Play/Pause key. If music was playing, it should pause." if ok else "I could not send the media pause key."
    if any(x in c for x in ["resume song", "resume music", "continue song", "continue music"]):
        ok = send_media_key("play")
        return True, "Pressed Windows media Play/Pause key. If music was paused, it should resume." if ok else "I could not send the media play key."
    return False, ""


def play_media(c):
    c = normalize_words(c)
    platform_name = "spotify" if "spotify" in c else "youtube" if "youtube" in c else "spotify"
    query = remove_words(c, ["play", "song", "songs", "music", "on", "in", "youtube", "spotify", "google", "app", "that", "please", "any", "koi", "kuch", "available", "current"]).strip()
    if platform_name == "spotify":
        if not settings.get("spotify_control_enabled", True):
            return True, "Spotify Control is OFF in Nova Settings. Turn it ON and allow the popup before I control Spotify."
        if query:
            run_start("spotify:search:" + quote_plus(query))
            threading.Thread(target=delayed_spotify_play_attempt, args=(query,), daemon=True).start()
            return True, f"I am searching Spotify for '{query}' and trying to press Play on the first result only once. If Spotify login, an ad, or the app layout blocks automation, I will not claim fake success."
        run_start(SPOTIFY_DEFAULT_PLAYLIST_URI)
        threading.Thread(target=delayed_spotify_play_attempt, args=("", True), daemon=True).start()
        return True, "Generic Spotify play understood: I am opening an available playlist and trying to press Play only once. I will not send the extra Play/Pause key, so the song should not stop immediately after starting."
    # YouTube only when user says YouTube.
    if not query:
        query = "music"
    first = youtube_first_video_url(query)
    if first:
        webbrowser.open(first)
        return True, f"I opened the first YouTube video result for '{query}'. If browser autoplay is blocked, press the Play button."
    webbrowser.open("https://www.youtube.com/results?search_query=" + quote_plus(query))
    return True, f"I opened YouTube search results for '{query}'. I could not verify that a video started, so I am not saying it is playing."

def paste_text(text):
    if pyperclip:
        pyperclip.copy(text)
        if pyautogui:
            pyautogui.hotkey("ctrl", "v")
            return True
    if pyautogui:
        pyautogui.write(text, interval=0.01)
        return True
    return False


def type_text(c):
    text = remove_words(c, ["type", "write", "paste", "now", "please"])
    if not text:
        return False, "Tell me what to type."
    ok = paste_text(text)
    return ok, "Typed it." if ok else "Typing module is not installed."


def make_file(kind, title="Nova File", content=None):
    kind = kind.lower()
    content = content or f"Created by {APP_NAME}\nDate: {datetime.datetime.now()}\n"
    if kind in ["text", "txt", "note"]:
        p = safe_filename(title, "txt")
        p.write_text(content, encoding="utf-8")
    elif kind in ["word", "doc", "docx"]:
        p = safe_filename(title, "docx")
        try:
            from docx import Document
            doc = Document()
            doc.add_heading(title, 0)
            doc.add_paragraph(content)
            doc.save(p)
        except Exception:
            p = safe_filename(title, "rtf")
            p.write_text(r"{\rtf1\ansi " + content.replace("\n", r"\line ") + "}", encoding="utf-8")
    elif kind in ["excel", "xlsx", "sheet"]:
        p = safe_filename(title, "xlsx")
        try:
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = "Nova"
            ws.append(["Created By", APP_NAME])
            ws.append(["Date", str(datetime.datetime.now())])
            wb.save(p)
        except Exception:
            p = safe_filename(title, "csv")
            p.write_text("Created By,Nova AI\n", encoding="utf-8")
    elif kind in ["powerpoint", "ppt", "pptx", "presentation"]:
        p = safe_filename(title, "pptx")
        try:
            from pptx import Presentation
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
            slide.placeholders[1].text = f"Created by {APP_NAME}"
            prs.save(p)
        except Exception:
            p = safe_filename(title, "txt")
            p.write_text("Install python-pptx to create PPTX files.", encoding="utf-8")
    else:
        return False, "Supported files: text, Word, Excel, PowerPoint."
    try:
        os.startfile(p)
    except Exception:
        pass
    return True, f"Created {p.name} on Desktop/Nova AI Files."



def smart_content_for(c):
    c = normalize_words(c)
    if "date" in c and "time" in c and "code" in c:
        return 'Python date and time code:\n\nfrom datetime import datetime\n\nnow = datetime.now()\nprint("Current date:", now.strftime("%d-%m-%Y"))\nprint("Current time:", now.strftime("%H:%M:%S"))\nprint("Date and time:", now.strftime("%d-%m-%Y %H:%M:%S"))\n'
    if "python" in c and "code" in c:
        return "Python code example:\n\nprint('Hello from Nova AI')\n"
    if "code" in c:
        return "Code note created by Nova AI. Tell me the language and topic for exact code.\n"
    return "Created by Nova AI.\nDate and time: " + str(datetime.datetime.now())


def create_code_file(c):
    content = smart_content_for(c)
    title = "date_time_code" if "date" in c and "time" in c else "nova_code"
    p = FILES_DIR / (title + "_" + datetime.datetime.now().strftime("%Y%m%d_%H%M%S") + ".py")
    p.write_text(content.replace("Python date and time code:\n\n", ""), encoding="utf-8")
    try:
        os.startfile(p)
    except Exception:
        pass
    return True, f"Created Python code file {p.name} on Desktop/Nova AI Files."


def create_document_from_command(c):
    c = normalize_words(c)
    if "word" in c:
        content = ""
        if "write" in c:
            content = c.split("write", 1)[1].strip()
        elif "type" in c:
            content = c.split("type", 1)[1].strip()
        if not content or "code" in c or "date" in c or "time" in c:
            content = smart_content_for(c)
        return make_file("word", "Nova Word Document", content)
    if "excel" in c or "sheet" in c:
        return make_file("excel", "Nova Excel Sheet")
    if "powerpoint" in c or "ppt" in c or "presentation" in c or "slide" in c:
        return make_file("powerpoint", "Nova Presentation")
    return make_file("text", "Nova Note")


# -------------------- File Studio --------------------
def file_studio_help():
    return (
        "File Studio is ready.\n\n"
        "Analyze commands:\n"
        "- analyze latest file\n"
        "- analyze uploaded file\n"
        "- summarize report.docx\n"
        "- read data.csv\n\n"
        "Create commands:\n"
        "- create python file date time code\n"
        "- make html file for portfolio\n"
        "- create json file with name Adarsh\n"
        "- make Word file with date and time code\n\n"
        "Supported local analysis: text/code, JSON, CSV, DOCX, XLSX, PPTX, PDF basic info, images, and videos."
    )


def safe_general_filename(name, default_ext=".txt", folder=None):
    raw = Path(str(name or "nova_file")).name
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", Path(raw).stem).strip(" ._") or "nova_file"
    ext = Path(raw).suffix.lower() or default_ext
    if not ext.startswith("."):
        ext = "." + ext
    if ext not in GENERAL_FILE_EXTS:
        ext = default_ext
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return (folder or FILE_CREATED_DIR) / f"{stem}_{stamp}{ext}"


def file_report_path(stem):
    return safe_general_filename(f"{stem}_file_analysis.txt", ".txt", FILE_REPORT_DIR)


def file_roots():
    roots = [
        FILE_UPLOAD_DIR,
        FILE_CREATED_DIR,
        FILES_DIR,
        MEDIA_UPLOAD_DIR,
        MEDIA_IMAGE_DIR,
        MEDIA_VIDEO_DIR,
        MEDIA_EDIT_DIR,
        SCREENSHOT_DIR,
        DESKTOP,
        Path.home() / "Documents",
        Path.home() / "Downloads",
        Path.home() / "Pictures",
        Path.home() / "Videos",
    ]
    return [Path(p) for p in roots if p and Path(p).exists()]


def iter_general_files(kind="any"):
    exts = GENERAL_FILE_EXTS
    if kind == "text":
        exts = TEXT_FILE_EXTS
    elif kind == "document":
        exts = DOC_FILE_EXTS
    elif kind == "media":
        exts = IMAGE_EXTS | VIDEO_EXTS
    seen = set()
    deep_roots = {FILE_UPLOAD_DIR, FILE_CREATED_DIR, FILE_REPORT_DIR, FILES_DIR, MEDIA_DIR}
    for root in file_roots():
        iterator = root.rglob("*") if root in deep_roots else root.glob("*")
        try:
            for path in iterator:
                if not path.is_file() or path.suffix.lower() not in exts:
                    continue
                key = str(path.resolve()).lower()
                if key in seen:
                    continue
                seen.add(key)
                yield path
        except Exception:
            continue


def find_general_file(command, kind="any"):
    raw = str(command or "")
    ext_group = "|".join(re.escape(x.lstrip(".")) for x in sorted(GENERAL_FILE_EXTS, key=len, reverse=True))
    quoted = re.findall(rf'"([^"]+\.({ext_group}))"', raw, flags=re.I)
    candidates = [item[0] if isinstance(item, tuple) else item for item in quoted]
    candidates += re.findall(rf"([A-Za-z]:\\[^<>|?*\n\r]+?\.({ext_group}))", raw, flags=re.I)
    flat = []
    for item in candidates:
        flat.append(item[0] if isinstance(item, tuple) else item)
    for item in flat:
        p = Path(str(item).strip())
        if p.exists() and p.is_file():
            return p
    files = list(iter_general_files(kind))
    if not files:
        return None
    c = normalize_words(raw)
    if any(x in c for x in ["latest", "recent", "last", "uploaded"]):
        return max(files, key=lambda p: p.stat().st_mtime)
    query = c
    for word in [
        "analyze", "analyse", "summarize", "summarise", "summary", "read", "open", "file",
        "document", "uploaded", "latest", "recent", "please", "karo", "ka", "ki", "ko", "naam", "named"
    ]:
        query = re.sub(rf"\b{re.escape(word)}\b", " ", query)
    query = re.sub(r"\s+", " ", query).strip()
    if not query:
        return max(files, key=lambda p: p.stat().st_mtime)
    names = {normalize_words(p.stem): p for p in files}
    for name, path in names.items():
        if query in name or name in query:
            return path
    match = difflib.get_close_matches(query, list(names.keys()), n=1, cutoff=0.52)
    if match:
        return names[match[0]]
    return max(files, key=lambda p: p.stat().st_mtime)


def file_size_text(path):
    size = path.stat().st_size
    for unit in ["B", "KB", "MB", "GB"]:
        if size < 1024 or unit == "GB":
            return f"{size:.1f} {unit}" if unit != "B" else f"{size} B"
        size /= 1024
    return str(path.stat().st_size) + " B"


def read_text_preview(path, limit=12000):
    data = path.read_bytes()[:limit]
    for enc in ["utf-8-sig", "utf-8", "cp1252", "latin-1"]:
        try:
            return data.decode(enc), enc
        except Exception:
            continue
    return data.decode("utf-8", errors="replace"), "utf-8-replace"


def analyze_text_file(path):
    text, enc = read_text_preview(path)
    lines = text.splitlines()
    words = re.findall(r"\S+", text)
    info = [
        f"File analyzed: {path.name}",
        f"- Type: text/code ({path.suffix.lower() or 'no extension'})",
        f"- Size: {file_size_text(path)}",
        f"- Encoding read as: {enc}",
        f"- Lines in preview: {len(lines)}",
        f"- Words in preview: {len(words)}",
        f"- Characters in preview: {len(text)}",
    ]
    ext = path.suffix.lower()
    if ext == ".json":
        try:
            parsed = json.loads(path.read_text(encoding=enc if enc != "utf-8-replace" else "utf-8"))
            if isinstance(parsed, dict):
                info.append("- JSON keys: " + ", ".join(list(parsed.keys())[:20]))
            elif isinstance(parsed, list):
                info.append(f"- JSON list items: {len(parsed)}")
            info.append("- JSON syntax: OK")
        except Exception as e:
            info.append(f"- JSON syntax issue: {e}")
    elif ext == ".csv":
        try:
            import csv
            rows = list(csv.reader(text.splitlines()))
            cols = max((len(r) for r in rows), default=0)
            info.append(f"- CSV preview rows: {len(rows)}")
            info.append(f"- CSV preview columns: {cols}")
            if rows:
                info.append("- CSV header/first row: " + ", ".join(rows[0][:12]))
        except Exception as e:
            info.append(f"- CSV read issue: {e}")
    if text.strip():
        preview = text.strip()
        if len(preview) > 1600:
            preview = preview[:1600].rstrip() + "\n..."
        info.extend(["", "Preview:", preview])
    else:
        info.append("- Preview is empty.")
    return info


def analyze_docx_file(path):
    try:
        from docx import Document
        doc = Document(str(path))
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        words = re.findall(r"\S+", "\n".join(paras))
        info = [
            f"File analyzed: {path.name}",
            "- Type: Word document",
            f"- Size: {file_size_text(path)}",
            f"- Paragraphs with text: {len(paras)}",
            f"- Tables: {len(doc.tables)}",
            f"- Words: {len(words)}",
        ]
        if paras:
            info.extend(["", "Preview:", "\n".join(paras[:12])[:1800]])
        return info
    except Exception as e:
        return [f"File analyzed: {path.name}", "- Type: Word document", f"- Size: {file_size_text(path)}", f"- DOCX text analysis failed: {e}"]


def analyze_xlsx_file(path):
    try:
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        info = [
            f"File analyzed: {path.name}",
            "- Type: Excel workbook",
            f"- Size: {file_size_text(path)}",
            "- Sheets: " + ", ".join(wb.sheetnames),
        ]
        for ws in wb.worksheets[:8]:
            info.append(f"- {ws.title}: {ws.max_row} rows x {ws.max_column} columns")
        return info
    except Exception as e:
        return [f"File analyzed: {path.name}", "- Type: Excel workbook", f"- Size: {file_size_text(path)}", f"- XLSX analysis failed: {e}"]


def analyze_pptx_file(path):
    try:
        from pptx import Presentation
        prs = Presentation(str(path))
        texts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
        info = [
            f"File analyzed: {path.name}",
            "- Type: PowerPoint presentation",
            f"- Size: {file_size_text(path)}",
            f"- Slides: {len(prs.slides)}",
            f"- Text blocks: {len(texts)}",
        ]
        if texts:
            info.extend(["", "Preview:", "\n".join(texts[:12])[:1800]])
        return info
    except Exception as e:
        return [f"File analyzed: {path.name}", "- Type: PowerPoint presentation", f"- Size: {file_size_text(path)}", f"- PPTX analysis failed: {e}"]


def analyze_pdf_file(path):
    info = [f"File analyzed: {path.name}", "- Type: PDF", f"- Size: {file_size_text(path)}"]
    try:
        if importlib.util.find_spec("pypdf"):
            from pypdf import PdfReader
        elif importlib.util.find_spec("PyPDF2"):
            from PyPDF2 import PdfReader
        else:
            info.append("- PDF text extraction tool is not installed. I can only report basic file info.")
            return info
        reader = PdfReader(str(path))
        info.append(f"- Pages: {len(reader.pages)}")
        text_parts = []
        for page in reader.pages[:5]:
            try:
                text_parts.append(page.extract_text() or "")
            except Exception:
                pass
        preview = "\n".join(x.strip() for x in text_parts if x.strip())
        if preview:
            info.extend(["", "Preview:", preview[:1800]])
        else:
            info.append("- No extractable text found in the first pages.")
        return info
    except Exception as e:
        info.append(f"- PDF analysis failed: {e}")
        return info


def analyze_general_file(path):
    path = Path(path)
    if not path.exists():
        return False, f"I could not find this file: {path}"
    ext = path.suffix.lower()
    if ext in IMAGE_EXTS:
        return analyze_image_file(path)
    if ext in VIDEO_EXTS:
        return analyze_video_file(path)
    try:
        if ext in TEXT_FILE_EXTS:
            info = analyze_text_file(path)
        elif ext == ".docx":
            info = analyze_docx_file(path)
        elif ext == ".xlsx":
            info = analyze_xlsx_file(path)
        elif ext == ".pptx":
            info = analyze_pptx_file(path)
        elif ext == ".pdf":
            info = analyze_pdf_file(path)
        elif ext == ".rtf":
            info = analyze_text_file(path)
        else:
            info = [f"File analyzed: {path.name}", f"- Type: {ext or 'unknown'}", f"- Size: {file_size_text(path)}", "- I do not have a parser for this file type yet."]
        modified = datetime.datetime.fromtimestamp(path.stat().st_mtime).strftime("%d-%m-%Y %H:%M:%S")
        info.insert(2, f"- Modified: {modified}")
        report = file_report_path(path.stem)
        report.write_text("\n".join(info), encoding="utf-8")
        info.append(f"- Report saved: {report.name}")
        return True, "\n".join(info)
    except Exception as e:
        return False, f"Could not analyze file: {e}"


def infer_file_extension(c):
    if any(x in c for x in ["python", ".py", "py file"]):
        return ".py"
    if any(x in c for x in ["html", "website page", "web page"]):
        return ".html"
    if "css" in c:
        return ".css"
    if any(x in c for x in ["javascript", "java script", ".js"]):
        return ".js"
    if "typescript" in c or ".ts" in c:
        return ".ts"
    if "json" in c:
        return ".json"
    if "csv" in c:
        return ".csv"
    if any(x in c for x in ["markdown", ".md"]):
        return ".md"
    if any(x in c for x in ["word", "docx", "document"]):
        return ".docx"
    if any(x in c for x in ["excel", "xlsx", "sheet"]):
        return ".xlsx"
    if any(x in c for x in ["powerpoint", "ppt", "presentation", "slide"]):
        return ".pptx"
    return ".txt"


def infer_file_title(raw, ext):
    quoted = re.findall(r'"([^"]+)"', str(raw or ""))
    if quoted:
        return quoted[0]
    c = normalize_words(raw)
    m = re.search(r"(?:called|named|naam|name)\s+([a-z0-9 _.-]{2,60})", c)
    if m:
        title = m.group(1)
        title = re.split(r"\b(with|me|mein|mai|write|containing|for|ka|ki)\b", title)[0].strip()
        if title:
            return title
    if "date" in c and "time" in c and ("code" in c or ext == ".py"):
        return "date_time_code"
    if ext == ".html":
        return "nova_page"
    if ext == ".json":
        return "nova_data"
    if ext == ".csv":
        return "nova_table"
    if ext == ".md":
        return "nova_notes"
    return "nova_file"


def extract_creation_content(raw, c, ext):
    text = str(raw or "")
    for key in ["with content", "containing", "write", "type", "with", "for"]:
        idx = text.lower().find(key)
        if idx >= 0:
            content = text[idx + len(key):].strip(" :,-")
            if content:
                if ext == ".json":
                    try:
                        json.loads(content)
                        return json.dumps(json.loads(content), indent=2)
                    except Exception:
                        words = content.split()
                        if len(words) >= 2:
                            key_name = re.sub(r"[^A-Za-z0-9_]+", "_", words[0]).strip("_") or "value"
                            return json.dumps({key_name: " ".join(words[1:])}, indent=2)
                        return json.dumps({"value": content}, indent=2)
                return content
    if "date" in c and "time" in c and ("code" in c or ext == ".py"):
        return smart_content_for("python date time code").replace("Python date and time code:\n\n", "")
    if ext == ".py":
        return 'print("Hello from Nova AI")\n'
    if ext == ".html":
        return "<!doctype html>\n<html>\n<head>\n  <meta charset=\"utf-8\">\n  <title>Nova Page</title>\n</head>\n<body>\n  <h1>Hello from Nova AI</h1>\n</body>\n</html>\n"
    if ext == ".css":
        return "body {\n  font-family: Arial, sans-serif;\n  margin: 24px;\n}\n"
    if ext == ".js":
        return "console.log('Hello from Nova AI');\n"
    if ext == ".ts":
        return "const message: string = 'Hello from Nova AI';\nconsole.log(message);\n"
    if ext == ".json":
        return json.dumps({"created_by": APP_NAME, "created_at": datetime.datetime.now().isoformat()}, indent=2)
    if ext == ".csv":
        return "Name,Value\nCreated By,Nova AI\nDate," + datetime.datetime.now().strftime("%d-%m-%Y %H:%M:%S") + "\n"
    if ext == ".md":
        return "# Nova Notes\n\nCreated by Nova AI.\n"
    return smart_content_for(c)


def create_general_file_command(raw, c):
    ext = infer_file_extension(c)
    title = infer_file_title(raw, ext)
    content = extract_creation_content(raw, c, ext)
    out = safe_general_filename(title + ext, ext, FILE_CREATED_DIR)
    try:
        if ext == ".docx":
            try:
                from docx import Document
                doc = Document()
                doc.add_heading(title.replace("_", " ").title(), 0)
                doc.add_paragraph(content)
                doc.save(out)
            except Exception:
                out = safe_general_filename(title + ".rtf", ".rtf", FILE_CREATED_DIR)
                out.write_text(r"{\rtf1\ansi " + content.replace("\n", r"\line ") + "}", encoding="utf-8")
        elif ext == ".xlsx":
            try:
                from openpyxl import Workbook
                wb = Workbook()
                ws = wb.active
                ws.title = "Nova"
                for line in str(content).splitlines() or ["Created by Nova AI"]:
                    ws.append([line])
                wb.save(out)
            except Exception:
                out = safe_general_filename(title + ".csv", ".csv", FILE_CREATED_DIR)
                out.write_text("Created By,Nova AI\n" + str(content), encoding="utf-8")
        elif ext == ".pptx":
            try:
                from pptx import Presentation
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = title.replace("_", " ").title()
                slide.placeholders[1].text = str(content)[:800]
                prs.save(out)
            except Exception:
                out = safe_general_filename(title + ".txt", ".txt", FILE_CREATED_DIR)
                out.write_text("Install python-pptx to create PPTX files.\n\n" + str(content), encoding="utf-8")
        else:
            out.write_text(str(content).rstrip() + "\n", encoding="utf-8")
        try:
            os.startfile(out)
        except Exception:
            pass
        return True, f"Created file: {out.name}\nLocation: {out.parent}"
    except Exception as e:
        return False, f"Could not create file: {e}"


def looks_like_file_request(c):
    file_words = ["file", "document", "docx", "pdf", "csv", "json", "python file", "html file", "text file", "markdown", "xlsx", "pptx"]
    action_words = ["analyze", "analyse", "read", "summary", "summarize", "summarise", "create", "make", "generate", "write"]
    return any(x in c for x in file_words) and any(x in c for x in action_words)


def file_command(raw, c):
    if c in ["file help", "file studio", "file status", "document help"]:
        return True, file_studio_help()
    if any(x in c for x in ["analyze", "analyse", "read", "summary", "summarize", "summarise"]):
        path = find_general_file(raw)
        if not path:
            return False, "I could not find a file. Upload it in File/Media Studio or put it on Desktop, Documents, Downloads, Pictures, or Videos."
        return analyze_general_file(path)
    if any(x in c for x in ["create", "make", "generate", "write"]):
        return create_general_file_command(raw, c)
    return False, ""

def screenshot():
    if not pyautogui:
        return False, "PyAutoGUI is not installed."
    p = SCREENSHOT_DIR / f"screenshot_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.png"
    pyautogui.screenshot(str(p))
    return True, f"Screenshot saved: {p.name}."


# -------------------- Media Studio --------------------
def media_help():
    return (
        "Media Studio is ready.\n\n"
        "Use the one media button: select photo/video/file, write what to do, then press play.\n\n"
        "Examples:\n"
        "- analyze this photo\n"
        "- cinematic portrait, keep the same face\n"
        "- crop 9:16 for reels\n"
        "- merge selected photos collage\n"
        "- create image of futuristic blue car in rain\n"
        "- create video of neon intro for Nova\n"
        "- analyze selected document\n\n"
        "Nova now uses only the selected files for that button. It will not mix old/latest photos or videos. Honest limit: local style edits are Pillow/FFmpeg drafts; true identity-preserving background/outfit/object changes need an image/video editing AI API connected."
    )


def pil_available():
    return importlib.util.find_spec("PIL") is not None


def media_tool_status():
    ffmpeg = media_ffmpeg_path()
    ffprobe = media_ffprobe_path()
    status = {
        "pillow": pil_available(),
        "ffmpeg": bool(ffmpeg),
        "ffprobe": bool(ffprobe),
        "ffmpeg_path": ffmpeg,
        "ffprobe_path": ffprobe,
        "gemini_vision": bool(api_key("GEMINI_API_KEY") and requests),
        "media_folder": str(MEDIA_DIR),
    }
    return status


def media_ffmpeg_path():
    local_tool = BASE_DIR / "tools" / "ffmpeg" / "bin" / "ffmpeg.exe"
    if local_tool.exists():
        return str(local_tool)
    found = shutil.which("ffmpeg")
    if found:
        return found
    candidates = []
    for root in [
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "WinGet" / "Packages",
        Path(os.environ.get("PROGRAMFILES", "")),
        Path(os.environ.get("PROGRAMFILES(X86)", "")),
        Path(os.environ.get("PROGRAMDATA", "")) / "Microsoft" / "WinGet" / "Packages",
    ]:
        try:
            if root and root.exists():
                candidates.extend(root.rglob("ffmpeg.exe"))
        except Exception:
            pass
    if candidates:
        candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        return str(candidates[0])
    return ""


def media_ffprobe_path():
    local_tool = BASE_DIR / "tools" / "ffmpeg" / "bin" / "ffprobe.exe"
    if local_tool.exists():
        return str(local_tool)
    found = shutil.which("ffprobe")
    if found:
        return found
    candidates = []
    for root in [
        Path(os.environ.get("LOCALAPPDATA", "")) / "Microsoft" / "WinGet" / "Packages",
        Path(os.environ.get("PROGRAMFILES", "")),
        Path(os.environ.get("PROGRAMFILES(X86)", "")),
        Path(os.environ.get("PROGRAMDATA", "")) / "Microsoft" / "WinGet" / "Packages",
    ]:
        try:
            if root and root.exists():
                candidates.extend(root.rglob("ffprobe.exe"))
        except Exception:
            pass
    if candidates:
        candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        return str(candidates[0])
    return ""


def safe_media_filename(name, default_ext=".bin"):
    raw = Path(str(name or "uploaded_media")).name
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", Path(raw).stem).strip(" ._") or "uploaded_media"
    ext = Path(raw).suffix.lower() or default_ext
    if ext not in IMAGE_EXTS and ext not in VIDEO_EXTS:
        ext = default_ext
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{stem}_{stamp}{ext}"


def safe_upload_filename(name, default_ext=".bin"):
    raw = Path(str(name or "uploaded_file")).name
    stem = re.sub(r"[^A-Za-z0-9._ -]+", "_", Path(raw).stem).strip(" ._") or "uploaded_file"
    ext = Path(raw).suffix.lower() or default_ext
    if ext not in GENERAL_FILE_EXTS:
        ext = default_ext
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"{stem}_{stamp}{ext}"


def media_roots():
    roots = [
        MEDIA_UPLOAD_DIR,
        MEDIA_IMAGE_DIR,
        MEDIA_VIDEO_DIR,
        MEDIA_EDIT_DIR,
        SCREENSHOT_DIR,
        FILES_DIR,
        DESKTOP,
        Path.home() / "Pictures",
        Path.home() / "Downloads",
        Path.home() / "Videos",
    ]
    return [p for p in roots if p and Path(p).exists()]


def iter_media_files(kind="any", recursive_media=True):
    exts = IMAGE_EXTS | VIDEO_EXTS
    if kind == "image":
        exts = IMAGE_EXTS
    elif kind == "video":
        exts = VIDEO_EXTS
    seen = set()
    for root in media_roots():
        root = Path(root)
        iterator = root.rglob("*") if recursive_media and root in [MEDIA_DIR, MEDIA_UPLOAD_DIR, MEDIA_IMAGE_DIR, MEDIA_VIDEO_DIR, MEDIA_EDIT_DIR] else root.glob("*")
        try:
            for path in iterator:
                if not path.is_file() or path.suffix.lower() not in exts:
                    continue
                key = str(path.resolve()).lower()
                if key in seen:
                    continue
                seen.add(key)
                yield path
        except Exception:
            continue


def find_media_file(command, kind="any"):
    raw = str(command or "")
    quoted = re.findall(r'"([^"]+\.(?:png|jpe?g|webp|bmp|gif|tiff?|mp4|mov|mkv|avi|webm|m4v|wmv))"', raw, flags=re.I)
    candidates = quoted[:]
    candidates += re.findall(r"([A-Za-z]:\\[^<>|?*\n\r]+?\.(?:png|jpe?g|webp|bmp|gif|tiff?|mp4|mov|mkv|avi|webm|m4v|wmv))", raw, flags=re.I)
    for item in candidates:
        p = Path(item.strip())
        if p.exists():
            return p
    files = list(iter_media_files(kind))
    if not files:
        return None
    c = normalize_words(raw)
    if any(x in c for x in ["latest", "last", "recent", "uploaded", "photo", "image", "picture", "video"]) or len(c.split()) < 4:
        return max(files, key=lambda p: p.stat().st_mtime)
    query = clean_entity_name(remove_words(c, [
        "analyze", "analyse", "describe", "photo", "image", "picture", "pic", "video", "crop", "edit", "merge",
        "latest", "last", "recent", "file", "make", "create", "generate", "using", "prompt"
    ]))
    if not query:
        return max(files, key=lambda p: p.stat().st_mtime)
    names = {p.stem.lower(): p for p in files}
    for name, path in names.items():
        if query in name or name in query:
            return path
    match = difflib.get_close_matches(query, list(names.keys()), n=1, cutoff=0.55)
    if match:
        return names[match[0]]
    return max(files, key=lambda p: p.stat().st_mtime)


def wrap_text_for_draw(draw, text, font, max_width):
    words = str(text or "").split()
    lines = []
    current = ""
    for word in words:
        trial = (current + " " + word).strip()
        try:
            width = draw.textbbox((0, 0), trial, font=font)[2]
        except Exception:
            width = len(trial) * 12
        if width <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def load_media_font(size=36, bold=False):
    font_candidates = [
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
    ]
    try:
        from PIL import ImageFont
        for font_path in font_candidates:
            if font_path.exists():
                return ImageFont.truetype(str(font_path), size)
        return ImageFont.load_default()
    except Exception:
        return None


def dominant_image_colors(image, count=5):
    small = image.convert("RGB").resize((80, 80))
    quantized = small.quantize(colors=count).convert("RGB")
    colors = quantized.getcolors(80 * 80) or []
    colors.sort(reverse=True, key=lambda item: item[0])
    return [f"#{r:02x}{g:02x}{b:02x}" for _, (r, g, b) in colors[:count]]


def aspect_ratio_label(width, height):
    ratio = width / max(1, height)
    known = [(1.0, "1:1 square"), (16 / 9, "16:9 landscape"), (9 / 16, "9:16 vertical"), (4 / 5, "4:5 portrait"), (3 / 2, "3:2 photo"), (4 / 3, "4:3 classic")]
    best = min(known, key=lambda item: abs(item[0] - ratio))
    if abs(best[0] - ratio) < 0.08:
        return best[1]
    return f"{ratio:.2f}:1 custom"


def image_sharpness_score(image):
    try:
        from PIL import ImageFilter, ImageStat
        edges = image.convert("L").resize((220, 220)).filter(ImageFilter.FIND_EDGES)
        stat = ImageStat.Stat(edges)
        return round(stat.mean[0], 1)
    except Exception:
        return 0


def image_analysis_suggestions(width, height, brightness, sharpness, has_alpha):
    suggestions = []
    ratio = width / max(1, height)
    if brightness < 70:
        suggestions.append("try brighten/enhance")
    elif brightness > 205:
        suggestions.append("try contrast or reduce brightness")
    if sharpness < 8:
        suggestions.append("try sharpen")
    if abs(ratio - 16 / 9) > 0.08:
        suggestions.append("crop 16:9 for video/YouTube")
    if abs(ratio - 9 / 16) > 0.08:
        suggestions.append("crop 9:16 for reels/shorts")
    if has_alpha:
        suggestions.append("transparent image detected; save as PNG to keep alpha")
    if not suggestions:
        suggestions.append("image looks technically ready")
    return suggestions


def open_output_file(path):
    try:
        os.startfile(path)
    except Exception:
        pass


def media_report_path(stem, ext="txt"):
    return MEDIA_DIR / safe_media_filename(stem + "." + ext, "." + ext)


def ffmpeg_run(args, timeout=300):
    exe = media_ffmpeg_path()
    if not exe:
        return None, "FFmpeg is not installed."
    result = subprocess.run([exe] + list(args), capture_output=True, text=True, timeout=timeout)
    if result.returncode != 0:
        return result, (result.stderr or result.stdout or "FFmpeg failed")[-900:]
    return result, ""


def media_still_from_video(path, second=1.0):
    ffmpeg = media_ffmpeg_path()
    if not ffmpeg:
        return None
    out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_thumbnail.jpg", ".jpg")
    try:
        subprocess.run([ffmpeg, "-y", "-ss", str(second), "-i", str(path), "-frames:v", "1", "-q:v", "2", str(out)], capture_output=True, text=True, timeout=60)
        return out if out.exists() and out.stat().st_size > 0 else None
    except Exception as e:
        log(f"Video thumbnail failed: {e}")
        return None


def media_contact_sheet_from_video(path, duration=0.0):
    ffmpeg = media_ffmpeg_path()
    if not ffmpeg or not pil_available():
        return None
    try:
        from PIL import Image, ImageDraw
        seconds = []
        if duration and duration > 3:
            seconds = [max(0.2, duration * x) for x in [0.08, 0.25, 0.42, 0.58, 0.75, 0.92]]
        else:
            seconds = [0.2, 0.8, 1.4, 2.0]
        thumbs = []
        temp_paths = []
        for idx, sec in enumerate(seconds):
            out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_shot_{idx}.jpg", ".jpg")
            subprocess.run([ffmpeg, "-y", "-ss", f"{sec:.2f}", "-i", str(path), "-frames:v", "1", "-q:v", "4", str(out)], capture_output=True, text=True, timeout=45)
            if out.exists() and out.stat().st_size > 0:
                temp_paths.append(out)
                with Image.open(out) as img:
                    thumbs.append(img.convert("RGB").resize((320, 180)))
        if not thumbs:
            return None
        cols = 2
        rows = (len(thumbs) + cols - 1) // cols
        canvas = Image.new("RGB", (cols * 320, rows * 214), (18, 24, 38))
        draw = ImageDraw.Draw(canvas)
        font = load_media_font(16)
        for i, thumb in enumerate(thumbs):
            x = (i % cols) * 320
            y = (i // cols) * 214
            canvas.paste(thumb, (x, y))
            draw.text((x + 8, y + 184), f"{seconds[i]:.1f}s", font=font, fill=(230, 240, 255))
        out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_contact_sheet.jpg", ".jpg")
        canvas.save(out, quality=90)
        for p in temp_paths:
            try:
                p.unlink(missing_ok=True)
            except Exception:
                pass
        return out
    except Exception as e:
        log(f"Video contact sheet failed: {e}")
        return None


def gemini_analyze_image(path, question=""):
    if not settings.get("media_ai_vision_enabled", True) or not requests or not api_key("GEMINI_API_KEY"):
        return ""
    try:
        size_mb = path.stat().st_size / (1024 * 1024)
        if size_mb > 12:
            return "Gemini vision skipped because this image is over 12 MB."
        mime = mimetypes.guess_type(str(path))[0] or "image/png"
        data64 = base64.b64encode(path.read_bytes()).decode("ascii")
        prompt = question.strip() or "Describe this image clearly. Mention visible objects, text, scene, colors, and useful editing suggestions."
        payload = {
            "system_instruction": {"parts": [{"text": system_prompt()}]},
            "contents": [{
                "role": "user",
                "parts": [
                    {"text": prompt},
                    {"inline_data": {"mime_type": mime, "data": data64}},
                ],
            }],
            "generationConfig": {"temperature": 0.35, "maxOutputTokens": 450},
        }
        model = setting_value("gemini_model", "gemini-2.5-flash")
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{model}:generateContent"
        r = requests.post(url, headers={"x-goog-api-key": api_key("GEMINI_API_KEY"), "Content-Type": "application/json"}, json=payload, timeout=45)
        data = r.json()
        if r.status_code >= 400:
            return "Gemini vision failed: " + data.get("error", {}).get("message", "API error")
        return data["candidates"][0]["content"]["parts"][0]["text"].strip()
    except Exception as e:
        log(f"Gemini image analysis failed: {e}")
        return "Gemini vision failed locally: " + str(e)


def analyze_image_file(path, question=""):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    try:
        from PIL import Image, ImageStat
        with Image.open(path) as img:
            width, height = img.size
            stat = ImageStat.Stat(img.convert("L").resize((80, 80)))
            brightness = round(stat.mean[0], 1)
            sharpness = image_sharpness_score(img)
            colors = dominant_image_colors(img)
            has_alpha = img.mode in ["RGBA", "LA"] or ("transparency" in getattr(img, "info", {}))
            frames = getattr(img, "n_frames", 1)
            suggestions = image_analysis_suggestions(width, height, brightness, sharpness, has_alpha)
            info = [
                f"Image analyzed: {path.name}",
                f"- Size: {width} x {height}",
                f"- Aspect: {aspect_ratio_label(width, height)}",
                f"- Format: {img.format or path.suffix.upper().strip('.')}",
                f"- Mode: {img.mode}",
                f"- Frames: {frames}",
                f"- File size: {path.stat().st_size / (1024 * 1024):.2f} MB",
                f"- Brightness: {brightness}/255",
                f"- Sharpness estimate: {sharpness}/255",
                f"- Transparency: {'yes' if has_alpha else 'no'}",
                "- Dominant colors: " + ", ".join(colors),
                "- Suggestions: " + "; ".join(suggestions),
            ]
        ai = gemini_analyze_image(path, question)
        if ai:
            info.extend(["", "AI vision:", ai])
        else:
            info.append("- Semantic AI vision is not connected. Add GEMINI_API_KEY for object/text/scene understanding.")
        report = media_report_path(path.stem + "_image_analysis")
        report.write_text("\n".join(info), encoding="utf-8")
        info.append(f"- Report saved: {report.name}")
        return True, "\n".join(info)
    except Exception as e:
        return False, f"Could not analyze image: {e}"


def analyze_video_file(path):
    ffprobe = media_ffprobe_path()
    lines = [
        f"Video file: {path.name}",
        f"- File size: {path.stat().st_size / (1024 * 1024):.2f} MB",
    ]
    if not ffprobe:
        lines.append("- FFprobe/FFmpeg is not installed, so I can only read basic file info. Install FFmpeg for duration, resolution, crop, and merge.")
        return True, "\n".join(lines)
    try:
        result = subprocess.run([
            ffprobe, "-v", "error", "-print_format", "json", "-show_format", "-show_streams", str(path)
        ], capture_output=True, text=True, timeout=20)
        data = json.loads(result.stdout or "{}")
        video_stream = next((s for s in data.get("streams", []) if s.get("codec_type") == "video"), {})
        audio_stream = next((s for s in data.get("streams", []) if s.get("codec_type") == "audio"), {})
        duration = data.get("format", {}).get("duration", "")
        width = int(video_stream.get("width") or 0) if video_stream else 0
        height = int(video_stream.get("height") or 0) if video_stream else 0
        if video_stream:
            lines.append(f"- Resolution: {width} x {height}")
            if width and height:
                lines.append(f"- Aspect: {aspect_ratio_label(width, height)}")
            lines.append(f"- Video codec: {video_stream.get('codec_name', 'unknown')}")
            if video_stream.get("avg_frame_rate"):
                lines.append(f"- Frame rate: {video_stream.get('avg_frame_rate')}")
        if audio_stream:
            lines.append(f"- Audio codec: {audio_stream.get('codec_name', 'unknown')}")
        dur = float(duration) if duration else 0.0
        if dur:
            lines.append(f"- Duration: {dur:.1f} seconds")
        thumb = media_still_from_video(path, 1.0 if not dur else min(max(dur * 0.18, 0.3), 5.0))
        sheet = media_contact_sheet_from_video(path, dur)
        if thumb:
            lines.append(f"- Thumbnail saved: {thumb.name}")
            ok_img, img_summary = analyze_image_file(thumb, "Analyze this video thumbnail briefly.")
            if ok_img:
                thumb_lines = [line for line in img_summary.splitlines() if line.startswith("- Brightness") or line.startswith("- Dominant") or line.startswith("- Suggestions")]
                if thumb_lines:
                    lines.append("- Thumbnail quick read: " + " | ".join(x[2:] for x in thumb_lines))
        if sheet:
            lines.append(f"- Contact sheet saved: {sheet.name}")
        report = media_report_path(path.stem + "_video_analysis")
        report.write_text("\n".join(lines), encoding="utf-8")
        lines.append(f"- Report saved: {report.name}")
        return True, "\n".join(lines)
    except Exception as e:
        return False, f"Could not analyze video: {e}"


def media_analyze_command(raw, c):
    kind = "video" if "video" in c else "image" if any(x in c for x in ["photo", "image", "picture", "pic"]) else "any"
    path = find_media_file(raw, kind)
    if not path:
        return False, "I could not find an uploaded/local media file. Upload a photo/video in Media Studio or put it in Desktop/Downloads/Pictures/Videos."
    if path.suffix.lower() in IMAGE_EXTS:
        return analyze_image_file(path, raw)
    if path.suffix.lower() in VIDEO_EXTS:
        return analyze_video_file(path)
    return False, "Unsupported media type."


def prompt_from_command(raw, c):
    quoted = re.findall(r'"([^"]+)"', str(raw or ""))
    if quoted:
        return quoted[-1].strip()
    for marker in ["prompt", "of", "for", "about", "ka", "ki"]:
        m = re.search(rf"\b{re.escape(marker)}\b\s+(.+)$", str(raw or ""), flags=re.I)
        if m:
            return m.group(1).strip()
    cleaned = remove_words(c, ["make", "create", "generate", "photo", "image", "picture", "video", "using", "prompt", "ai", "one", "ek", "banao", "banado"])
    return cleaned.strip() or "Nova AI futuristic assistant"


def prompt_canvas_size(c, video=False):
    if "9:16" in c or "vertical" in c or "reel" in c or "short" in c:
        return (720, 1280) if video else (1080, 1920)
    if "16:9" in c or "wide" in c or "wallpaper" in c or "youtube" in c or "landscape" in c:
        return (1280, 720) if video else (1600, 900)
    if "4:5" in c or "portrait" in c:
        return (864, 1080) if video else (1080, 1350)
    return (960, 540) if video else (1024, 1024)


def prompt_palette(prompt):
    p = normalize_words(prompt)
    seed = int(hashlib.sha256(prompt.encode("utf-8")).hexdigest()[:8], 16)
    palettes = {
        "neon": [(44, 212, 255), (110, 231, 183), (255, 77, 109)],
        "rain": [(45, 55, 72), (80, 180, 255), (190, 230, 255)],
        "fire": [(255, 92, 64), (255, 178, 72), (70, 20, 12)],
        "nature": [(46, 125, 86), (142, 201, 107), (232, 245, 210)],
        "space": [(20, 24, 60), (96, 165, 250), (218, 190, 255)],
        "gold": [(255, 209, 102), (245, 158, 11), (42, 33, 18)],
    }
    for key, colors in palettes.items():
        if key in p or (key == "space" and any(x in p for x in ["galaxy", "cosmic", "universe"])):
            return colors
    return [
        ((seed >> 16) & 255, (seed >> 8) & 255, seed & 255),
        ((seed >> 5) & 255, (seed >> 13) & 255, (seed >> 21) & 255),
        (18, 24, 38),
    ]


def draw_prompt_art(prompt, size=(1024, 1024), frame_index=0, total_frames=1):
    from PIL import Image, ImageDraw, ImageFilter
    w, h = size
    seed = int(hashlib.sha256(prompt.encode("utf-8")).hexdigest()[:8], 16)
    colors = prompt_palette(prompt)
    gw, gh = max(160, w // 4), max(160, h // 4)
    img = Image.new("RGB", (gw, gh), colors[-1])
    px = img.load()
    motion = frame_index / max(1, total_frames - 1)
    for y in range(gh):
        for x in range(gw):
            t = (x + y) / max(1, gw + gh)
            wave = (1 + math.sin((x * 0.045) + (y * 0.033) + seed % 31 + motion * math.pi * 2)) / 2
            r = int(colors[0][0] * (1 - t) + colors[1][0] * t)
            g = int(colors[0][1] * (1 - t) + colors[1][1] * t)
            b = int(colors[0][2] * (1 - t) + colors[1][2] * t)
            px[x, y] = (min(255, int(r * (0.72 + wave * 0.36))), min(255, int(g * (0.72 + wave * 0.36))), min(255, int(b * (0.72 + wave * 0.36))))
    try:
        resample = Image.Resampling.BICUBIC
    except Exception:
        resample = Image.BICUBIC
    img = img.resize((w, h), resample).filter(ImageFilter.GaussianBlur(0.7))
    draw = ImageDraw.Draw(img, "RGBA")
    for i in range(10):
        radius = int((min(w, h) * 0.07) + ((seed >> (i % 16)) & int(min(w, h) * 0.16)))
        cx = int((seed * (i + 3) * 37 + motion * w * (i + 1) * 0.18) % w)
        cy = int((seed * (i + 5) * 29 + motion * h * (i + 2) * 0.10) % h)
        draw.ellipse((cx - radius, cy - radius, cx + radius, cy + radius), fill=(255, 255, 255, 14 + (i % 4) * 5))
    title_font = load_media_font(max(32, min(64, w // 18)), bold=True)
    small_font = load_media_font(max(16, min(28, w // 42)))
    panel_h = max(170, min(280, h // 4))
    margin = max(34, w // 18)
    draw.rounded_rectangle((margin, h - panel_h - margin, w - margin, h - margin), radius=28, fill=(5, 10, 18, 172), outline=(255, 255, 255, 70), width=2)
    draw.text((margin + 34, h - panel_h - margin + 32), "Nova Prompt Design", font=small_font, fill=(185, 230, 255, 230))
    lines = wrap_text_for_draw(draw, prompt, title_font, w - margin * 2 - 70)[:3]
    y = h - panel_h - margin + 72
    for line in lines:
        draw.text((margin + 34, y), line, font=title_font, fill=(255, 255, 255, 245))
        y += max(42, int(w // 17))
    return img


def create_prompt_image(raw, c):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    try:
        prompt = prompt_from_command(raw, c)
        img = draw_prompt_art(prompt, prompt_canvas_size(c, video=False))
        path = MEDIA_IMAGE_DIR / safe_media_filename("prompt_image.png", ".png")
        img.save(path)
        meta = MEDIA_IMAGE_DIR / safe_media_filename("prompt_image_prompt.txt", ".txt")
        meta.write_text(f"Prompt: {prompt}\nSize: {img.size[0]} x {img.size[1]}\nCreated: {datetime.datetime.now()}\n", encoding="utf-8")
        open_output_file(path)
        return True, f"Created local prompt image: {path.name} ({img.size[0]}x{img.size[1]}). Prompt saved: {meta.name}. This is local generated design, not photoreal AI model output."
    except Exception as e:
        return False, f"Could not create image: {e}"


def crop_box_for_ratio(width, height, ratio):
    current = width / max(1, height)
    if current > ratio:
        new_w = int(height * ratio)
        left = (width - new_w) // 2
        return (left, 0, left + new_w, height)
    new_h = int(width / ratio)
    top = (height - new_h) // 2
    return (0, top, width, top + new_h)


def parse_crop_ratio(c):
    if "square" in c or "1:1" in c:
        return 1.0, "square"
    if "9:16" in c or "vertical" in c or "reel" in c or "short" in c:
        return 9 / 16, "9x16"
    if "16:9" in c or "wide" in c or "landscape" in c:
        return 16 / 9, "16x9"
    if "4:5" in c or "portrait" in c:
        return 4 / 5, "4x5"
    return 1.0, "square"


def crop_image_command(raw, c):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    path = find_media_file(raw, "image")
    if not path:
        return False, "I could not find an image to crop. Upload/select a photo first."
    try:
        from PIL import Image
        ratio, label = parse_crop_ratio(c)
        with Image.open(path) as img:
            box = crop_box_for_ratio(*img.size, ratio)
            cropped = img.crop(box)
            out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_crop_{label}.png", ".png")
            cropped.save(out)
        try:
            os.startfile(out)
        except Exception:
            pass
        return True, f"Cropped {path.name} to {label}: {out.name}."
    except Exception as e:
        return False, f"Could not crop image: {e}"


def edit_image_command(raw, c):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    path = find_media_file(raw, "image")
    if not path:
        return False, "I could not find an image to edit. Upload/select a photo first."
    try:
        from PIL import Image, ImageEnhance, ImageFilter, ImageOps, ImageDraw
        with Image.open(path) as img:
            edited = img.convert("RGB")
            actions = []
            if "resize" in c or "compress" in c:
                max_side = 1280
                m = re.search(r"\b(\d{3,4})\s*(?:px|pixel|pixels)?\b", c)
                if m:
                    max_side = max(320, min(4096, int(m.group(1))))
                edited.thumbnail((max_side, max_side))
                actions.append(f"resize_{max_side}")
            if any(x in c for x in ["gray", "grey", "black and white", "bw"]):
                edited = ImageOps.grayscale(edited).convert("RGB")
                actions.append("grayscale")
            if "sepia" in c:
                gray = ImageOps.grayscale(edited)
                sepia = ImageOps.colorize(gray, "#2f1b0c", "#ffd79b")
                edited = sepia.convert("RGB")
                actions.append("sepia")
            if "blur" in c:
                edited = edited.filter(ImageFilter.GaussianBlur(3))
                actions.append("blur")
            if "sharpen" in c or "sharp" in c:
                edited = edited.filter(ImageFilter.SHARPEN)
                actions.append("sharpen")
            if "bright" in c or "brightness" in c:
                edited = ImageEnhance.Brightness(edited).enhance(1.25)
                actions.append("brighten")
            if "contrast" in c:
                edited = ImageEnhance.Contrast(edited).enhance(1.25)
                actions.append("contrast")
            if "rotate left" in c:
                edited = edited.rotate(90, expand=True)
                actions.append("rotate_left")
            elif "rotate" in c:
                edited = edited.rotate(-90, expand=True)
                actions.append("rotate_right")
            if "mirror" in c or "flip" in c:
                edited = ImageOps.mirror(edited)
                actions.append("mirror")
            if "vignette" in c:
                overlay = Image.new("RGBA", edited.size, (0, 0, 0, 0))
                mask = Image.new("L", edited.size, 0)
                d = ImageDraw.Draw(mask)
                margin = int(min(edited.size) * 0.08)
                d.ellipse((margin, margin, edited.width - margin, edited.height - margin), fill=255)
                mask = mask.filter(ImageFilter.GaussianBlur(int(min(edited.size) * 0.16)))
                vignette = Image.new("RGBA", edited.size, (0, 0, 0, 150))
                overlay.paste(vignette, (0, 0), ImageOps.invert(mask))
                edited = Image.alpha_composite(edited.convert("RGBA"), overlay).convert("RGB")
                actions.append("vignette")
            semantic_edit = any(x in c for x in ["remove background", "change background", "replace background", "remove object", "change face", "add object", "generative fill"])
            if semantic_edit and not actions:
                return False, "That is prompt-based AI photo editing. Local Nova can crop, merge, rotate, blur, sharpen, grayscale, brighten, and enhance now; semantic edits need an image-editing AI model/API connected."
            if not actions:
                edited = ImageEnhance.Color(edited).enhance(1.15)
                edited = ImageEnhance.Contrast(edited).enhance(1.10)
                actions.append("auto_enhance")
            out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_{'_'.join(actions)}.png", ".png")
            edited.save(out)
        try:
            os.startfile(out)
        except Exception:
            pass
        return True, f"Edited {path.name} with {', '.join(actions)}: {out.name}."
    except Exception as e:
        return False, f"Could not edit image: {e}"


def merge_images_command(raw, c):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    files = sorted(list(iter_media_files("image")), key=lambda p: p.stat().st_mtime, reverse=True)[:6]
    if len(files) < 2:
        return False, "I need at least two recent/uploaded images to merge."
    try:
        from PIL import Image, ImageOps, ImageDraw
        images = []
        labels = []
        for p in files:
            with Image.open(p) as img:
                images.append(ImageOps.contain(img.convert("RGB"), (640, 640)))
                labels.append(p.stem[:36])
        font = load_media_font(18)
        vertical = "vertical" in c or "up down" in c or "top bottom" in c
        horizontal = "horizontal" in c or "side by side" in c
        if vertical:
            w = max(img.width for img in images)
            h = sum(img.height + 34 for img in images)
            canvas = Image.new("RGB", (w, h), (18, 24, 38))
            draw = ImageDraw.Draw(canvas)
            y = 0
            for img, label in zip(images, labels):
                canvas.paste(img, ((w - img.width) // 2, y))
                draw.text((10, y + img.height + 7), label, font=font, fill=(220, 235, 255))
                y += img.height + 34
        elif horizontal:
            h = max(img.height for img in images) + 36
            w = sum(img.width for img in images)
            canvas = Image.new("RGB", (w, h), (18, 24, 38))
            draw = ImageDraw.Draw(canvas)
            x = 0
            for img, label in zip(images, labels):
                canvas.paste(img, (x, 0))
                draw.text((x + 10, img.height + 8), label, font=font, fill=(220, 235, 255))
                x += img.width
        else:
            cols = 2 if len(images) <= 4 else 3
            cell = 640
            rows = (len(images) + cols - 1) // cols
            canvas = Image.new("RGB", (cols * cell, rows * (cell + 34)), (18, 24, 38))
            draw = ImageDraw.Draw(canvas)
            for i, (img, label) in enumerate(zip(images, labels)):
                x = (i % cols) * cell + (cell - img.width) // 2
                y = (i // cols) * (cell + 34) + (cell - img.height) // 2
                canvas.paste(img, (x, y))
                draw.text(((i % cols) * cell + 12, (i // cols) * (cell + 34) + cell + 7), label, font=font, fill=(220, 235, 255))
        out = MEDIA_EDIT_DIR / safe_media_filename("merged_images.png", ".png")
        canvas.save(out)
        open_output_file(out)
        return True, f"Merged {len(images)} latest images into {out.name}."
    except Exception as e:
        return False, f"Could not merge images: {e}"


def create_prompt_video(raw, c):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    try:
        prompt = prompt_from_command(raw, c)
        size = prompt_canvas_size(c, video=True)
        frame_count = 72 if "long" in c else 48
        fps = 12
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        frame_dir = MEDIA_VIDEO_DIR / f"prompt_video_frames_{stamp}"
        frame_dir.mkdir(parents=True, exist_ok=True)
        frames = []
        for i in range(frame_count):
            frame = draw_prompt_art(prompt, size, i, frame_count)
            frame_path = frame_dir / f"frame_{i:04d}.png"
            frame.save(frame_path)
            if i < 36:
                frames.append(frame.resize((480, int(480 * size[1] / size[0]))))
        ffmpeg = media_ffmpeg_path()
        mp4 = MEDIA_VIDEO_DIR / safe_media_filename("prompt_video.mp4", ".mp4")
        if ffmpeg:
            result, err = ffmpeg_run([
                "-y", "-framerate", str(fps), "-i", str(frame_dir / "frame_%04d.png"),
                "-vf", "format=yuv420p", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-movflags", "+faststart", str(mp4)
            ], timeout=240)
            try:
                shutil.rmtree(frame_dir, ignore_errors=True)
            except Exception:
                pass
            if result and result.returncode == 0 and mp4.exists():
                meta = MEDIA_VIDEO_DIR / safe_media_filename("prompt_video_prompt.txt", ".txt")
                meta.write_text(f"Prompt: {prompt}\nSize: {size[0]} x {size[1]}\nFPS: {fps}\nFrames: {frame_count}\nCreated: {datetime.datetime.now()}\n", encoding="utf-8")
                open_output_file(mp4)
                return True, f"Created MP4 prompt video: {mp4.name} ({size[0]}x{size[1]}, {frame_count / fps:.1f}s). Prompt saved: {meta.name}. This is local generated design video."
            log("Prompt video MP4 failed: " + err)
        gif = MEDIA_VIDEO_DIR / safe_media_filename("prompt_video.gif", ".gif")
        frames[0].save(gif, save_all=True, append_images=frames[1:], duration=int(1000 / fps), loop=0)
        try:
            shutil.rmtree(frame_dir, ignore_errors=True)
        except Exception:
            pass
        open_output_file(gif)
        return True, f"Created animated GIF video draft: {gif.name}. MP4 creation was unavailable or failed; FFmpeg detail is in nova_log.txt."
    except Exception as e:
        return False, f"Could not create video draft: {e}"


def merge_videos_command(raw, c):
    ffmpeg = media_ffmpeg_path()
    if not ffmpeg:
        return False, "FFmpeg is not installed, so I cannot merge videos yet. Install FFmpeg, then say merge latest videos again."
    files = sorted(list(iter_media_files("video")), key=lambda p: p.stat().st_mtime, reverse=True)[:6]
    if len(files) < 2:
        return False, "I need at least two recent/uploaded videos to merge."
    try:
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        list_file = MEDIA_VIDEO_DIR / f"concat_{stamp}.txt"
        out = MEDIA_VIDEO_DIR / f"merged_videos_{stamp}.mp4"
        lines = []
        for p in reversed(files):
            safe = str(p.resolve()).replace("'", "'\\''")
            lines.append(f"file '{safe}'")
        list_file.write_text("\n".join(lines), encoding="utf-8")
        result = subprocess.run([ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", str(out)], capture_output=True, text=True, timeout=180)
        if result.returncode != 0:
            result = subprocess.run([ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c:v", "libx264", "-preset", "veryfast", "-crf", "21", "-c:a", "aac", "-movflags", "+faststart", str(out)], capture_output=True, text=True, timeout=300)
        if result.returncode != 0:
            normalized_dir = MEDIA_VIDEO_DIR / f"normalized_{stamp}"
            normalized_dir.mkdir(parents=True, exist_ok=True)
            normalized = []
            for idx, src in enumerate(reversed(files)):
                norm = normalized_dir / f"part_{idx:03d}.mp4"
                subprocess.run([
                    ffmpeg, "-y", "-i", str(src), "-vf", "scale=1280:-2,fps=30,format=yuv420p",
                    "-c:v", "libx264", "-preset", "veryfast", "-crf", "23", "-c:a", "aac", "-ar", "44100", "-ac", "2", str(norm)
                ], capture_output=True, text=True, timeout=300)
                if norm.exists() and norm.stat().st_size > 0:
                    normalized.append(norm)
            if len(normalized) >= 2:
                list_file.write_text("\n".join(f"file '{str(p.resolve()).replace(chr(39), chr(39) + chr(92) + chr(39) + chr(39))}'" for p in normalized), encoding="utf-8")
                result = subprocess.run([ffmpeg, "-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", "-movflags", "+faststart", str(out)], capture_output=True, text=True, timeout=300)
        if result.returncode != 0:
            return False, "FFmpeg tried to merge the videos but failed: " + (result.stderr or result.stdout)[-500:]
        sheet = media_contact_sheet_from_video(out)
        open_output_file(out)
        extra = f" Contact sheet: {sheet.name}." if sheet else ""
        return True, f"Merged {len(files)} latest videos into {out.name}.{extra}"
    except Exception as e:
        return False, f"Could not merge videos: {e}"


def crop_video_command(raw, c):
    ffmpeg = media_ffmpeg_path()
    if not ffmpeg:
        return False, "FFmpeg is not installed, so I cannot crop videos yet. Install FFmpeg, then try again."
    path = find_media_file(raw, "video")
    if not path:
        return False, "I could not find a video to crop. Upload/select a video first."
    ratio, label = parse_crop_ratio(c)
    if label == "square":
        filt = "crop=w='min(iw,ih)':h='min(iw,ih)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    elif label == "9x16":
        filt = "crop=w='min(iw,ih*9/16)':h='min(ih,iw*16/9)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    elif label == "16x9":
        filt = "crop=w='min(iw,ih*16/9)':h='min(ih,iw*9/16)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    else:
        filt = "crop=w='min(iw,ih*4/5)':h='min(ih,iw*5/4)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    out = MEDIA_VIDEO_DIR / safe_media_filename(f"{path.stem}_crop_{label}.mp4", ".mp4")
    try:
        result = subprocess.run([ffmpeg, "-y", "-i", str(path), "-vf", filt, "-c:v", "libx264", "-preset", "veryfast", "-crf", "20", "-c:a", "aac", "-movflags", "+faststart", str(out)], capture_output=True, text=True, timeout=300)
        if result.returncode != 0:
            return False, "FFmpeg tried to crop the video but failed: " + (result.stderr or result.stdout)[-500:]
        thumb = media_still_from_video(out, 1.0)
        open_output_file(out)
        if thumb:
            return True, f"Cropped video to {label}: {out.name}. Preview thumbnail: {thumb.name}."
        return True, f"Cropped video to {label}: {out.name}."
    except Exception as e:
        return False, f"Could not crop video: {e}"


PHOTO_STYLE_PRESETS = [
    ("cinematic", ["cinematic", "movie poster", "dramatic lighting"], "cinematic portrait with dramatic lighting, shallow depth of field, professional color grading"),
    ("instagram", ["instagram", "model style", "influencer", "fashion"], "stylish Instagram fashion portrait with soft lighting and clean modern colors"),
    ("royal", ["royal", "king", "queen", "palace", "crown", "wedding"], "royal luxury portrait with warm gold lighting and premium editorial mood"),
    ("cyberpunk", ["cyberpunk", "neon", "futuristic city", "sci fi", "sci-fi"], "cyberpunk neon style with blue and pink glow, cinematic shadows"),
    ("business", ["business", "linkedin", "professional", "passport", "corporate"], "professional business portrait with clean light and natural sharp details"),
    ("bollywood", ["bollywood", "hero style", "movie hero"], "Bollywood hero movie-poster style with cinematic glow and sharp face details"),
    ("anime", ["anime", "manga"], "high quality anime-inspired character style with stronger colors and clean edges"),
    ("luxury", ["luxury car", "sports car", "premium look"], "premium luxury photoshoot look with glossy contrast and city-night mood"),
    ("warrior", ["warrior", "armor", "battlefield"], "powerful warrior poster mood with intense contrast and warm highlights"),
    ("vintage", ["90s", "vintage", "retro", "old camera", "film grain"], "90s vintage film style with warm faded colors and grain"),
    ("streetwear", ["streetwear", "hoodie", "sneakers", "urban"], "modern streetwear fashion look with urban contrast"),
    ("ai futuristic", ["ai futuristic", "hologram", "robotic", "tech background"], "futuristic AI character look with cool blue tech lighting"),
    ("black studio", ["black background", "studio portrait", "spotlight"], "premium black-background studio portrait with soft spotlight"),
    ("nature", ["nature", "forest", "mountain", "beach", "travel blogger"], "bright natural outdoor aesthetic with golden light"),
    ("mafia", ["mafia", "gangster", "villain", "dark villain", "horror"], "dark cinematic poster look with deep shadows and controlled red mood"),
    ("fitness", ["gym", "fitness", "athletic"], "fitness photoshoot look with strong contrast and crisp details"),
    ("realistic", ["4k", "realistic", "dslr", "enhance", "ultra realistic", "high quality"], "realistic 4K DSLR upgrade with natural texture, sharp detail, and clean colors"),
    ("painting", ["oil painting", "painting", "brush strokes"], "oil-painting inspired texture and museum-style color"),
    ("pixar", ["pixar", "3d animated", "3d render"], "cute 3D animated character-inspired color and smoothness"),
    ("superhero", ["superhero"], "superhero poster mood with punchy contrast and cinematic glow"),
]


def matched_photo_style_prompt(prompt):
    c = normalize_words(prompt)
    matches = []
    for name, aliases, description in PHOTO_STYLE_PRESETS:
        if any(alias in c for alias in aliases):
            matches.append((name, description))
    direct_style = any(x in c for x in ["transform this photo", "turn this photo", "make this photo", "edit this photo", "keep the face", "same face"])
    if matches:
        names = ", ".join(name for name, _ in matches[:3])
        descriptions = "; ".join(desc for _, desc in matches[:3])
        return names, descriptions
    if direct_style:
        return "prompt style", str(prompt).strip()
    return "", ""


def local_semantic_limit_note(prompt):
    c = normalize_words(prompt)
    semantic_words = [
        "change background", "replace background", "remove background", "add background", "new background",
        "crown", "palace", "armor", "sports car", "car", "outfit", "clothes", "suit", "city background",
        "forest background", "beach background", "mountain background", "add object", "remove object",
        "face identity", "keep same face", "same face", "do not change face",
    ]
    if any(word in c for word in semantic_words):
        return (
            "Honest limit: I made a local style draft on the selected photo only. "
            "Local tools can enhance, grade, crop, blur, sharpen, rotate, mirror, and add simple poster mood. "
            "True background/outfit/object replacement or identity-preserving AI editing needs an image-editing AI API/model connected."
        )
    return ""


def apply_photo_style_grade(image, prompt):
    from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageOps
    c = normalize_words(prompt)
    edited = image.convert("RGB")
    actions = []

    def blend(color, alpha):
        nonlocal edited
        edited = Image.blend(edited, Image.new("RGB", edited.size, color), alpha)

    def vignette(strength=130):
        nonlocal edited
        overlay = Image.new("RGBA", edited.size, (0, 0, 0, 0))
        mask = Image.new("L", edited.size, 0)
        draw = ImageDraw.Draw(mask)
        margin = int(min(edited.size) * 0.08)
        draw.ellipse((margin, margin, edited.width - margin, edited.height - margin), fill=255)
        mask = mask.filter(ImageFilter.GaussianBlur(int(min(edited.size) * 0.18)))
        dark = Image.new("RGBA", edited.size, (0, 0, 0, strength))
        overlay.paste(dark, (0, 0), ImageOps.invert(mask))
        edited = Image.alpha_composite(edited.convert("RGBA"), overlay).convert("RGB")

    if "resize" in c or "compress" in c:
        max_side = 1280
        m = re.search(r"\b(\d{3,4})\s*(?:px|pixel|pixels)?\b", c)
        if m:
            max_side = max(320, min(4096, int(m.group(1))))
        edited.thumbnail((max_side, max_side))
        actions.append(f"resize_{max_side}")
    if any(x in c for x in ["gray", "grey", "black and white", "bw"]):
        edited = ImageOps.grayscale(edited).convert("RGB")
        actions.append("grayscale")
    if any(x in c for x in ["sepia", "vintage", "90s", "retro", "old camera", "film grain"]):
        gray = ImageOps.grayscale(edited)
        edited = ImageOps.colorize(gray, "#2d1a10", "#ffd6a3").convert("RGB")
        edited = ImageEnhance.Contrast(edited).enhance(0.95)
        actions.append("vintage_sepia")
    if "blur" in c:
        edited = edited.filter(ImageFilter.GaussianBlur(3))
        actions.append("blur")
    if "sharpen" in c or "sharp" in c:
        edited = edited.filter(ImageFilter.SHARPEN)
        actions.append("sharpen")
    if "bright" in c or "brightness" in c:
        edited = ImageEnhance.Brightness(edited).enhance(1.22)
        actions.append("brighten")
    if "contrast" in c:
        edited = ImageEnhance.Contrast(edited).enhance(1.22)
        actions.append("contrast")
    if "rotate left" in c:
        edited = edited.rotate(90, expand=True)
        actions.append("rotate_left")
    elif "rotate" in c:
        edited = edited.rotate(-90, expand=True)
        actions.append("rotate_right")
    if "mirror" in c or "flip" in c:
        edited = ImageOps.mirror(edited)
        actions.append("mirror")

    style_name, _ = matched_photo_style_prompt(prompt)
    if style_name:
        if any(x in style_name for x in ["cinematic", "bollywood", "superhero"]):
            edited = ImageEnhance.Contrast(edited).enhance(1.20)
            edited = ImageEnhance.Color(edited).enhance(1.12)
            vignette(120)
            actions.append("cinematic_grade")
        if any(x in style_name for x in ["instagram", "business", "realistic", "fitness", "streetwear", "luxury"]):
            edited = ImageEnhance.Brightness(edited).enhance(1.06)
            edited = ImageEnhance.Contrast(edited).enhance(1.12)
            edited = ImageEnhance.Color(edited).enhance(1.10)
            edited = edited.filter(ImageFilter.SHARPEN)
            actions.append("clean_dslr_enhance")
        if any(x in style_name for x in ["royal", "nature", "warrior"]):
            blend((255, 203, 104), 0.10)
            edited = ImageEnhance.Contrast(edited).enhance(1.12)
            actions.append("warm_gold_grade")
        if any(x in style_name for x in ["cyberpunk", "ai futuristic"]):
            blend((24, 180, 255), 0.12)
            edited = ImageEnhance.Color(edited).enhance(1.35)
            edited = ImageEnhance.Contrast(edited).enhance(1.22)
            vignette(105)
            actions.append("neon_cyberpunk_grade")
        if "anime" in style_name:
            edited = ImageEnhance.Color(edited).enhance(1.45)
            edited = edited.filter(ImageFilter.SMOOTH_MORE).filter(ImageFilter.EDGE_ENHANCE)
            actions.append("anime_inspired")
        if "painting" in style_name:
            edited = edited.filter(ImageFilter.SMOOTH_MORE).filter(ImageFilter.CONTOUR)
            actions.append("painting_texture")
        if "pixar" in style_name:
            edited = ImageEnhance.Color(edited).enhance(1.28)
            edited = edited.filter(ImageFilter.SMOOTH_MORE)
            actions.append("soft_3d_color")
        if any(x in style_name for x in ["black studio", "mafia"]):
            edited = ImageEnhance.Contrast(edited).enhance(1.28)
            vignette(160)
            if "mafia" in style_name:
                blend((80, 8, 16), 0.08)
            actions.append("dark_studio_grade")

    if not actions:
        edited = ImageEnhance.Color(edited).enhance(1.12)
        edited = ImageEnhance.Contrast(edited).enhance(1.10)
        edited = edited.filter(ImageFilter.SHARPEN)
        actions.append("auto_enhance")
    return edited, actions


def selected_crop_images(paths, prompt):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    try:
        from PIL import Image
        c = normalize_words(prompt)
        ratio, label = parse_crop_ratio(c)
        outputs = []
        for idx, path in enumerate(paths, 1):
            with Image.open(path) as img:
                cropped = img.crop(crop_box_for_ratio(*img.size, ratio))
                out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_selected_{idx}_crop_{label}.png", ".png")
                cropped.save(out)
            outputs.append(out.name)
        if outputs:
            open_output_file(MEDIA_EDIT_DIR / outputs[0])
        return True, f"Cropped {len(outputs)} selected photo(s) to {label}: " + ", ".join(outputs)
    except Exception as e:
        return False, f"Could not crop selected photo(s): {e}"


def selected_edit_images(paths, prompt):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    try:
        from PIL import Image
        style_name, expanded = matched_photo_style_prompt(prompt)
        outputs = []
        action_names = []
        for idx, path in enumerate(paths, 1):
            with Image.open(path) as img:
                edited, actions = apply_photo_style_grade(img, prompt)
                tag = (style_name or "_".join(actions[:2]) or "edit").replace(" ", "_").replace(",", "")
                out = MEDIA_EDIT_DIR / safe_media_filename(f"{path.stem}_selected_{idx}_{tag}.png", ".png")
                edited.save(out)
            outputs.append(out.name)
            action_names.extend(actions)
        if outputs:
            open_output_file(MEDIA_EDIT_DIR / outputs[0])
        unique_actions = ", ".join(dict.fromkeys(action_names))
        lines = [
            f"Edited {len(outputs)} selected photo(s) only: " + ", ".join(outputs),
            f"Local actions: {unique_actions}.",
        ]
        if expanded:
            lines.append("Style understood: " + expanded)
        limit = local_semantic_limit_note(prompt)
        if limit:
            lines.append(limit)
        return True, "\n".join(lines)
    except Exception as e:
        return False, f"Could not edit selected photo(s): {e}"


def selected_merge_images(paths, prompt):
    if not pil_available():
        return False, "Pillow is missing. Run INSTALL_AND_RUN.bat or install Pillow."
    if len(paths) < 2:
        return False, "Select at least two photos to merge."
    try:
        from PIL import Image, ImageDraw, ImageOps
        c = normalize_words(prompt)
        images = []
        labels = []
        for p in paths[:12]:
            with Image.open(p) as img:
                images.append(ImageOps.contain(img.convert("RGB"), (720, 720)))
                labels.append(p.stem[:36])
        font = load_media_font(18)
        vertical = "vertical" in c or "up down" in c or "top bottom" in c
        horizontal = "horizontal" in c or "side by side" in c
        if vertical:
            w = max(img.width for img in images)
            h = sum(img.height + 34 for img in images)
            canvas = Image.new("RGB", (w, h), (18, 24, 38))
            draw = ImageDraw.Draw(canvas)
            y = 0
            for img, label in zip(images, labels):
                canvas.paste(img, ((w - img.width) // 2, y))
                draw.text((10, y + img.height + 7), label, font=font, fill=(220, 235, 255))
                y += img.height + 34
        elif horizontal:
            h = max(img.height for img in images) + 36
            w = sum(img.width for img in images)
            canvas = Image.new("RGB", (w, h), (18, 24, 38))
            draw = ImageDraw.Draw(canvas)
            x = 0
            for img, label in zip(images, labels):
                canvas.paste(img, (x, 0))
                draw.text((x + 10, img.height + 8), label, font=font, fill=(220, 235, 255))
                x += img.width
        else:
            cols = 2 if len(images) <= 4 else 3
            cell = 720
            rows = (len(images) + cols - 1) // cols
            canvas = Image.new("RGB", (cols * cell, rows * (cell + 34)), (18, 24, 38))
            draw = ImageDraw.Draw(canvas)
            for i, (img, label) in enumerate(zip(images, labels)):
                x = (i % cols) * cell + (cell - img.width) // 2
                y = (i // cols) * (cell + 34) + (cell - img.height) // 2
                canvas.paste(img, (x, y))
                draw.text(((i % cols) * cell + 12, (i // cols) * (cell + 34) + cell + 7), label, font=font, fill=(220, 235, 255))
        out = MEDIA_EDIT_DIR / safe_media_filename("selected_merged_photos.png", ".png")
        canvas.save(out)
        open_output_file(out)
        return True, f"Merged {len(images)} selected photos only into {out.name}."
    except Exception as e:
        return False, f"Could not merge selected photos: {e}"


def video_crop_filter_for_label(label):
    if label == "square":
        return "crop=w='min(iw,ih)':h='min(iw,ih)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    if label == "9x16":
        return "crop=w='min(iw,ih*9/16)':h='min(ih,iw*16/9)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    if label == "16x9":
        return "crop=w='min(iw,ih*16/9)':h='min(ih,iw*9/16)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"
    return "crop=w='min(iw,ih*4/5)':h='min(ih,iw*5/4)':x='(iw-ow)/2':y='(ih-oh)/2',scale=trunc(iw/2)*2:trunc(ih/2)*2"


def selected_crop_videos(paths, prompt):
    if not media_ffmpeg_path():
        return False, "FFmpeg is not installed, so I cannot crop selected videos yet. Run INSTALL_MEDIA_TOOLS.cmd first."
    c = normalize_words(prompt)
    ratio, label = parse_crop_ratio(c)
    filt = video_crop_filter_for_label(label)
    outputs = []
    try:
        for idx, path in enumerate(paths, 1):
            out = MEDIA_VIDEO_DIR / safe_media_filename(f"{path.stem}_selected_{idx}_crop_{label}.mp4", ".mp4")
            result, err = ffmpeg_run(["-y", "-i", str(path), "-vf", filt, "-c:v", "libx264", "-preset", "veryfast", "-crf", "20", "-c:a", "aac", "-movflags", "+faststart", str(out)], timeout=300)
            if not result or result.returncode != 0:
                return False, "FFmpeg tried to crop the selected video but failed: " + err
            outputs.append(out.name)
        if outputs:
            open_output_file(MEDIA_VIDEO_DIR / outputs[0])
        return True, f"Cropped {len(outputs)} selected video(s) to {label}: " + ", ".join(outputs)
    except Exception as e:
        return False, f"Could not crop selected video(s): {e}"


def selected_merge_videos(paths, prompt):
    if not media_ffmpeg_path():
        return False, "FFmpeg is not installed, so I cannot merge videos yet. Run INSTALL_MEDIA_TOOLS.cmd first."
    if len(paths) < 2:
        return False, "Select at least two videos to merge."
    try:
        stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        list_file = MEDIA_VIDEO_DIR / f"selected_concat_{stamp}.txt"
        out = MEDIA_VIDEO_DIR / f"selected_merged_videos_{stamp}.mp4"
        list_file.write_text("\n".join(f"file '{str(p.resolve()).replace(chr(92), '/').replace(chr(39), chr(39) + chr(92) + chr(39) + chr(39))}'" for p in paths), encoding="utf-8")
        result, err = ffmpeg_run(["-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c", "copy", "-movflags", "+faststart", str(out)], timeout=240)
        if not result or result.returncode != 0:
            result, err = ffmpeg_run(["-y", "-f", "concat", "-safe", "0", "-i", str(list_file), "-c:v", "libx264", "-preset", "veryfast", "-crf", "21", "-c:a", "aac", "-movflags", "+faststart", str(out)], timeout=360)
        if not result or result.returncode != 0:
            return False, "FFmpeg tried to merge the selected videos but failed: " + err
        open_output_file(out)
        sheet = media_contact_sheet_from_video(out)
        extra = f" Preview contact sheet: {sheet.name}." if sheet else ""
        return True, f"Merged {len(paths)} selected videos only into {out.name}.{extra}"
    except Exception as e:
        return False, f"Could not merge selected videos: {e}"


def selected_edit_videos(paths, prompt):
    if not media_ffmpeg_path():
        return False, "FFmpeg is not installed, so I cannot edit selected videos yet. Run INSTALL_MEDIA_TOOLS.cmd first."
    c = normalize_words(prompt)
    filters = []
    actions = []
    if any(x in c for x in ["gray", "grey", "black and white", "bw"]):
        filters.append("format=gray")
        actions.append("grayscale")
    if any(x in c for x in ["bright", "brightness", "enhance", "contrast"]):
        filters.append("eq=contrast=1.16:saturation=1.12:brightness=0.04")
        actions.append("enhance")
    if "mirror" in c or "flip" in c:
        filters.append("hflip")
        actions.append("mirror")
    if "rotate" in c:
        filters.append("transpose=1")
        actions.append("rotate")
    mute = any(x in c for x in ["mute", "remove audio", "no audio"])
    if mute:
        actions.append("mute")
    if not filters and not mute:
        return False, "For selected videos I can crop, merge, grayscale, brighten/enhance, rotate, mirror, or mute locally. Prompt-based semantic video editing needs a video AI model/API."
    outputs = []
    try:
        for idx, path in enumerate(paths, 1):
            out = MEDIA_VIDEO_DIR / safe_media_filename(f"{path.stem}_selected_{idx}_{'_'.join(actions)}.mp4", ".mp4")
            args = ["-y", "-i", str(path)]
            if filters:
                args += ["-vf", ",".join(filters), "-c:v", "libx264", "-preset", "veryfast", "-crf", "21"]
            else:
                args += ["-c:v", "copy"]
            if mute:
                args += ["-an"]
            else:
                args += ["-c:a", "aac"]
            args += ["-movflags", "+faststart", str(out)]
            result, err = ffmpeg_run(args, timeout=300)
            if not result or result.returncode != 0:
                return False, "FFmpeg tried to edit the selected video but failed: " + err
            outputs.append(out.name)
        if outputs:
            open_output_file(MEDIA_VIDEO_DIR / outputs[0])
        return True, f"Edited {len(outputs)} selected video(s) with {', '.join(actions)}: " + ", ".join(outputs)
    except Exception as e:
        return False, f"Could not edit selected video(s): {e}"


def analyze_selected_paths(paths, prompt=""):
    replies = []
    ok_all = True
    for path in paths:
        ext = path.suffix.lower()
        if ext in IMAGE_EXTS:
            ok, msg = analyze_image_file(path, prompt)
        elif ext in VIDEO_EXTS:
            ok, msg = analyze_video_file(path)
        else:
            ok, msg = analyze_general_file(path)
        ok_all = ok_all and ok
        replies.append(msg)
    return ok_all, "\n\n".join(replies)


def smart_selected_media_task(prompt, paths):
    raw = str(prompt or "").strip()
    c = normalize_words(raw)
    paths = [Path(p) for p in paths if p]
    images = [p for p in paths if p.suffix.lower() in IMAGE_EXTS]
    videos = [p for p in paths if p.suffix.lower() in VIDEO_EXTS]
    files = [p for p in paths if p.suffix.lower() not in IMAGE_EXTS | VIDEO_EXTS]

    if not paths:
        if any(x in c for x in ["file", "document", "docx", "xlsx", "pptx", "pdf", "csv", "json", "python file", "html file"]) and any(x in c for x in ["create", "make", "generate", "write"]):
            return create_general_file_command(raw, c)
        if "video" in c or "gif" in c:
            return create_prompt_video(raw, c)
        if any(x in c for x in ["image", "photo", "picture", "wallpaper", "poster", "thumbnail", "create", "make", "generate"]):
            return create_prompt_image(raw, c)
        return False, "Type what to create or select a file first. Example: select a photo and write cinematic portrait, or type create image cyberpunk Nova assistant."

    if files and (images or videos):
        return False, "You selected files mixed with photos/videos. I will not mix them. Select only files, only photos, or only videos for one task."
    if images and videos:
        wants_photo = any(x in c for x in ["photo", "image", "picture", "pic"]) and "video" not in c
        wants_video = "video" in c and not any(x in c for x in ["photo", "image", "picture", "pic"])
        if wants_photo:
            videos = []
        elif wants_video:
            images = []
        else:
            return False, "You selected both photos and videos. I will not mix them. Select only photos or only videos, or say clearly: use photos / use videos."

    if files:
        if any(x in c for x in ["create", "make", "generate", "write"]) and not any(x in c for x in ["analyze", "analyse", "read", "summary", "summarize"]):
            return create_general_file_command(raw, c)
        return analyze_selected_paths(files, raw)

    if images:
        if any(x in c for x in ["analyze", "analyse", "describe", "what is in", "what's in", "read"]):
            return analyze_selected_paths(images, raw)
        if any(x in c for x in ["merge", "combine", "collage"]):
            return selected_merge_images(images, raw)
        if "crop" in c or any(x in c for x in ["square", "16:9", "9:16", "4:5", "reel", "short", "portrait crop", "landscape crop"]):
            return selected_crop_images(images, raw)
        return selected_edit_images(images, raw or "enhance realistic 4k")

    if videos:
        if any(x in c for x in ["analyze", "analyse", "describe", "what is in", "what's in", "read", "info"]):
            return analyze_selected_paths(videos, raw)
        if any(x in c for x in ["merge", "combine", "join"]):
            return selected_merge_videos(videos, raw)
        if "crop" in c or any(x in c for x in ["square", "16:9", "9:16", "4:5", "reel", "short"]):
            return selected_crop_videos(videos, raw)
        return selected_edit_videos(videos, raw or "enhance")

    return False, "I could not understand this selected media task."


def save_uploaded_media(filename, data_url):
    if "," not in str(data_url):
        return False, "Upload failed: file data was not received.", None
    header, data = str(data_url).split(",", 1)
    try:
        raw = base64.b64decode(data)
    except Exception:
        return False, "Upload failed: invalid file data.", None
    max_bytes = int(settings.get("max_media_upload_mb", 80)) * 1024 * 1024
    if len(raw) > max_bytes:
        return False, f"Upload is too large. Limit is {settings.get('max_media_upload_mb', 80)} MB.", None
    ext = Path(str(filename or "")).suffix.lower()
    if not ext:
        if "image/png" in header:
            ext = ".png"
        elif "image/jpeg" in header:
            ext = ".jpg"
        elif "video/mp4" in header:
            ext = ".mp4"
        else:
            ext = ".bin"
    if ext not in GENERAL_FILE_EXTS:
        return False, "Unsupported file type. Use image, video, text/code, PDF, Word, Excel, PowerPoint, JSON, or CSV.", None
    if ext in IMAGE_EXTS or ext in VIDEO_EXTS:
        out = MEDIA_UPLOAD_DIR / safe_media_filename(filename, ext)
        studio = "Media Studio"
    else:
        out = FILE_UPLOAD_DIR / safe_upload_filename(filename, ext)
        studio = "File Studio"
    out.write_bytes(raw)
    return True, f"Uploaded {out.name} to {studio}.", out


def looks_like_media_request(c):
    return any(x in c for x in ["photo", "image", "picture", "pic", "video", "media", "wallpaper", "thumbnail", "poster", "crop", "merge photos", "merge images", "merge videos"]) and any(
        x in c for x in ["analyze", "analyse", "describe", "what is", "make", "create", "generate", "edit", "crop", "merge", "resize", "compress", "enhance", "blur", "sharpen", "grayscale", "gray", "sepia", "vignette"]
    )


def media_command(raw, c):
    if c in ["media help", "photo help", "video help", "image help", "media studio"]:
        return True, media_help()
    if "status" in c and any(x in c for x in ["media", "photo", "video", "image"]):
        status = media_tool_status()
        return True, (
            "Media status\n"
            f"- Pillow image tools: {'OK' if status['pillow'] else 'missing'}\n"
            f"- FFmpeg video tools: {'OK' if status['ffmpeg'] else 'missing'}\n"
            f"- FFprobe video analysis: {'OK' if status.get('ffprobe') else 'missing'}\n"
            f"- Gemini image vision: {'OK' if status['gemini_vision'] else 'not connected'}\n"
            f"- FFmpeg path: {status.get('ffmpeg_path') or 'not found'}\n"
            f"- Folder: {status['media_folder']}"
        )
    if any(x in c for x in ["analyze", "analyse", "describe", "what is in", "what's in"]):
        return media_analyze_command(raw, c)
    if any(x in c for x in ["merge videos", "merge video"]) or ("merge" in c and "video" in c):
        return merge_videos_command(raw, c)
    if any(x in c for x in ["merge photos", "merge images", "merge pictures"]) or ("merge" in c and any(x in c for x in ["photo", "image", "picture"])):
        return merge_images_command(raw, c)
    if "crop" in c and "video" in c:
        return crop_video_command(raw, c)
    if "crop" in c:
        return crop_image_command(raw, c)
    if any(x in c for x in ["edit", "enhance", "blur", "sharpen", "grayscale", "gray", "black and white", "rotate", "mirror", "brighten", "contrast", "sepia", "vignette", "resize", "compress"]) and any(x in c for x in ["photo", "image", "picture", "pic"]):
        return edit_image_command(raw, c)
    if any(x in c for x in ["make video", "create video", "generate video"]) or (any(x in c for x in ["make", "create", "generate"]) and "video" in c):
        return create_prompt_video(raw, c)
    if any(x in c for x in ["make photo", "create photo", "generate photo", "make image", "create image", "generate image", "make picture", "create picture"]) or (any(x in c for x in ["make", "create", "generate"]) and any(x in c for x in ["photo", "image", "picture", "wallpaper", "thumbnail", "poster"])):
        return create_prompt_image(raw, c)
    return False, ""


def volume(c):
    if not pyautogui:
        return False, "PyAutoGUI is not installed."
    if "mute" in c:
        pyautogui.press("volumemute")
        return True, "Volume muted or unmuted."
    key = "volumeup" if any(x in c for x in ["up", "increase", "badhao"]) else "volumedown"
    for _ in range(5):
        pyautogui.press(key)
    return True, "Volume changed."


def open_folder(c):
    c = normalize_words(c)
    name = clean_entity_name(remove_words(c, ["open", "show", "go", "to", "folder", "files", "file", "my", "the"]))
    folders = common_folders()
    aliases = {
        "desktop": ["desktop", "desk top"],
        "documents": ["documents", "document"],
        "downloads": ["downloads", "download"],
        "pictures": ["pictures", "photos", "images"],
        "videos": ["videos", "video"],
        "music": ["music", "songs"],
        "nova_files": ["nova", "nova files", "created files"],
        "onedrive": ["onedrive", "one drive"],
    }
    for key, words in aliases.items():
        if name == key or any(w in c for w in words):
            path = folders.get(key)
            if path and Path(path).exists():
                os.startfile(path)
                return True, f"Opening {key.replace('_', ' ')} folder."
    if c in ["open c drive", "c drive", "open drive c"]:
        os.startfile("C:\\")
        return True, "Opening C drive."
    return False, ""


def keyboard_mouse_control(c):
    if not pyautogui:
        return False, "Keyboard/mouse automation is not installed."
    c = normalize_words(c)
    hotkeys = {
        "copy": ("ctrl", "c"), "paste": ("ctrl", "v"), "cut": ("ctrl", "x"),
        "select all": ("ctrl", "a"), "save": ("ctrl", "s"), "undo": ("ctrl", "z"),
        "redo": ("ctrl", "y"), "new tab": ("ctrl", "t"), "reopen tab": ("ctrl", "shift", "t"),
        "close tab": ("ctrl", "w"), "refresh": ("ctrl", "r"), "reload": ("ctrl", "r"),
        "find": ("ctrl", "f"), "print": ("ctrl", "p"), "task manager": ("ctrl", "shift", "esc"),
        "lock laptop": ("win", "l"), "lock screen": ("win", "l"),
    }
    for phrase, keys in hotkeys.items():
        if c == phrase or c.startswith(phrase + " ") or f" {phrase}" in c:
            pyautogui.hotkey(*keys)
            return True, f"Sent {phrase} command."
    press_map = {
        "enter": "enter", "press enter": "enter", "tab": "tab", "press tab": "tab",
        "escape": "esc", "esc": "esc", "space": "space", "backspace": "backspace",
        "delete": "delete", "up": "up", "down": "down", "left": "left", "right": "right",
        "home": "home", "end": "end",
    }
    for phrase, key in press_map.items():
        if c == phrase or c == "press " + phrase or c.endswith(" " + phrase):
            pyautogui.press(key)
            return True, f"Pressed {key}."
    if "double click" in c:
        pyautogui.doubleClick()
        return True, "Double clicked."
    if "right click" in c:
        pyautogui.rightClick()
        return True, "Right clicked."
    if c == "click" or "left click" in c:
        pyautogui.click()
        return True, "Clicked."
    if "move mouse" in c or "mouse move" in c:
        amount_match = re.search(r"\b(\d{1,4})\b", c)
        amount = int(amount_match.group(1)) if amount_match else 120
        dx = dy = 0
        if "left" in c:
            dx = -amount
        elif "right" in c:
            dx = amount
        elif "up" in c:
            dy = -amount
        elif "down" in c:
            dy = amount
        if dx or dy:
            pyautogui.moveRel(dx, dy, duration=0.12)
            return True, "Mouse moved."
        return False, "Tell direction: move mouse left/right/up/down."
    if "drag mouse" in c or c.startswith("drag "):
        amount_match = re.search(r"\b(\d{1,4})\b", c)
        amount = int(amount_match.group(1)) if amount_match else 160
        dx = amount if "right" in c else -amount if "left" in c else 0
        dy = amount if "down" in c else -amount if "up" in c else 0
        if dx or dy:
            pyautogui.dragRel(dx, dy, duration=0.25, button="left")
            return True, "Mouse dragged."
        return False, "Tell direction: drag left/right/up/down."
    if "scroll down" in c:
        pyautogui.scroll(-6)
        return True, "Scrolled down."
    if "scroll up" in c:
        pyautogui.scroll(6)
        return True, "Scrolled up."
    if "go back" in c or c == "back":
        pyautogui.hotkey("alt", "left")
        return True, "Went back."
    if "go forward" in c or c == "forward":
        pyautogui.hotkey("alt", "right")
        return True, "Went forward."
    return False, ""


def window_control(c):
    if not pyautogui:
        return False, "PyAutoGUI is not installed."
    if "minimize" in c:
        pyautogui.hotkey("win", "down")
        return True, "Window minimized."
    if "maximize" in c:
        pyautogui.hotkey("win", "up")
        return True, "Window maximized."
    if "switch" in c or "next window" in c:
        pyautogui.hotkey("alt", "tab")
        return True, "Switched window."
    if "close window" in c:
        pyautogui.hotkey("alt", "f4")
        return True, "Window close command sent."
    return False, ""


def browser_page_control(c):
    if not pyautogui:
        return False, "PyAutoGUI is not installed."
    c = normalize_words(c)
    if any(x in c for x in ["close all slide", "close all slides", "close all page", "close all pages", "close all tab", "close all tabs"]):
        pyautogui.hotkey("ctrl", "shift", "w")
        return True, "Closed all pages or tabs in the current browser window."
    positions = {
        "first": "1", "1st": "1", "one": "1",
        "second": "2", "2nd": "2", "two": "2",
        "third": "3", "thride": "3", "3rd": "3", "three": "3",
        "fourth": "4", "4th": "4", "four": "4",
        "fifth": "5", "5th": "5", "five": "5",
        "sixth": "6", "6th": "6", "six": "6",
        "seventh": "7", "7th": "7", "seven": "7",
        "eighth": "8", "8th": "8", "eight": "8",
    }
    if any(x in c for x in ["last", "front", "current", "this"]):
        if "last" in c:
            pyautogui.hotkey("ctrl", "9")
            time.sleep(0.1)
        pyautogui.hotkey("ctrl", "w")
        return True, "Closed the requested current or last page."
    for word, number in positions.items():
        if word in c:
            pyautogui.hotkey("ctrl", number)
            time.sleep(0.1)
            pyautogui.hotkey("ctrl", "w")
            return True, f"Closed page number {number}."
    if any(x in c for x in ["close slide", "close page", "close tab"]):
        pyautogui.hotkey("ctrl", "w")
        return True, "Closed the current page or tab."
    return False, ""



def open_whatsapp_app():
    # Strict local-app rule: never open browser/Google for WhatsApp unless user explicitly asks.
    if focus_window("whatsapp"):
        return True
    shortcut = find_start_menu_app("whatsapp")
    if shortcut:
        open_shortcut(shortcut)
        time.sleep(3)
        if focus_window("whatsapp"):
            return True
    exe = find_known_executable("whatsapp")
    if exe:
        try:
            subprocess.Popen([str(exe)])
            time.sleep(3)
            if focus_window("whatsapp"):
                return True
        except Exception as e:
            log(f"WhatsApp exe open error: {e}")
    launchers = [
        ["cmd", "/c", "start", "", "whatsapp:"],
        ["explorer.exe", "shell:AppsFolder\\5319275A.WhatsAppDesktop_cv1g1gvanyjgm!App"],
        ["explorer.exe", "shell:AppsFolder\\5319275A.WhatsAppBeta_cv1g1gvanyjgm!App"],
    ]
    launched_any = False
    for launcher in launchers:
        try:
            subprocess.Popen(launcher, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
            launched_any = True
            time.sleep(3)
            if focus_window("whatsapp"):
                return True
        except Exception as e:
            log(f"WhatsApp launcher error: {launcher}: {e}")
    # Some Windows builds launch WhatsApp but do not expose the window title to automation.
    # Treat a local app launch as success so Nova never opens Google for WhatsApp.
    return launched_any


def search_whatsapp(c):
    if not settings.get("whatsapp_control_enabled", True):
        return True, "WhatsApp Control is OFF in Nova Settings. Turn it ON and allow the popup before I control WhatsApp."
    c = normalize_words(c)
    if "browser" in c or "web" in c:
        return False, "WhatsApp browser/web command disabled by default. Say clearly: open WhatsApp in browser, if you want browser."
    name = resolve_contact_name(remove_words(c, ["search", "find", "whatsapp", "in", "on", "contact", "chat"]))
    app_ok = open_whatsapp_app()
    if not app_ok:
        return True, "WhatsApp Desktop was not found or is not logged in. I will not open Google for this. Install or log in to WhatsApp Desktop first."
    if not pyautogui:
        return True, "WhatsApp is open, but keyboard automation is not available. Type the contact name manually in WhatsApp search."
    if name:
        def do_search():
            try:
                time.sleep(2)
                focus_window("whatsapp")
                # WhatsApp Desktop search shortcuts vary; try both safe shortcuts.
                pyautogui.hotkey("ctrl", "f")
                time.sleep(0.2)
                pyautogui.hotkey("ctrl", "k")
                time.sleep(0.2)
                paste_text(name)
            except Exception as e:
                log(f"WhatsApp search automation stopped: {e}")
                speak("WhatsApp search automation stopped. Move the mouse away from the screen corner and try again.")
        threading.Thread(target=do_search, daemon=True).start()
        return True, f"WhatsApp is open. I am trying to paste {name} into the search box."
    return True, "WhatsApp is open. Say the contact name clearly, for example: search Dishant in WhatsApp."


def whatsapp_call_contact(c):
    if not settings.get("whatsapp_control_enabled", True):
        return True, "WhatsApp Control is OFF in Nova Settings. Turn it ON and allow the popup before I control WhatsApp."
    c = normalize_words(c)
    name = resolve_contact_name(remove_words(c, ["call", "voice", "video", "whatsapp", "on", "ko", "par", "karo", "lagao"]))
    app_ok = open_whatsapp_app()
    if not app_ok:
        return True, "WhatsApp Desktop was not found or is not logged in. I will not open Google for this. Install or log in to WhatsApp Desktop first."
    if not pyautogui:
        return True, "WhatsApp is open, but keyboard automation is not available. Open the contact and click the call button manually."
    if name:
        def do_call():
            try:
                time.sleep(2.2)
                focus_window("whatsapp")
                pyautogui.hotkey("ctrl", "f")
                time.sleep(0.25)
                pyautogui.hotkey("ctrl", "k")
                time.sleep(0.25)
                paste_text(name)
                time.sleep(0.7)
                pyautogui.press("enter")
                time.sleep(1.1)
                # Try common WhatsApp Desktop call shortcuts/navigation. If layout differs,
                # WhatsApp will stay open with the contact selected.
                if "video" in c:
                    pyautogui.hotkey("ctrl", "shift", "v")
                else:
                    pyautogui.hotkey("ctrl", "shift", "p")
                time.sleep(0.7)
                pyautogui.press("enter")
                time.sleep(0.8)
                click_whatsapp_call_button(video=("video" in c))
            except Exception as e:
                log(f"WhatsApp call automation stopped: {e}")
                speak("WhatsApp call automation stopped. Move the mouse away from the screen corner and try again.")
        threading.Thread(target=do_call, daemon=True).start()
        return True, f"I am trying to open the {name} chat in WhatsApp and click the call button. This is best-effort automation; if the call does not start, WhatsApp layout, login, contact selection, or permissions blocked it."
    return True, "WhatsApp is open. Say the contact name clearly, for example: call Dishant on WhatsApp."


def whatsapp_command(c):
    if not settings.get("whatsapp_control_enabled", True):
        return True, "WhatsApp Control is OFF in Nova Settings. Turn it ON and allow the popup before I control WhatsApp."
    c = normalize_words(c)
    if "whatsapp" not in c and not any(x in c for x in ["message", "call"]):
        return False, ""
    if "browser" in c or "web" in c:
        webbrowser.open("https://web.whatsapp.com")
        return True, "Opening WhatsApp in browser because you clearly asked for browser/web."
    if "call" in c:
        return whatsapp_call_contact(c)
    app_ok = open_whatsapp_app()
    if "message" in c or "send" in c:
        return True, "WhatsApp local app launch command was sent. Say the contact and message clearly, for example: message Ritik hello in WhatsApp." if app_ok else "WhatsApp Desktop was not found. Install or log in to WhatsApp Desktop first."
    return True, "WhatsApp local app launch command was sent. I will not open Google." if app_ok else "WhatsApp Desktop was not found. I will not open Google. Install or log in to WhatsApp Desktop first."


def ps_literal(value) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def is_admin_user():
    if platform.system().lower() != "windows":
        return False
    try:
        return bool(ctypes.windll.shell32.IsUserAnAdmin())
    except Exception:
        return False


def run_elevated_powershell(script_body: str, timeout=90):
    """Run a short admin PowerShell task through a real Windows UAC popup.
    The task writes a result file so Nova can verify before replying.
    """
    if platform.system().lower() != "windows":
        return False, "This action is Windows only."
    action_dir = BASE_DIR / ".nova_admin_actions"
    action_dir.mkdir(exist_ok=True)
    token = secrets.token_hex(8)
    ps1 = action_dir / f"nova_admin_{token}.ps1"
    result_file = action_dir / f"nova_admin_{token}.txt"
    wrapper = f"""
$ErrorActionPreference = 'Stop'
$NovaOk = $false
$NovaMessage = ''
try {{
{script_body}
    if ([string]::IsNullOrWhiteSpace($NovaMessage)) {{ $NovaMessage = 'Action completed.' }}
    if ($NovaOk) {{
        Set-Content -Path {ps_literal(result_file)} -Value ('OK|' + $NovaMessage) -Encoding UTF8
        exit 0
    }} else {{
        Set-Content -Path {ps_literal(result_file)} -Value ('FAIL|' + $NovaMessage) -Encoding UTF8
        exit 2
    }}
}} catch {{
    Set-Content -Path {ps_literal(result_file)} -Value ('ERROR|' + $_.Exception.Message) -Encoding UTF8
    exit 1
}}
"""
    ps1.write_text(wrapper, encoding="utf-8")
    if is_admin_user():
        try:
            subprocess.Popen(
                [_powershell_exe(), "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(ps1)],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
            )
        except Exception as e:
            return False, f"Admin action start failed: {e}"
    else:
        try:
            params = f'-NoProfile -ExecutionPolicy Bypass -File "{ps1}"'
            rc = ctypes.windll.shell32.ShellExecuteW(None, "runas", _powershell_exe(), params, str(BASE_DIR), 1)
            if int(rc) <= 32:
                return False, "Windows permission popup was cancelled or blocked."
        except Exception as e:
            return False, f"Windows permission popup failed: {e}"

    deadline = time.time() + timeout
    while time.time() < deadline:
        if result_file.exists():
            raw = result_file.read_text(encoding="utf-8", errors="replace").strip()
            if "|" in raw:
                status, message = raw.split("|", 1)
            else:
                status, message = "FAIL", raw
            ok = status == "OK"
            return ok, message.strip() or ("Action verified." if ok else "Action failed.")
        time.sleep(0.4)
    return False, "Windows permission popup was not approved, or the admin action timed out."


def wifi_adapter_status():
    """Return Wi-Fi adapter state from netsh without requiring admin."""
    if platform.system().lower() != "windows":
        return None
    try:
        r = subprocess.run(["netsh", "interface", "show", "interface"], capture_output=True, text=True, timeout=8)
        output = (r.stdout or "") + "\n" + (r.stderr or "")
    except Exception as e:
        log(f"Wi-Fi status read failed: {e}")
        return None
    candidates = []
    for line in output.splitlines():
        parts = re.split(r"\s{2,}", line.strip())
        if len(parts) >= 4 and parts[0].lower() in ["enabled", "disabled"]:
            item = {
                "admin": parts[0],
                "state": parts[1],
                "type": parts[2],
                "name": " ".join(parts[3:]),
            }
            name_l = item["name"].lower()
            if any(x in name_l for x in ["wi-fi", "wifi", "wireless", "wlan"]):
                candidates.append(item)
    return candidates[0] if candidates else None


def set_wifi(enable=True):
    status = wifi_adapter_status()
    target = "Enabled" if enable else "Disabled"
    label = "ON" if enable else "OFF"
    if not status:
        open_uri("ms-settings:network-wifi")
        return True, "I could not identify the Wi-Fi adapter automatically, so I opened Wi-Fi settings. I did not claim Wi-Fi changed."
    name = status["name"]
    if status["admin"].lower() == target.lower():
        return True, f"Wi-Fi is already {label}. Verified adapter '{name}' admin state is {status['admin']}."

    admin_value = "enabled" if enable else "disabled"
    expected = "Enabled" if enable else "Disabled"
    script = f"""
$name = {ps_literal(name)}
$target = {ps_literal(admin_value)}
$expected = {ps_literal(expected)}
& netsh interface set interface name="$name" admin=$target | Out-Null
Start-Sleep -Seconds 2
$rows = & netsh interface show interface
$pattern = '\\s' + [regex]::Escape($name) + '$'
$line = $rows | Where-Object {{ $_ -match $pattern }} | Select-Object -First 1
if ($line -match ('^\\s*' + [regex]::Escape($expected) + '\\s+')) {{
    $NovaOk = $true
    $NovaMessage = "Verified Wi-Fi {label}. Windows reports adapter '$name' admin state $expected."
}} else {{
    $NovaOk = $false
    $NovaMessage = "I tried to turn Wi-Fi {label}, but Windows still reports: $line"
}}
"""
    ok, msg = run_elevated_powershell(script, timeout=75)
    if ok:
        return True, msg
    open_uri("ms-settings:network-wifi")
    return True, msg + " Wi-Fi settings are open so you can check it manually."



def powershell_json(command: str, timeout=12):
    if platform.system().lower() != "windows":
        return None, "Windows only."
    try:
        result = subprocess.run(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", command], capture_output=True, text=True, timeout=timeout)
        if result.returncode != 0:
            return None, (result.stderr or result.stdout or "PowerShell command failed.").strip()
        out = (result.stdout or "").strip()
        if not out:
            return None, ""
        try:
            return json.loads(out), ""
        except Exception:
            return out, ""
    except Exception as e:
        return None, str(e)


def set_bluetooth(enable=True):
    """Enable/disable the Bluetooth radio with admin permission, then verify.
    Windows has no simple public Bluetooth toggle API; disabling the radio device is
    the most honest local-control method and requires the user to allow UAC.
    """
    state = "ON" if enable else "OFF"
    action = "Enable-PnpDevice" if enable else "Disable-PnpDevice"
    expected_ok = "$true" if enable else "$false"
    script = f"""
$devices = @(Get-PnpDevice -Class Bluetooth -ErrorAction SilentlyContinue | Where-Object {{
    $_.InstanceId -notlike 'BTHENUM*' -and
    $_.FriendlyName -notmatch 'Enumerator|Protocol|Service|LE Device|Hands-Free|Audio|Remote Control|Personal Area|RFCOMM' -and
    ($_.FriendlyName -match 'Bluetooth|Wireless|Intel|Realtek|MediaTek|Qualcomm|RZ')
}})
if (-not $devices -or $devices.Count -eq 0) {{
    Start-Process 'ms-settings:bluetooth'
    $NovaOk = $false
    $NovaMessage = 'I could not find the Bluetooth radio device automatically. Bluetooth settings are open. I did not claim Bluetooth changed.'
}} else {{
    foreach ($d in $devices) {{
        {action} -InstanceId $d.InstanceId -Confirm:$false -ErrorAction Stop
    }}
    Start-Sleep -Seconds 3
    $after = @()
    foreach ($d in $devices) {{
        $one = Get-PnpDevice -InstanceId $d.InstanceId -ErrorAction SilentlyContinue
        if ($one) {{ $after += $one }}
    }}
    $names = ($devices | ForEach-Object {{ $_.FriendlyName }}) -join ', '
    if ({expected_ok}) {{
        $bad = @($after | Where-Object {{ $_.Status -ne 'OK' }})
        $NovaOk = ($bad.Count -eq 0)
    }} else {{
        $stillOn = @($after | Where-Object {{ $_.Status -eq 'OK' }})
        $NovaOk = ($stillOn.Count -eq 0)
    }}
    if ($NovaOk) {{
        $NovaMessage = "Verified Bluetooth {state}. Radio device changed: $names"
    }} else {{
        $current = ($after | ForEach-Object {{ $_.FriendlyName + '=' + $_.Status }}) -join ', '
        $NovaMessage = "I tried to turn Bluetooth {state}, but Windows still reports: $current"
        Start-Process 'ms-settings:bluetooth'
    }}
}}
"""
    ok, msg = run_elevated_powershell(script, timeout=90)
    if ok:
        return True, msg
    open_uri("ms-settings:bluetooth")
    return True, msg + " Bluetooth settings are open so you can check it manually."


def firewall_status_lines():
    lines = []
    if platform.system().lower() != "windows":
        return lines
    data, err = powershell_json("Get-NetFirewallProfile | Select-Object Name,Enabled | ConvertTo-Json -Compress", timeout=8)
    if data:
        if isinstance(data, dict):
            data = [data]
        for item in data:
            lines.append(f"- Firewall {item.get('Name','profile')}: {'ON' if item.get('Enabled') else 'OFF'}")
    elif err:
        lines.append("- Firewall status: open Windows Security to check.")
    return lines


def run_defender_quick_scan():
    if platform.system().lower() != "windows":
        return False, "Defender scan is Windows only."
    try:
        subprocess.Popen(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", "Start-MpScan -ScanType QuickScan"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        run_start("windowsdefender:")
        return True, "Windows Defender quick scan started. I also opened Windows Security."
    except Exception as e:
        run_start("windowsdefender:")
        return True, f"Windows Security opened. Automatic quick scan did not start: {e}"


def update_defender_signatures():
    if platform.system().lower() != "windows":
        return False, "Defender update is Windows only."
    try:
        subprocess.Popen(["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", "Update-MpSignature"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return True, "Windows Defender security intelligence update started."
    except Exception as e:
        return False, f"Defender update did not start: {e}"


def enable_firewall_profiles():
    if platform.system().lower() != "windows":
        return False, "Firewall control is Windows only."
    script = """
Set-NetFirewallProfile -Profile Domain,Public,Private -Enabled True
Start-Sleep -Milliseconds 700
$profiles = @(Get-NetFirewallProfile | Select-Object Name,Enabled)
$off = @($profiles | Where-Object { -not $_.Enabled })
if ($off.Count -eq 0) {
    $NovaOk = $true
    $NovaMessage = 'Verified: Windows Firewall is ON for Domain, Private, and Public profiles.'
} else {
    $NovaOk = $false
    $NovaMessage = 'Firewall command ran, but these profiles still look OFF: ' + (($off | ForEach-Object { $_.Name }) -join ', ')
}
"""
    ok, msg = run_elevated_powershell(script, timeout=70)
    if ok:
        return True, msg
    run_start("windowsdefender:")
    return True, msg + " Windows Security opened."

def security_status():
    lines = [
        "Security Center",
        "- Local-only mode: " + ("ON" if settings.get("local_only_mode", True) else "OFF"),
        "- Assistant switch: " + ("ON" if settings.get("assistant_enabled", True) else "OFF"),
        "- Local PIN lock: " + ("ON" if access_lock_enabled() else "not set"),
        "- Dangerous action confirmation: " + ("ON" if settings.get("confirm_dangerous_actions", True) else "OFF"),
        "- Laptop profile: " + ("connected" if (settings.get("laptop_connected") or LAPTOP_PROFILE_FILE.exists()) else "not connected"),
        "- Network check: " + network_status().get("message", "unknown"),
    ]
    if platform.system().lower() == "windows":
        try:
            ps = (
                "$s=Get-MpComputerStatus; "
                "[pscustomobject]@{"
                "AntivirusEnabled=$s.AntivirusEnabled;"
                "RealTimeProtectionEnabled=$s.RealTimeProtectionEnabled;"
                "Firewall='Open Windows Security for firewall status';"
                "QuickScanAge=$s.QuickScanAge;"
                "SignatureAge=$s.AntivirusSignatureAge"
                "} | ConvertTo-Json -Compress"
            )
            result = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=8)
            if result.returncode == 0 and result.stdout.strip():
                data = json.loads(result.stdout)
                lines.extend([
                    "- Defender antivirus: " + ("ON" if data.get("AntivirusEnabled") else "check Windows Security"),
                    "- Real-time protection: " + ("ON" if data.get("RealTimeProtectionEnabled") else "needs attention"),
                    f"- Defender signature age: {data.get('SignatureAge', 'unknown')} days",
                    f"- Quick scan age: {data.get('QuickScanAge', 'unknown')} days",
                ])
            else:
                lines.append("- Defender status: open Windows Security to check.")
        except Exception as e:
            lines.append(f"- Defender status: could not read automatically ({e}).")
    lines.extend(firewall_status_lines())
    lines.append("- Tool web server: bound to 127.0.0.1 only, so other devices cannot connect by default.")
    lines.append("- VPN note: I can open VPN settings, but I will not add stealth/location-hiding automation.")
    return "\n".join(lines)


def defender_realtime_on():
    if platform.system().lower() != "windows":
        return False, "Defender control is Windows only."
    script = """
Set-MpPreference -DisableRealtimeMonitoring $false
Set-MpPreference -PUAProtection Enabled
Start-Sleep -Seconds 1
$s = Get-MpComputerStatus
if ($s.RealTimeProtectionEnabled) {
    $NovaOk = $true
    $NovaMessage = 'Verified: Microsoft Defender real-time protection is ON and PUA protection is enabled.'
} else {
    $NovaOk = $false
    $NovaMessage = 'Defender command ran, but real-time protection still looks OFF. Open Windows Security and check Virus & threat protection.'
}
"""
    ok, msg = run_elevated_powershell(script, timeout=70)
    if ok:
        return True, msg
    run_start("windowsdefender:")
    return True, msg + " Windows Security opened."


def nova_shield_help():
    return (
        "Nova Shield defensive commands:\n"
        "- nova shield status\n"
        "- enable nova shield\n"
        "- security repair\n"
        "- quick virus scan\n"
        "- update defender\n"
        "- enable firewall\n"
        "- dark web safety\n\n"
        "Nova Shield uses Windows Defender, Windows Firewall, local-only mode, dangerous-action confirmation, and honest status checks. It is defensive only."
    )


def nova_shield_status():
    lines = [
        "Nova Shield",
        "- Mode: defensive only",
        "- Local-only server: " + ("ON" if settings.get("local_only_mode", True) else "OFF"),
        "- Control permission profile: " + ("saved" if settings.get("local_control_permission", False) else "not saved"),
        "- Dangerous-action confirmation: " + ("ON" if settings.get("confirm_dangerous_actions", True) else "OFF"),
        "- PIN lock: " + ("ON" if access_lock_enabled() else "not set"),
        "",
        security_status(),
    ]
    return "\n".join(lines)


def enable_nova_shield():
    settings["local_only_mode"] = True
    settings["confirm_dangerous_actions"] = True
    settings["control_mode"] = True
    settings["local_control_permission"] = True
    settings["nova_shield_enabled"] = True
    settings["nova_shield_updated_at"] = datetime.datetime.now().isoformat()
    save_json(SETTINGS_FILE, settings)
    script = """
Set-NetFirewallProfile -Profile Domain,Public,Private -Enabled True
Set-MpPreference -DisableRealtimeMonitoring $false
Set-MpPreference -PUAProtection Enabled
Update-MpSignature -ErrorAction SilentlyContinue
Start-MpScan -ScanType QuickScan -ErrorAction SilentlyContinue
Start-Sleep -Seconds 1
$fw = @(Get-NetFirewallProfile | Select-Object Name,Enabled)
$fwOff = @($fw | Where-Object { -not $_.Enabled })
$mp = Get-MpComputerStatus
$checks = @()
if ($fwOff.Count -eq 0) { $checks += 'Firewall verified ON for all profiles' } else { $checks += ('Firewall still needs attention: ' + (($fwOff | ForEach-Object { $_.Name }) -join ', ')) }
if ($mp.RealTimeProtectionEnabled) { $checks += 'Defender real-time protection verified ON' } else { $checks += 'Defender real-time protection still needs attention' }
$checks += 'Defender signature update requested'
$checks += 'Quick scan requested'
$NovaOk = ($fwOff.Count -eq 0 -and $mp.RealTimeProtectionEnabled)
$NovaMessage = ($checks -join '; ')
"""
    ok_admin, msg_admin = run_elevated_powershell(script, timeout=180)
    if not ok_admin:
        run_start("windowsdefender:")
        msg_admin = msg_admin + " Windows Security opened."
    status_file = BASE_DIR / "nova_shield_status.json"
    status_file.write_text(json.dumps({
        "enabled_at": datetime.datetime.now().isoformat(),
        "local_only_mode": True,
        "dangerous_action_confirmation": True,
        "admin_action_ok": ok_admin,
        "admin_action": msg_admin,
    }, indent=2), encoding="utf-8")
    return True, (
        "Nova Shield enabled defensively.\n"
        "- " + "\n- ".join([
            msg_admin,
            "Local-only mode ON",
            "Dangerous action confirmation ON",
            f"Shield status saved: {status_file.name}",
        ])
    )


def dark_web_safety_info():
    if not settings.get("dark_web_safety_enabled", True):
        return "Dark Web Safety is OFF in Nova Settings. Turn it ON and allow the popup before I discuss defensive dark-web safety."
    return (
        "Dark web safety overview:\n"
        "- I can explain risks, privacy basics, scam signs, malware risks, phishing, credential leaks, and how to secure your laptop.\n"
        "- I will not access hidden services, illegal marketplaces, stolen-data sources, or give buying/selling, hacking, evasion, or credential-access instructions.\n"
        "- Your permission does not change that safety rule; Nova Shield stays defensive only.\n"
        "- Best protection: keep Defender and Firewall ON, update Windows, use strong unique passwords, turn on 2FA, avoid unknown downloads, and scan suspicious files before opening.\n"
        "- If you think your data leaked, change passwords from a clean browser session, revoke unknown logins, enable 2FA, and run: enable nova shield."
    )



# -------------------- Casting / Wireless Display / Settings Control --------------------
def open_windows_cast_panel(device_name: str = ""):
    """Open Windows Cast/Wireless Display panel safely.
    This does not bypass TV pairing prompts. Windows/Android TV must be on the same Wi-Fi
    and the TV must support Miracast/Chromecast or Wireless Display.
    """
    if platform.system().lower() != "windows":
        return False, "Cast control is Windows only."
    try:
        # Win+K is the official Windows Cast quick panel shortcut.
        if pyautogui:
            pyautogui.hotkey("win", "k")
            time.sleep(1.0)
            if device_name:
                # Best-effort only: some Windows builds accept typing/search in this panel.
                paste_text(device_name)
                time.sleep(0.3)
                pyautogui.press("enter")
        else:
            run_start("ms-settings-connectabledevices:devicediscovery")
        msg = "Cast panel opened. Make sure the Android TV is on the same Wi-Fi and Cast/Wireless Display is enabled on the TV."
        if device_name:
            msg += f" I also made a best-effort attempt to select device '{device_name}'."
        msg += " If the TV shows an Allow/Pair prompt, approve it with the TV remote."
        return True, msg
    except Exception as e:
        try:
            run_start("ms-settings-connectabledevices:devicediscovery")
        except Exception:
            pass
        return True, f"Cast settings opened. Win+K automation failed: {e}"


def project_screen_mode(mode: str):
    """Use Windows DisplaySwitch safely: internal/clone/extend/external."""
    if platform.system().lower() != "windows":
        return False, "Projection control is Windows only."
    mode_map = {
        "pc": "/internal",
        "internal": "/internal",
        "stop": "/internal",
        "disconnect": "/internal",
        "duplicate": "/clone",
        "mirror": "/clone",
        "clone": "/clone",
        "extend": "/extend",
        "external": "/external",
        "second": "/external",
    }
    arg = mode_map.get(mode, "/clone")
    try:
        subprocess.Popen(["DisplaySwitch.exe", arg], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        if arg == "/internal":
            return True, "Casting/projection stop command sent. Screen mode is set to PC screen only."
        if arg == "/clone":
            return True, "Duplicate/Mirror screen mode was set. Use Win+K to select the Android TV if needed."
        if arg == "/extend":
            return True, "Extend screen mode was set."
        return True, "Second screen only mode was set."
    except Exception as e:
        return False, f"Projection command failed: {e}"


def cast_command(c):
    c = normalize_words(c)
    # Stop/disconnect must be checked first.
    if any(x in c for x in ["stop casting", "disconnect cast", "stop mirror", "stop mirroring", "pc screen only", "laptop screen only"]):
        return project_screen_mode("stop")
    if any(x in c for x in ["duplicate screen", "mirror screen", "screen mirror", "mirror laptop", "duplicate display"]):
        ok, msg = project_screen_mode("duplicate")
        ok2, msg2 = open_windows_cast_panel("")
        return True, msg + "\n" + msg2
    if any(x in c for x in ["extend screen", "extend display"]):
        return project_screen_mode("extend")
    if any(x in c for x in ["second screen only", "tv screen only"]):
        return project_screen_mode("external")
    if any(x in c for x in ["cast", "android tv", "wireless display", "connect display", "connect to tv", "screen cast", "screen casting", "project to tv", "project screen"]):
        name = c
        for w in ["cast", "my", "laptop", "to", "android", "tv", "connect", "wireless", "display", "screen", "project", "mirror", "mirroring", "on", "with"]:
            name = re.sub(r"\b" + re.escape(w) + r"\b", " ", name)
        name = re.sub(r"\s+", " ", name).strip()
        return open_windows_cast_panel(name)
    return False, ""


WINDOWS_SETTINGS_COMMANDS = {
    "display": "ms-settings:display",
    "sound": "ms-settings:sound",
    "volume": "ms-settings:sound",
    "wifi": "ms-settings:network-wifi",
    "wi fi": "ms-settings:network-wifi",
    "network": "ms-settings:network",
    "bluetooth": "ms-settings:bluetooth",
    "cast": "ms-settings-connectabledevices:devicediscovery",
    "wireless display": "ms-settings-connectabledevices:devicediscovery",
    "printer": "ms-settings:printers",
    "mouse": "ms-settings:mousetouchpad",
    "keyboard": "ms-settings:keyboard",
    "touchpad": "ms-settings:devices-touchpad",
    "camera": "ms-settings:camera",
    "microphone": "ms-settings:privacy-microphone",
    "mic": "ms-settings:privacy-microphone",
    "privacy": "ms-settings:privacy",
    "apps": "ms-settings:appsfeatures",
    "installed apps": "ms-settings:appsfeatures",
    "default apps": "ms-settings:defaultapps",
    "startup apps": "ms-settings:startupapps",
    "storage": "ms-settings:storagesense",
    "battery": "ms-settings:batterysaver",
    "power": "ms-settings:powersleep",
    "notification": "ms-settings:notifications",
    "focus": "ms-settings:quiethours",
    "date": "ms-settings:dateandtime",
    "time": "ms-settings:dateandtime",
    "language": "ms-settings:regionlanguage",
    "region": "ms-settings:regionformatting",
    "account": "ms-settings:accounts",
    "sign in": "ms-settings:signinoptions",
    "personalization": "ms-settings:personalization",
    "theme": "ms-settings:themes",
    "taskbar": "ms-settings:taskbar",
    "accessibility": "ms-settings:easeofaccess",
    "windows update": "ms-settings:windowsupdate",
    "update": "ms-settings:windowsupdate",
    "security": "windowsdefender:",
    "firewall": "windowsdefender:",
    "vpn": "ms-settings:network-vpn",
    "proxy": "ms-settings:network-proxy",
    "troubleshoot": "ms-settings:troubleshoot",
    "recovery": "ms-settings:recovery",
    "activation": "ms-settings:activation",
    "about": "ms-settings:about",
}


def windows_settings_command(c):
    c = normalize_words(c)
    if c in ["open all settings", "all settings", "control all settings", "open setting control", "open settings control", "settings dashboard"]:
        run_start("ms-settings:")
        return True, "Windows Settings dashboard opened. You can say: open display settings, open sound settings, cast laptop to Android TV, or open Bluetooth settings."
    if "setting" not in c and "settings" not in c:
        return False, ""
    for key, uri in sorted(WINDOWS_SETTINGS_COMMANDS.items(), key=lambda x: len(x[0]), reverse=True):
        if key in c:
            run_start(uri)
            return True, f"Opening {key} settings."
    if c.startswith("open") or "control" in c:
        run_start("ms-settings:")
        return True, "Windows Settings opened."
    return False, ""


# -------------------- Safe Windows ON/OFF Setting Control --------------------
def ps_run(command: str, timeout=12):
    """Run PowerShell safely and return (ok, output)."""
    if platform.system().lower() != "windows":
        return False, "Windows only."
    try:
        r = subprocess.run(
            ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-Command", command],
            capture_output=True, text=True, timeout=timeout
        )
        out = ((r.stdout or "") + ("\n" + r.stderr if r.stderr else "")).strip()
        return r.returncode == 0, out
    except Exception as e:
        return False, str(e)


def open_uri(uri: str):
    try:
        run_start(uri)
        return True
    except Exception:
        try:
            os.startfile(uri)
            return True
        except Exception:
            return False


def set_volume_mute(enable=True):
    # enable=True means mute ON; enable=False means unmute.
    if not pyautogui:
        open_uri("ms-settings:sound")
        return True, "Sound settings opened. PyAutoGUI is missing, so direct mute/unmute could not run."
    try:
        # Windows media mute key toggles; we do best-effort. For exact state, use sound settings.
        pyautogui.press("volumemute")
        return True, "Mute toggle command sent. If the state is opposite, say mute or unmute again."
    except Exception as e:
        open_uri("ms-settings:sound")
        return True, f"Sound settings opened. Mute key failed: {e}"


def change_volume(direction="up", steps=5):
    if not pyautogui:
        open_uri("ms-settings:sound")
        return True, "Sound settings opened. PyAutoGUI is missing."
    key = "volumeup" if direction == "up" else "volumedown"
    try:
        for _ in range(max(1, min(int(steps), 20))):
            pyautogui.press(key)
            time.sleep(0.03)
        return True, f"Volume {direction} command sent."
    except Exception as e:
        open_uri("ms-settings:sound")
        return True, f"Sound settings opened. Volume key failed: {e}"


def set_brightness_percent(percent: int):
    percent = max(0, min(100, int(percent)))
    cmd = f"(Get-WmiObject -Namespace root/WMI -Class WmiMonitorBrightnessMethods).WmiSetBrightness(1,{percent})"
    ok, out = ps_run(cmd, timeout=8)
    if ok:
        return True, f"Brightness was set to {percent}%."
    open_uri("ms-settings:display")
    return True, f"Display settings opened. Direct brightness control did not work: {out}"


def set_dark_mode(enable=True):
    # enable=True means dark mode ON; enable=False means light mode.
    val = 0 if enable else 1
    cmd = (
        f"Set-ItemProperty -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\Themes\\Personalize "
        f"-Name AppsUseLightTheme -Type DWord -Value {val}; "
        f"Set-ItemProperty -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\Themes\\Personalize "
        f"-Name SystemUsesLightTheme -Type DWord -Value {val}"
    )
    ok, out = ps_run(cmd, timeout=8)
    open_uri("ms-settings:personalization-colors")
    return True, "Dark mode was turned ON." if enable and ok else ("Light mode was turned ON." if ok else f"Color settings opened. Direct theme change failed: {out}")


def set_notifications(enable=True):
    # Windows stores notification master switch in HKCU. Explorer/settings may need reopen to reflect it.
    val = 1 if enable else 0
    cmd = "New-Item -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\PushNotifications -Force | Out-Null; " \
          f"Set-ItemProperty -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\PushNotifications -Name ToastEnabled -Type DWord -Value {val}"
    ok, out = ps_run(cmd, timeout=8)
    open_uri("ms-settings:notifications")
    return True, "Notifications were turned ON." if enable and ok else ("Notifications were turned OFF." if ok else f"Notifications settings opened. Direct change failed: {out}")


def set_firewall(enable=True):
    if enable:
        return enable_firewall_profiles()
    # For safety, do not silently turn off firewall. User can still open settings manually.
    open_uri("windowsdefender:")
    return True, "I did not turn the firewall OFF automatically because that can reduce laptop security. Windows Security opened; confirm manually only if you really need it."


def set_defender_realtime(enable=True):
    if enable:
        return defender_realtime_on()
    open_uri("windowsdefender:")
    return True, "I did not turn Defender OFF automatically because it is unsafe and increases attack risk. Windows Security opened."


def set_battery_saver(enable=True):
    # Windows battery saver exact API varies. This opens the page and uses a safe command where possible.
    # /setdcvalueindex can tune power behavior, but not always flip Battery Saver instantly on desktops.
    open_uri("ms-settings:batterysaver")
    state = "ON" if enable else "OFF"
    return True, f"Battery saver settings opened. Windows may block the direct Battery Saver {state} toggle depending on the device/version."


def set_mobile_hotspot(enable=True):
    open_uri("ms-settings:network-mobilehotspot")
    state = "ON" if enable else "OFF"
    return True, f"Mobile hotspot settings opened. Windows does not reliably allow direct Hotspot {state} automation without secure permission."


def set_airplane_mode(enable=True):
    open_uri("ms-settings:network-airplanemode")
    state = "ON" if enable else "OFF"
    return True, f"Airplane mode settings opened. Windows does not provide a reliable safe public API for direct Airplane mode {state} automation."


def set_night_light(enable=True):
    open_uri("ms-settings:nightlight")
    state = "ON" if enable else "OFF"
    return True, f"Night light settings opened. Direct Night Light {state} automation may be restricted by this Windows version."


def set_camera_privacy(enable=True):
    # enable=True means allow camera access; enable=False means block camera access.
    allow = "Allow" if enable else "Deny"
    cmd = (
        "New-Item -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\CapabilityAccessManager\\ConsentStore\\webcam -Force | Out-Null; "
        f"Set-ItemProperty -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\CapabilityAccessManager\\ConsentStore\\webcam -Name Value -Value {allow}"
    )
    ok, out = ps_run(cmd, timeout=8)
    open_uri("ms-settings:privacy-webcam")
    return True, "Camera access was turned ON." if enable and ok else ("Camera access was turned OFF." if ok else f"Camera privacy settings opened. Direct change failed: {out}")


def set_microphone_privacy(enable=True):
    allow = "Allow" if enable else "Deny"
    cmd = (
        "New-Item -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\CapabilityAccessManager\\ConsentStore\\microphone -Force | Out-Null; "
        f"Set-ItemProperty -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\CapabilityAccessManager\\ConsentStore\\microphone -Name Value -Value {allow}"
    )
    ok, out = ps_run(cmd, timeout=8)
    open_uri("ms-settings:privacy-microphone")
    return True, "Microphone access was turned ON." if enable and ok else ("Microphone access was turned OFF." if ok else f"Microphone privacy settings opened. Direct change failed: {out}")


def set_location_privacy(enable=True):
    allow = "Allow" if enable else "Deny"
    cmd = (
        "New-Item -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\CapabilityAccessManager\\ConsentStore\\location -Force | Out-Null; "
        f"Set-ItemProperty -Path HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\CapabilityAccessManager\\ConsentStore\\location -Name Value -Value {allow}"
    )
    ok, out = ps_run(cmd, timeout=8)
    open_uri("ms-settings:privacy-location")
    return True, "Location access was turned ON." if enable and ok else ("Location access was turned OFF." if ok else f"Location privacy settings opened. Direct change failed: {out}")


def set_proxy(enable=True):
    if enable:
        open_uri("ms-settings:network-proxy")
        return True, "Proxy settings opened. To turn Proxy ON, I need an address and port, so I did not enable it without details."
    cmd = (
        "Set-ItemProperty -Path 'HKCU:\\Software\\Microsoft\\Windows\\CurrentVersion\\Internet Settings' -Name ProxyEnable -Value 0; "
        "netsh winhttp reset proxy"
    )
    ok, out = ps_run(cmd, timeout=10)
    open_uri("ms-settings:network-proxy")
    return True, "Proxy was turned OFF." if ok else f"Proxy settings opened. Direct proxy OFF failed: {out}"


def set_vpn(enable=True):
    open_uri("ms-settings:network-vpn")
    state = "connect" if enable else "disconnect"
    return True, f"VPN settings opened. To turn VPN {state}, Windows needs you to select a saved VPN profile."


def all_settings_toggle_command(c):
    """Voice on/off/toggle router for Windows settings.
    It directly changes only safe/reliable settings and opens exact page when Windows blocks direct control.
    """
    c = normalize_words(c)

    # Detect on/off words. Also support Hindi/Hinglish.
    on_words = ["turn on", "switch on", "enable", "start", "on karo", "chalu", "chalu karo", "activate"]
    off_words = ["turn off", "switch off", "disable", "stop", "off karo", "band", "band karo", "deactivate"]
    wants_on = any(x in c for x in on_words) or c.endswith(" on")
    wants_off = any(x in c for x in off_words) or c.endswith(" off")
    if not (wants_on or wants_off):
        return False, ""
    enable = wants_on and not wants_off

    # Volume special commands.
    if any(x in c for x in ["volume up", "increase volume", "voice badhao", "sound badhao"]):
        return change_volume("up", 6)
    if any(x in c for x in ["volume down", "decrease volume", "voice kam", "sound kam"]):
        return change_volume("down", 6)
    if any(x in c for x in ["mute", "unmute", "sound mute", "voice mute"]):
        return set_volume_mute(True)

    # Percentage brightness commands.
    m = re.search(r"brightness[^0-9]*(\d{1,3})", c)
    if m:
        return set_brightness_percent(int(m.group(1)))

    # Common setting toggles.
    if "wifi" in c or "wi fi" in c or "wireless" in c:
        return set_wifi(enable)
    if "bluetooth" in c or "blue tooth" in c:
        return set_bluetooth(enable)
    if "firewall" in c:
        return set_firewall(enable)
    if "defender" in c or "real time protection" in c or "virus protection" in c:
        return set_defender_realtime(enable)
    if "dark mode" in c or "dark theme" in c:
        return set_dark_mode(enable)
    if "light mode" in c or "light theme" in c:
        return set_dark_mode(False if enable else True)
    if "notification" in c:
        return set_notifications(enable)
    if "battery saver" in c:
        return set_battery_saver(enable)
    if "hotspot" in c or "mobile hotspot" in c:
        return set_mobile_hotspot(enable)
    if "airplane" in c or "flight mode" in c:
        return set_airplane_mode(enable)
    if "night light" in c:
        return set_night_light(enable)
    if "camera" in c:
        return set_camera_privacy(enable)
    if "microphone" in c or " mic" in c or c.startswith("mic "):
        return set_microphone_privacy(enable)
    if "location" in c:
        return set_location_privacy(enable)
    if "proxy" in c:
        return set_proxy(enable)
    if "vpn" in c:
        return set_vpn(enable)

    # Display/project commands.
    if any(x in c for x in ["cast", "android tv", "wireless display", "screen mirror", "project"]):
        if enable:
            return cast_command("cast laptop to android tv")
        return cast_command("stop casting")

    return False, ""

def security_command(c):
    c = normalize_words(c)
    if any(x in c for x in ["shield help", "nova shield help", "security help", "protective help"]):
        return True, nova_shield_help()
    if any(x in c for x in ["dark web", "darkweb", "deep web"]):
        return True, dark_web_safety_info()
    if any(x in c for x in ["shield status", "nova shield status", "protection status", "protective status"]):
        return True, nova_shield_status()
    if any(x in c for x in ["enable nova shield", "turn on nova shield", "add shield", "strong shield", "protective shield", "make shield", "shield on"]):
        return enable_nova_shield()
    if any(x in c for x in ["quick scan", "virus scan", "scan laptop", "scan my laptop", "defender scan"]):
        return run_defender_quick_scan()
    if any(x in c for x in ["update defender", "update antivirus", "security update", "defender update"]):
        return update_defender_signatures()
    if any(x in c for x in ["enable firewall", "firewall on", "turn on firewall"]):
        return enable_firewall_profiles()
    if any(x in c for x in ["enable antivirus", "antivirus on", "turn on antivirus", "defender on", "real time protection on"]):
        return defender_realtime_on()
    if any(x in c for x in ["secure my laptop", "protect my laptop", "make laptop secure", "security repair", "fix security", "add antivirus", "add strong firewall", "make secure", "add protective"]):
        return enable_nova_shield()
    if any(x in c for x in ["status", "check", "scan report", "security center", "secure status"]):
        return True, security_status()
    if "vpn" in c:
        run_start("ms-settings:network-vpn")
        return True, "Opening Windows VPN settings. Use a trusted VPN provider only; I will not automate stealth location rotation."
    if "proxy" in c:
        run_start("ms-settings:network-proxy")
        return True, "Opening proxy settings."
    if "firewall" in c:
        run_start("windowsdefender:")
        return True, "Opening Windows Security. Check Firewall & network protection there."
    if "update" in c:
        run_start("ms-settings:windowsupdate")
        return True, "Opening Windows Update settings."
    if "privacy" in c:
        run_start("ms-settings:privacy")
        return True, "Opening Windows Privacy settings."
    if any(x in c for x in ["virus", "defender", "windows security", "security"]):
        run_start("windowsdefender:")
        return True, "Opening Windows Security."
    return False, ""


def install_app(c):
    name = remove_words(c, ["install", "download", "app", "game"])
    if not name:
        return False, "Tell me what to install."
    if shutil.which("winget"):
        subprocess.Popen(f'winget search "{name}" --accept-source-agreements', shell=True)
        return True, f"I opened winget search for {name}. Check the correct package before installing."
    webbrowser.open("ms-windows-store://search/?query=" + quote_plus(name))
    return True, f"Opening Microsoft Store search for {name}."


def uninstall_app(c):
    name = remove_words(c, ["uninstall", "remove", "delete", "app", "game"])
    if not name:
        return False, "Tell me what to uninstall."
    subprocess.Popen("appwiz.cpl", shell=True)
    return True, f"Opening uninstall window. Select {name} manually to avoid deleting the wrong app."


def remember_note(c):
    note = remove_words(c, ["remember", "note", "save", "that"])
    if not note:
        return False, "Tell me what to remember."
    memory.setdefault("notes", []).append({"text": note, "time": datetime.datetime.now().isoformat()})
    save_json(MEMORY_FILE, memory)
    return True, "I saved that note in local memory."

def tell_datetime():
    now = datetime.datetime.now()
    return f"Today is {now.strftime('%d %B %Y')} and the time is {now.strftime('%I:%M %p')}."


def open_named_file(c):
    c = normalize_words(c)
    # Example: satya naam ki word file open karo
    name = c
    for w in ["open", "word", "file", "naam", "ki", "hai", "usko", "karo", "research paper"]:
        name = name.replace(w, " ")
    name = re.sub(r"\s+", " ", name).strip()
    exts = ["*.docx", "*.doc", "*.pdf", "*.txt", "*.xlsx", "*.pptx"]
    candidates = []
    roots = [DESKTOP, Path.home() / "Documents", FILES_DIR, Path.home() / "Downloads"]
    for root in roots:
        if root.exists():
            for ext in exts:
                candidates.extend(root.rglob(ext))
    if name:
        best = [p for p in candidates if name.lower() in p.stem.lower()]
    else:
        best = candidates
    if best:
        best.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        try:
            os.startfile(best[0])
        except Exception:
            pass
        return True, f"Opening file: {best[0].name}"
    return False, f"I could not find the {name or 'requested'} file in Desktop, Documents, or Downloads."


# -------------------- Autopilot Automation --------------------
AUTOPILOT_ACTION_STARTS = [
    "open", "start", "launch", "play", "search", "find", "close", "focus", "write", "type",
    "paste", "create", "make", "turn", "switch", "enable", "disable", "call", "message",
    "take", "analyze", "analyse", "edit", "crop", "merge", "read", "copy", "run", "scan",
    "update", "connect", "learn", "repair", "fix", "set", "volume", "screenshot",
]


def autopilot_help():
    return (
        "Nova Autopilot is ready.\n\n"
        "Use one command with steps:\n"
        "- autopilot open spotify then play jhol song in spotify\n"
        "- auto open whatsapp then search Dishant in WhatsApp\n"
        "- open settings then open bluetooth settings\n\n"
        "Autopilot runs steps in order, stops on safety confirmation or permission blocks, and stays honest if Windows/app security blocks a step."
    )


def plan_mode_response(command):
    goal = str(command or "").strip()
    if not goal:
        return "Plan mode is ON. Tell me the goal first."
    prompt = (
        "Plan mode: do not execute any desktop action. "
        "Make a short practical plan for this user goal. "
        "Use 3 to 6 numbered steps, mention permissions/tools needed, and keep it clear. "
        f"Goal: {goal}"
    )
    if api_status().get("online"):
        reply = ai_reply(prompt)
        if reply and "online brain failed" not in reply.lower():
            return reply
    return "\n".join([
        f"Plan for: {goal}",
        "1. Understand exactly what app, file, or setting is involved.",
        "2. Check whether Nova has the required local permission in Settings.",
        "3. Run the safest direct laptop action first.",
        "4. If Windows or an app blocks automation, show a permission popup and wait for you.",
        "5. Verify the result before saying it is done.",
    ])


def looks_like_autopilot(c):
    if not settings.get("autopilot_enabled", True):
        return c.startswith(("auto ", "automate ", "autopilot ")) or " then " in c or " phir " in c or " fir " in c or " uske baad " in c
    if c in ["autopilot", "automation", "automation help", "autopilot help", "auto help"]:
        return True
    if c.startswith(("auto ", "automate ", "autopilot ")):
        return True
    return any(x in c for x in [" then ", " and then ", " phir ", " fir ", " uske baad ", ";"])


def split_autopilot_steps(raw):
    text = normalize_words(raw)
    for prefix in ["autopilot", "automate", "automatic", "auto"]:
        if text.startswith(prefix + " "):
            text = text[len(prefix):].strip()
            break
    text = text.replace(" uske baad ", " then ").replace(" and then ", " then ").replace(" after that ", " then ")
    text = text.replace(" phir ", " then ").replace(" fir ", " then ")
    text = text.replace(";", " then ")
    text = re.sub(r"\s+and\s+(?=(" + "|".join(AUTOPILOT_ACTION_STARTS) + r")\b)", " then ", text)
    steps = [s.strip(" .,-") for s in re.split(r"\bthen\b", text) if s.strip(" .,-")]
    clean = []
    for step in steps:
        if step in ["and", "then", "please"]:
            continue
        clean.append(step)
    return clean[:8]


def autopilot_command(raw, ask_confirm=None):
    c = normalize_words(raw)
    if c in ["autopilot", "automation", "automation help", "autopilot help", "auto help"]:
        return autopilot_help()
    if not settings.get("autopilot_enabled", True):
        return "Autopilot is OFF in Nova Settings. Turn it ON and allow the popup before I run multi-step automation."
    steps = split_autopilot_steps(raw)
    if len(steps) < 2:
        return autopilot_help()
    results = ["Autopilot started."]
    for index, step in enumerate(steps, 1):
        reply = handle_command(step, ask_confirm=ask_confirm, _autopilot_depth=1)
        if reply == "EXIT":
            results.append(f"{index}. {step}: closing Nova.")
            break
        results.append(f"{index}. {step}: {reply}")
        lower = str(reply).lower()
        if str(reply).startswith("For safety I need") or " is off in nova settings" in lower or "permission is off" in lower or "permission needed" in lower:
            results.append("Autopilot paused because this step needs your permission.")
            break
        if any(x in lower for x in ["i tried, but got an error", "could not", "nahi mila", "not found", "blocked"]):
            results.append("Autopilot stopped because this step did not complete cleanly.")
            break
        time.sleep(0.35)
    return "\n".join(results)


# -------------------- Router --------------------
def needs_confirmation(c):
    return settings.get("confirm_dangerous_actions", True) and any(w in c for w in DANGEROUS_WORDS)


def handle_command(raw, ask_confirm=None, _autopilot_depth=0):
    c = resolve_followup_command(normalize_words(raw))
    if not c:
        return "I am listening."

    confirmed = False
    if c.startswith("confirm "):
        confirmed = True
        c = c.replace("confirm ", "", 1).strip()
    elif c.startswith("yes "):
        confirmed = True
        c = c.replace("yes ", "", 1).strip()

    if _autopilot_depth == 0 and looks_like_autopilot(c):
        return autopilot_command(raw, ask_confirm=ask_confirm)

    if needs_confirmation(c) and not confirmed:
        if ask_confirm:
            if not ask_confirm("This can change or remove something. Do you want to continue?"):
                return "Cancelled."
        else:
            return f"For safety I need a clear confirmation. Say: confirm {c}"

    ok, msg = False, ""
    try:
        if c in ["exit", "quit", "stop assistant"]:
            return "EXIT"
        if c in ["assistant on", "nova on", "turn on assistant", "enable assistant", "on assistant"]:
            return set_assistant_power(True)
        if c in ["assistant off", "nova off", "turn off assistant", "disable assistant", "off assistant"]:
            return set_assistant_power(False)
        if c in ["assistant switch", "assistant power", "assistant on off", "on off status"]:
            return "Assistant is ON." if settings.get("assistant_enabled", True) else "Assistant is OFF."
        if c.startswith(("set pin ", "change pin ", "set access pin ", "change access pin ")):
            pin = remove_words(c, ["set", "change", "access", "pin"])
            ok, msg = set_access_pin(pin)
            return msg
        if not settings.get("assistant_enabled", True):
            allowed_when_off = ["api status", "online status", "brain status", "check api", "assistant doctor", "doctor", "self check", "check assistant", "assistant status", "health check", "working assistant status", "control status", "laptop control status", "full function check", "check all functions", "media status", "media help", "fix nova network", "allow nova network", "fix api network", "allow python internet", "network permission fix"]
            if c not in allowed_when_off:
                return "Assistant is OFF. Press the ON button or say: assistant on."
        if c in ["api status", "online status", "brain status", "check api"]:
            status = api_status()
            if status["online"]:
                return "Online brain connected: " + ", ".join(status["order"]) + "."
            if not status["requests_installed"]:
                return "Online brain is offline because Python requests is not installed."
            if not status.get("order"):
                return "Online brain is offline because no API key was found in .env."
            if not status.get("python_http", False):
                return "Online brain has API keys, but Python/API internet is blocked or filtered. " + (status.get("network_message") or "Run ALLOW_NOVA_NETWORK.cmd and approve the Windows popup.")
            return "Online brain is configured, but provider status could not be verified. Try: assistant doctor."
        if c in ["privacy status", "privacy guard", "privacy guard status", "owner privacy", "safe privacy"]:
            return privacy_status_text()
        if c in ["assistant doctor", "doctor", "self check", "check assistant", "assistant status", "status", "health check", "working assistant status"]:
            return assistant_diagnostics(repair=False)
        if c in ["full function check", "check all functions", "check all function", "test all functions", "all function status", "deep check", "smooth check"]:
            return full_function_audit(repair=False)
        if c in ["repair all functions", "fix all functions", "repair full assistant", "fix everything", "smooth repair"]:
            if not confirmed and ask_confirm is None:
                return "For safety I need a clear confirmation. Say: confirm repair all functions"
            return full_function_audit(repair=True)
        if c in ["fix nova network", "allow nova network", "fix api network", "allow python internet", "network permission fix"]:
            if not confirmed and ask_confirm is None:
                return "For safety I need a clear confirmation. Say: confirm fix nova network"
            helper = BASE_DIR / "ALLOW_NOVA_NETWORK.cmd"
            if helper.exists():
                try:
                    os.startfile(helper)
                    return "Opened Nova network permission helper. Press Yes on the Windows administrator popup if you want to allow Python/API internet and mobile LAN access."
                except Exception as e:
                    return f"Could not open network helper: {e}"
            return "Network helper file is missing."
        if c in ["repair assistant", "fix assistant", "make assistant work", "make it work", "assistant repair", "working assistant", "pura work assistant"]:
            if not confirmed and ask_confirm is None:
                return "For safety I need a clear confirmation. Say: confirm repair assistant"
            return assistant_diagnostics(repair=True)
        if c in ["connect laptop", "connect my laptop", "connect full laptop", "full laptop connect", "connect whole laptop", "connect pure laptop"] or ("connect" in c and "laptop" in c):
            ok, msg = connect_laptop()
        if c in [
            "learn laptop", "scan laptop", "learn my laptop", "scan apps", "learn apps",
            "scan system", "learn system", "scan full system", "refresh laptop profile",
            "refresh system profile", "startup scan", "rescan laptop", "learn full laptop",
            "scan my laptop", "scan this laptop", "learn all apps", "learn all laptop",
        ] or ("learn" in c and "laptop" in c) or ("scan" in c and any(x in c for x in ["laptop", "system", "apps", "settings"])):
            ok, msg = learn_laptop()
        if ("permission" in c or "permision" in c or "control" in c or "allow" in c) and any(x in c for x in ["mouse", "keyboard", "gmail", "google", "tor", "laptop", "full", "all", "sabhi"]):
            settings["control_mode"] = True
            settings["local_control_permission"] = True
            settings["local_only_mode"] = True
            if "google" in c:
                settings["google_browser_permission"] = True
            if "gmail" in c or "email" in c or "mail" in c:
                settings["gmail_browser_permission"] = True
            if "tor" in c:
                settings["tor_browser_permission"] = True
            settings["permission_updated_at"] = datetime.datetime.now().isoformat()
            save_json(SETTINGS_FILE, settings)
            extra = " Tor Browser permission is ON for explicit legal privacy browsing commands." if settings.get("tor_browser_permission", False) else ""
            return "Permission profile saved locally on this laptop. I will use mouse/keyboard automation and your logged-in browser session for Google/Gmail when you give clear commands." + extra + " I still cannot bypass passwords, UAC, browser mic permission, or app security."
        if c in ["control mode on", "full control on", "laptop control on", "do what i say", "do what i say do what i say"] or c.startswith("do what i say"):
            settings["control_mode"] = True
            settings["local_control_permission"] = True
            settings["local_only_mode"] = True
            save_json(SETTINGS_FILE, settings)
            return "Laptop Control Mode is ON. I will execute clear laptop commands directly, with safety limits for passwords, UAC, login, and dangerous actions."
        if c in ["control mode off", "full control off", "laptop control off"]:
            settings["control_mode"] = False
            save_json(SETTINGS_FILE, settings)
            return "Laptop Control Mode is OFF. I will chat more and execute only clear app commands."
        if c in ["control status", "laptop control status"]:
            return "Laptop Control Mode is ON." if settings.get("control_mode", True) else "Laptop Control Mode is OFF."
        if c in ["tor status", "tor browser status", "privacy browser status"] or "tor browser" in c or c.startswith("tor ") or (" tor " in f" {c} "):
            ok, msg = tor_browser_command(raw)
            return msg
        if c in ["mute", "mute voice"]:
            settings["voice_enabled"] = False; save_json(SETTINGS_FILE, settings); return "Voice muted."
        if c in ["unmute", "unmute voice"]:
            settings["voice_enabled"] = True; save_json(SETTINGS_FILE, settings); return "Voice enabled."
        if c in ["voice status", "speech status", "audio status"]:
            return voice_status_text()
        if c in ["fix voice", "clear voice", "voice clear", "fix hindi voice", "hindi voice clear", "make voice clear", "voice clarity"]:
            ok, msg = fix_voice_clarity()
            return msg
        if c in ["setup hindi voice", "install hindi voice", "add hindi voice", "hindi voice setup", "install hindi speech"]:
            ok, msg = setup_hindi_voice()
            return msg
        if c in ["slow voice", "voice slow", "speak slow", "speak slowly"]:
            settings["speak_rate"] = 130; settings["voice_language_mode"] = "auto"; save_json(SETTINGS_FILE, settings); speak("Voice speed is now slower."); return "Voice speed set to slow and clear."
        if c in ["normal voice speed", "voice normal", "speak normal"]:
            settings["speak_rate"] = 145; settings["voice_language_mode"] = "auto"; save_json(SETTINGS_FILE, settings); speak("Voice speed is normal and clear."); return "Voice speed set to normal clear mode."
        if "change voice" in c or c.startswith("voice "):
            voice_name = remove_words(c, ["change", "voice", "to", "set", "select", "please"])
            ok, msg = set_voice_by_name(voice_name)
            return msg
        if c in ["test voice", "voice test"]:
            speak("Hello Adarsh, this is my current voice.")
            return "Testing current voice."
        if any(x in c for x in ["chatgpt", "chat assistant", "copy last answer", "read last answer", "fork last answer", "regenerate", "export chat", "summarize chat", "summary chat"]) and not settings.get("chatgpt_features_enabled", True):
            return "ChatGPT Tools are OFF in Nova Settings. Turn them ON and allow the popup before I use those tools."
        if c in ["chatgpt", "chatgpt help", "chatgpt functions", "chatgpt mode", "chat assistant help", "chat functions"]:
            return chatgpt_feature_help()
        if c in ["new chat", "clear chat", "clear conversation", "start new chat", "new conversation"]:
            return start_new_chat_session()
        if c in ["export chat", "download chat", "save chat", "export history", "save history"]:
            ok, msg = export_chat_history()
            return msg
        if c in ["summarize chat", "summary chat", "chat summary", "summarize history", "summarise chat"]:
            return summarize_chat_history()
        if c in ["copy last answer", "copy last response", "copy answer"]:
            ok, msg = copy_last_answer()
            return msg
        if c in ["read last answer", "read last response", "speak last answer"]:
            ok, msg = read_last_answer()
            return msg
        if c in ["fork last answer", "fork last response", "save fork"]:
            ok, msg = fork_last_answer()
            return msg
        if c in ["regenerate", "regenerate answer", "regenerate last answer", "try again", "answer again"]:
            return regenerate_last_answer()
        if c in ["media help", "photo help", "video help", "image help", "media studio", "media status", "photo status", "video status", "image status"] or looks_like_media_request(c):
            if not settings.get("media_studio_enabled", True):
                return "Media Studio is OFF in Nova Settings. Turn it ON and allow the popup before I use photo/video tools."
            ok, msg = media_command(raw, c)
            if msg:
                return msg
        if c in ["file help", "file studio", "file status", "document help"] or looks_like_file_request(c):
            if not settings.get("file_studio_enabled", True):
                return "File Studio is OFF in Nova Settings. Turn it ON and allow the popup before I create or analyze files."
            ok, msg = file_command(raw, c)
            if msg:
                return msg
        if "calculator" in c and any(x in c for x in ["code", "program", "python", "script", "make", "create", "write", "show"]):
            return calculator_code_reply()
        if any(x in c for x in ["date", "time", "aaj", "today", "current time", "current date"]) and not any(x in c for x in ["code", "file", "python", "create", "make", "write"]):
            return tell_datetime()
        ok_playbook, playbook_msg = app_playbook_command(raw)
        if playbook_msg:
            return playbook_msg

        # Specific actions first, general Google search last.
        if msg:
            pass
        elif any(x in c for x in ["slide", "slides", "page", "pages", "tab", "tabs"]) and "close" in c:
            ok, msg = browser_page_control(c)
        elif settings.get("control_mode", True) and any(x in c for x in ["copy", "paste", "cut", "select all", "save", "undo", "redo", "press ", "enter", "tab", "escape", "space", "backspace", "delete", "click", "scroll", "go back", "go forward", "new tab", "refresh", "reload", "find", "print", "task manager", "lock laptop"]):
            ok, msg = keyboard_mouse_control(c)
        elif c in ["close", "close this", "close it", "close current", "close window"]:
            ok, msg = window_control("close window")
        elif c.startswith(("close ", "stop ", "exit ")) or c.endswith(" close") or c.startswith("band "):
            ok, msg = close_app(c)
        elif c.startswith("focus ") or " focus" in c:
            ok, msg = focus_app(c)
        elif any(x in c for x in ["cast", "android tv", "wireless display", "connect display", "connect to tv", "screen cast", "screen mirror", "project to tv", "project screen", "stop casting", "disconnect cast"]):
            ok, msg = cast_command(c)
        elif any(x in c for x in ["turn on", "turn off", "switch on", "switch off", "enable", "disable", "on karo", "off karo", "chalu", "band karo", "band "]) and any(x in c for x in ["wifi", "wi fi", "bluetooth", "blue tooth", "firewall", "defender", "real time protection", "dark mode", "light mode", "notification", "battery saver", "hotspot", "airplane", "flight mode", "night light", "camera", "microphone", " mic", "location", "proxy", "vpn", "cast", "android tv", "wireless display", "screen mirror", "project"]):
            ok, msg = all_settings_toggle_command(c)
        elif any(x in c for x in ["brightness"]):
            ok, msg = all_settings_toggle_command(c)
        elif any(x in c for x in ["all settings", "setting control", "settings control"]) or (("setting" in c or "settings" in c) and any(x in c for x in ["display", "sound", "wifi", "bluetooth", "cast", "printer", "mouse", "keyboard", "camera", "microphone", "privacy", "apps", "storage", "battery", "power", "notification", "date", "time", "language", "account", "taskbar", "update", "vpn", "proxy", "troubleshoot", "recovery", "about"])):
            ok, msg = windows_settings_command(c)
        elif any(x in c for x in ["security", "secure", "defender", "firewall", "vpn", "proxy", "privacy", "windows update", "virus protection", "antivirus", "shield", "protective", "protect my laptop", "dark web", "darkweb", "deep web"]):
            if not settings.get("security_tools_enabled", True):
                return "Nova Shield/Security Tools are OFF in Nova Settings. Turn them ON and allow the popup before I use security tools."
            ok, msg = security_command(c)
        elif "turn on bluetooth" in c or "enable bluetooth" in c or "bluetooth on" in c:
            ok, msg = set_bluetooth(True)
        elif "turn off bluetooth" in c or "disable bluetooth" in c or "bluetooth off" in c:
            ok, msg = set_bluetooth(False)
        elif "turn on wifi" in c or "enable wifi" in c:
            ok, msg = set_wifi(True)
        elif "turn off wifi" in c or "disable wifi" in c:
            ok, msg = set_wifi(False)
        elif "system settings" in c or "about pc" in c:
            ok, msg = open_app("system settings")
        elif "bluetooth settings" in c or c == "bluetooth":
            ok, msg = open_app("bluetooth settings")
        elif "wifi settings" in c or c == "wifi":
            ok, msg = open_app("wifi settings")
        elif c in ["settings", "open settings"]:
            ok, msg = open_app("settings")
        elif c.startswith(("open ", "show ")) and any(x in c for x in ["desktop", "documents", "downloads", "pictures", "photos", "videos", "music", "onedrive", "c drive", "nova files"]):
            ok, msg = open_folder(c)
        elif looks_like_local_app_request(c):
            ok, msg = open_app("open " + c)
        elif "gmail" in c or (re.search(r"\b(email|mail)\b", c) and any(x in c for x in ["open", "search", "find", "compose", "write", "send"])):
            ok, msg = gmail_command(raw)
        elif "google" in c and any(x in c for x in ["open", "account", "drive", "docs", "sheets", "calendar", "photos"]):
            ok, msg = google_command(raw)
        elif "uninstall" in c or "remove app" in c:
            ok, msg = uninstall_app(c)
        elif re.search(r"\binstall\b", c) or "download app" in c or "download game" in c:
            ok, msg = install_app(c)
        elif "whatsapp" in c and ("browser" in c or "web" in c):
            ok, msg = whatsapp_command(c)
        elif ("whatsapp" in c and ("search" in c or "find" in c)):
            ok, msg = search_whatsapp(c)
        elif "whatsapp" in c or c.startswith("call ") or c.startswith("message "):
            ok, msg = whatsapp_command(c)
        elif ("word file" in c or "research paper" in c or "naam" in c or "named" in c) and "open" in c:
            ok, msg = open_named_file(c)
        elif ("make" in c or "create" in c or ("file" in c and ("write" in c or "code" in c))) and any(x in c for x in ["word", "excel", "sheet", "powerpoint", "ppt", "presentation", "text", "note", "file", "slide"]):
            ok, msg = create_document_from_command(c)
        elif "open" in c and "in browser" in c:
            target = remove_words(c, ["open", "in", "browser", "app"])
            url_map = {"youtube": "https://www.youtube.com", "google": "https://www.google.com", "gmail": "https://mail.google.com"}
            if target in url_map:
                webbrowser.open(url_map[target]); ok, msg = True, f"Opening {target} in browser."
            else:
                webbrowser.open("https://www.google.com/search?q=" + quote_plus(target)); ok, msg = True, f"Searching browser for {target}."
        elif c.startswith(("open ", "start ", "launch ")):
            ok, msg = open_app(c)
        elif any(x in c for x in ["pause song", "pause music", "pause spotify", "pause youtube", "next song", "next track", "previous song", "previous track", "resume song", "resume music", "continue song", "continue music"]):
            ok, msg = media_control_command(c)
        elif c.startswith("play") or " play " in c or "song" in c or ("spotify" in c and not c.startswith(("open ", "start ", "launch "))):
            ok, msg = play_media(c)
        elif "youtube" in c and ("search" in c or "find" in c):
            ok, msg = search_web(c, "youtube")
        elif c.startswith("search") or c.startswith("google") or "search google" in c:
            ok, msg = search_web(c, "google")
        elif c.startswith("type") or c.startswith("write") or c.startswith("paste"):
            ok, msg = type_text(c)
        elif "screenshot" in c:
            ok, msg = screenshot()
        elif "volume" in c or c in ["mute volume", "volume up", "volume down"] or c in ["up up up", "down down down"]:
            ok, msg = volume(c)
        elif any(x in c for x in ["minimize", "maximize", "switch window", "next window", "close window"]):
            ok, msg = window_control(c)
        elif c.startswith("remember") or c.startswith("note"):
            ok, msg = remember_note(c)
        elif "shutdown" in c:
            subprocess.Popen("shutdown /s /t 30", shell=True); ok, msg = True, "Shutdown scheduled in 30 seconds. Run shutdown /a to cancel."
        elif "restart" in c:
            subprocess.Popen("shutdown /r /t 30", shell=True); ok, msg = True, "Restart scheduled in 30 seconds. Run shutdown /a to cancel."
    except Exception as e:
        log(f"Command error: {e}")
        return f"I tried, but got an error: {e}"

    if msg:
        remember_command_context(c, msg)
        return msg
    return ai_reply(raw)

# -------------------- GUI --------------------
WEB_GUI_HTML = r"""<!doctype html>
<html lang="en">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Nova Assistant</title>
  <style>
    :root {
      color-scheme: light dark;
      --bg: #050b08;
      --panel: rgba(4, 18, 12, .92);
      --text: #d8ffe9;
      --muted: #83b49c;
      --line: rgba(0,255,156,.22);
      --blue: #2dd4ff;
      --red: #ff4d6d;
      --yellow: #f8d66d;
      --green: #00ff9c;
    }
    * { box-sizing: border-box; }
    body {
      margin: 0;
      min-height: 100vh;
      font-family: "Cascadia Code", "Consolas", "Segoe UI", monospace;
      background:
        linear-gradient(rgba(0, 255, 156, .04) 1px, transparent 1px),
        linear-gradient(90deg, rgba(0, 255, 156, .04) 1px, transparent 1px),
        radial-gradient(circle at 15% 12%, rgba(45,212,255,.16), transparent 28%),
        radial-gradient(circle at 85% 10%, rgba(0,255,156,.14), transparent 24%),
        #020604;
      background-size: 28px 28px, 28px 28px, auto, auto, auto;
      color: var(--text);
      display: grid;
      place-items: center;
    }
    .app {
      width: min(1180px, calc(100vw - 24px));
      height: min(860px, calc(100vh - 24px));
      display: grid;
      grid-template-rows: auto 1fr auto;
      gap: 14px;
      padding: 18px;
      border: 1px solid var(--line);
      border-radius: 8px;
      background: var(--panel);
      box-shadow: 0 0 0 1px rgba(0,255,156,.08), 0 28px 80px rgba(0,0,0,.42), 0 0 60px rgba(0,255,156,.08);
      backdrop-filter: blur(14px);
    }
    header {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
    }
    .brand {
      display: flex;
      align-items: center;
      gap: 10px;
      font-size: 18px;
      font-weight: 600;
    }
    .dots {
      display: grid;
      grid-template-columns: repeat(2, 9px);
      gap: 3px;
    }
    .dots span { width: 9px; height: 9px; border-radius: 50%; display: block; }
    .b { background: var(--blue); } .r { background: var(--red); }
    .y { background: var(--yellow); } .g { background: var(--green); }
    .status {
      color: var(--muted);
      font-size: 14px;
      min-width: 120px;
      text-align: right;
    }
    .nav {
      display: flex;
      gap: 8px;
      align-items: center;
      justify-content: center;
    }
    .nav button {
      border: 1px solid var(--line);
      border-radius: 999px;
      padding: 9px 13px;
      background: rgba(0,0,0,.24);
      color: var(--muted);
      cursor: pointer;
      font: inherit;
      font-size: 13px;
    }
    .nav button.active {
      background: var(--blue);
      border-color: var(--blue);
      color: #001113;
      font-weight: 600;
    }
    .power-toggle {
      border: 1px solid var(--line);
      border-radius: 999px;
      padding: 9px 14px;
      background: var(--green);
      color: #00150d;
      cursor: pointer;
      font: inherit;
      font-size: 13px;
      font-weight: 700;
      min-width: 72px;
    }
    .power-toggle.off {
      background: var(--red);
    }
    main {
      min-height: 0;
      display: grid;
      grid-template-columns: minmax(190px, 230px) minmax(190px, 250px) minmax(480px, 1fr);
      gap: 18px;
      align-items: stretch;
    }
    body.voice-page main { grid-template-columns: minmax(210px, 260px) minmax(360px, 1fr); }
    body.voice-page .work-panel { display: none; }
    body.voice-page .assistant-panel { border-right: 0; padding-right: 0; }
    body.chat-page main { grid-template-columns: minmax(220px, 280px) minmax(560px, 1fr); }
    body.chat-page .assistant-panel { display: none; }
    .history-panel {
      min-height: 0;
      display: grid;
      grid-template-rows: auto 1fr;
      gap: 10px;
      border-right: 1px solid var(--line);
      padding-right: 14px;
    }
    .assistant-panel {
      display: grid;
      align-content: center;
      justify-items: center;
      gap: 18px;
      text-align: center;
      border-right: 1px solid var(--line);
      padding-right: 18px;
    }
    .orb {
      width: clamp(118px, 20vw, 168px);
      aspect-ratio: 1;
      border: 0;
      border-radius: 50%;
      background:
        radial-gradient(circle at 32% 28%, #fff 0 10%, transparent 11%),
        conic-gradient(from 210deg, var(--blue), var(--red), var(--yellow), var(--green), var(--blue));
      box-shadow: 0 20px 45px rgba(66, 133, 244, .18), inset 0 0 0 18px rgba(255,255,255,.72);
      cursor: pointer;
      transition: transform .18s ease, filter .18s ease;
    }
    .orb:hover { transform: translateY(-2px); }
    .orb.listening {
      animation: pulse 1.1s infinite;
      filter: saturate(1.2);
    }
    @keyframes pulse {
      0%, 100% { box-shadow: 0 20px 45px rgba(66, 133, 244, .18), 0 0 0 0 rgba(66,133,244,.24), inset 0 0 0 18px rgba(255,255,255,.72); }
      50% { box-shadow: 0 20px 45px rgba(66, 133, 244, .22), 0 0 0 22px rgba(66,133,244,0), inset 0 0 0 14px rgba(255,255,255,.64); }
    }
    .prompt {
      font-size: clamp(22px, 3vw, 34px);
      font-weight: 500;
      line-height: 1.12;
      letter-spacing: 0;
      margin: 0;
    }
    .sub {
      color: var(--muted);
      font-size: 16px;
      min-height: 24px;
      margin: 0;
    }
    .chat {
      width: 100%;
      min-height: 0;
      overflow: auto;
      display: grid;
      align-content: start;
      gap: 10px;
      padding: 16px;
      border: 1px solid var(--line);
      border-radius: 8px;
      background: rgba(0,0,0,.22);
    }
    .history-list {
      min-height: 0;
      overflow: auto;
      display: grid;
      align-content: start;
      gap: 8px;
    }
    .history-item {
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 9px 10px;
      background: rgba(0,0,0,.22);
      color: var(--text);
      cursor: pointer;
      text-align: left;
      font: inherit;
      font-size: 13px;
      line-height: 1.3;
    }
    .history-item small {
      display: block;
      color: var(--muted);
      margin-top: 4px;
      white-space: nowrap;
      overflow: hidden;
      text-overflow: ellipsis;
    }
    .work-panel {
      min-width: 0;
      min-height: 0;
      display: grid;
      grid-template-rows: auto 1fr auto;
      gap: 12px;
    }
    .panel-title {
      display: flex;
      align-items: center;
      justify-content: space-between;
      gap: 12px;
      padding-bottom: 8px;
      border-bottom: 1px solid var(--line);
      color: var(--muted);
      font-size: 14px;
    }
    .panel-actions {
      display: flex;
      align-items: center;
      justify-content: flex-end;
      gap: 6px;
      flex-wrap: wrap;
    }
    .msg {
      width: fit-content;
      max-width: min(820px, 96%);
      padding: 12px 15px;
      border-radius: 8px;
      line-height: 1.45;
      text-align: left;
      border: 1px solid var(--line);
      background: rgba(255,255,255,.74);
      box-shadow: inset 0 0 0 1px rgba(0,255,156,.04);
      white-space: pre-wrap;
    }
    .msg pre {
      margin: 8px 0 0;
      padding: 12px;
      border-radius: 8px;
      overflow: auto;
      background: #202124;
      color: #f8fafd;
      white-space: pre-wrap;
      word-break: break-word;
      tab-size: 4;
    }
    .msg code {
      font-family: Consolas, "Cascadia Mono", monospace;
      font-size: 13px;
    }
    .msg-actions {
      display: flex;
      gap: 5px;
      margin-top: 8px;
      flex-wrap: wrap;
    }
    .icon-btn {
      border: 1px solid var(--line);
      border-radius: 8px;
      background: rgba(0,0,0,.28);
      color: var(--muted);
      cursor: pointer;
      min-width: 32px;
      height: 30px;
      padding: 0 8px;
      display: grid;
      place-items: center;
      font-size: 12px;
    }
    .you { margin-left: auto; border-color: rgba(45,212,255,.34); background: rgba(45,212,255,.12); }
    .nova { margin-right: auto; }
    .controls {
      display: grid;
      grid-template-columns: 1fr auto auto;
      gap: 10px;
      align-items: center;
    }
    .reader {
      display: none;
      grid-template-columns: 38px 38px 48px 38px 38px minmax(120px, 1fr) 38px;
      gap: 8px;
      align-items: center;
      padding: 8px 10px;
      border: 1px solid var(--line);
      border-radius: 8px;
      background: #202124;
      color: #fff;
    }
    .reader.active { display: grid; }
    .reader button {
      border: 0;
      border-radius: 999px;
      min-width: 38px;
      padding: 0 8px;
      height: 32px;
      display: grid;
      place-items: center;
      cursor: pointer;
      background: rgba(255,255,255,.14);
      color: #fff;
      font-size: 12px;
    }
    .reader button:disabled {
      cursor: not-allowed;
      opacity: .45;
    }
    .reader button.play {
      min-width: 48px;
      height: 40px;
      background: #fff;
      color: #202124;
      font-size: 13px;
    }
    .reader input[type="range"] {
      width: 100%;
      padding: 0;
      border: 0;
      accent-color: #fff;
      background: transparent;
    }
    .reader .reader-label {
      min-width: 54px;
      color: rgba(255,255,255,.78);
      font-size: 13px;
      text-align: center;
    }
    .modal {
      position: fixed;
      inset: 0;
      display: none;
      place-items: center;
      padding: 18px;
      background: rgba(32,33,36,.44);
      z-index: 20;
    }
    .modal.active { display: grid; }
    .lock-screen {
      position: fixed;
      inset: 0;
      display: none;
      place-items: center;
      padding: 20px;
      background:
        linear-gradient(rgba(0, 255, 156, .06) 1px, transparent 1px),
        linear-gradient(90deg, rgba(0, 255, 156, .06) 1px, transparent 1px),
        radial-gradient(circle at 50% 20%, rgba(0,255,156,.18), transparent 34%),
        #020805;
      background-size: 26px 26px, 26px 26px, auto, auto;
      color: #d8ffe9;
      z-index: 40;
    }
    .lock-screen.active { display: grid; }
    .lock-card {
      width: min(420px, 100%);
      border: 1px solid rgba(0,255,156,.34);
      border-radius: 8px;
      padding: 22px;
      background: rgba(3, 14, 10, .92);
      box-shadow: 0 0 42px rgba(0,255,156,.12);
    }
    .lock-card h2 {
      margin: 0 0 8px;
      font-size: 22px;
      letter-spacing: 0;
    }
    .lock-card p {
      margin: 0 0 16px;
      color: rgba(216,255,233,.72);
      line-height: 1.45;
    }
    .lock-card input {
      width: 100%;
      border-radius: 8px;
      border-color: rgba(0,255,156,.32);
      background: #06120d;
      color: #d8ffe9;
      margin-bottom: 12px;
    }
    .modal-card {
      width: min(440px, 100%);
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 18px;
      background: var(--bg);
      box-shadow: 0 24px 80px rgba(32,33,36,.26);
      color: var(--text);
    }
    .modal-card h2 {
      margin: 0 0 8px;
      font-size: 20px;
      letter-spacing: 0;
    }
    .modal-card p {
      margin: 0 0 16px;
      color: var(--muted);
      line-height: 1.45;
      white-space: pre-wrap;
    }
    .modal-actions {
      display: flex;
      justify-content: flex-end;
      gap: 10px;
    }
    input, textarea {
      border: 1px solid var(--line);
      border-radius: 18px;
      padding: 13px 16px;
      font: inherit;
      background: rgba(0,0,0,.28);
      color: var(--text);
      outline: none;
    }
    textarea {
      min-height: 54px;
      max-height: 132px;
      resize: vertical;
      line-height: 1.35;
    }
    button.small {
      border: 1px solid var(--line);
      border-radius: 999px;
      padding: 12px 18px;
      font: inherit;
      background: rgba(0,0,0,.26);
      color: var(--text);
      cursor: pointer;
    }
    button.primary {
      border-color: transparent;
      background: var(--blue);
      color: #001113;
      font-weight: 600;
    }
    .voice-controls {
      width: 100%;
      display: grid;
      gap: 8px;
    }
    select {
      width: 100%;
      border: 1px solid var(--line);
      border-radius: 8px;
      padding: 10px;
      font: inherit;
      background: rgba(0,0,0,.26);
      color: var(--text);
    }
    .chips {
      display: flex;
      gap: 8px;
      flex-wrap: wrap;
      justify-content: center;
    }
    .chip {
      border: 1px solid var(--line);
      border-radius: 999px;
      padding: 8px 12px;
      background: rgba(0,0,0,.22);
      color: var(--muted);
      cursor: pointer;
      font-size: 13px;
    }
    @media (prefers-color-scheme: dark) {
      :root { --bg: #050b08; --text: #d8ffe9; --muted: #83b49c; --line: rgba(0,255,156,.22); }
      .app, .chat { background: rgba(4,18,12,.92); }
      input, textarea, button.small, .chip, .msg, select, .history-item, .icon-btn, .nav button { background: rgba(32,33,36,.78); }
      .you { background: rgba(45,212,255,.14); }
    }
    @media (max-width: 760px) {
      .app { height: calc(100vh - 16px); padding: 12px; }
      main { grid-template-columns: 1fr; }
      .history-panel { border-right: 0; border-bottom: 1px solid var(--line); padding: 0 0 14px; max-height: 180px; }
      .assistant-panel { border-right: 0; border-bottom: 1px solid var(--line); padding: 0 0 14px; }
      .orb { width: 110px; }
      .controls { grid-template-columns: 1fr; }
      .reader { grid-template-columns: repeat(5, 38px) 1fr 38px; overflow-x: auto; }
    }
    /* Clean assistant visual refresh: keeps all old functions, removes the hacker/terminal look. */
    :root {
      color-scheme: dark;
      --bg: #10141f;
      --panel: rgba(20, 25, 38, .94);
      --surface: rgba(255, 255, 255, .075);
      --surface-2: rgba(255, 255, 255, .11);
      --text: #f4f7fb;
      --muted: #aeb9c9;
      --line: rgba(255, 255, 255, .14);
      --blue: #62a8ff;
      --red: #ff6b7a;
      --yellow: #ffd166;
      --green: #6ee7b7;
      --shadow: 0 24px 80px rgba(0, 0, 0, .34);
    }
    body {
      font-family: "Segoe UI", Inter, system-ui, -apple-system, BlinkMacSystemFont, sans-serif;
      background:
        radial-gradient(circle at 16% 12%, rgba(98, 168, 255, .24), transparent 34%),
        radial-gradient(circle at 86% 18%, rgba(110, 231, 183, .18), transparent 30%),
        linear-gradient(145deg, #111827 0%, #151b2c 46%, #0f172a 100%);
      background-size: auto;
      letter-spacing: 0;
    }
    .app {
      width: min(1280px, calc(100vw - 22px));
      height: min(900px, calc(100vh - 22px));
      border-color: rgba(255, 255, 255, .16);
      background: linear-gradient(180deg, rgba(25, 31, 47, .96), rgba(17, 24, 39, .94));
      box-shadow: var(--shadow);
    }
    header {
      padding-bottom: 6px;
      border-bottom: 1px solid rgba(255, 255, 255, .1);
    }
    .brand {
      font-family: "Segoe UI", system-ui, sans-serif;
      font-size: 19px;
      letter-spacing: 0;
    }
    .dots {
      width: 34px;
      height: 34px;
      grid-template-columns: repeat(2, 10px);
      place-content: center;
      border: 1px solid rgba(255, 255, 255, .14);
      border-radius: 8px;
      background: rgba(255, 255, 255, .08);
    }
    .dots span { width: 10px; height: 10px; }
    .nav {
      padding: 4px;
      border: 1px solid rgba(255, 255, 255, .12);
      border-radius: 999px;
      background: rgba(255, 255, 255, .07);
    }
    .nav button,
    button.small,
    .chip,
    .icon-btn,
    .history-item {
      background: rgba(255, 255, 255, .08);
      border-color: rgba(255, 255, 255, .14);
      color: var(--text);
      transition: transform .15s ease, background .15s ease, border-color .15s ease;
    }
    .nav button:hover,
    button.small:hover,
    .chip:hover,
    .icon-btn:hover,
    .history-item:hover {
      transform: translateY(-1px);
      background: rgba(255, 255, 255, .13);
      border-color: rgba(255, 255, 255, .24);
    }
    .nav button.active,
    button.primary {
      background: linear-gradient(135deg, #62a8ff, #6ee7b7);
      color: #07111f;
      border-color: transparent;
      box-shadow: 0 10px 28px rgba(98, 168, 255, .22);
    }
    .power-toggle {
      background: linear-gradient(135deg, #6ee7b7, #b9fbc0);
      color: #07111f;
      border-color: transparent;
      box-shadow: 0 10px 26px rgba(110, 231, 183, .2);
    }
    .power-toggle.off {
      background: linear-gradient(135deg, #ff6b7a, #ffadad);
      color: #27060a;
    }
    main {
      grid-template-columns: minmax(220px, 280px) minmax(220px, 300px) minmax(560px, 1fr);
      gap: 16px;
    }
    .history-panel,
    .assistant-panel {
      border-color: rgba(255, 255, 255, .12);
    }
    .history-item {
      border-radius: 8px;
      padding: 11px 12px;
      box-shadow: none;
    }
    .history-item small,
    .sub,
    .status,
    .panel-title,
    .chip {
      color: var(--muted);
    }
    .orb {
      background:
        radial-gradient(circle at 32% 28%, #ffffff 0 10%, transparent 11%),
        conic-gradient(from 210deg, #62a8ff, #ff6b7a, #ffd166, #6ee7b7, #62a8ff);
      box-shadow: 0 18px 48px rgba(98, 168, 255, .18), inset 0 0 0 18px rgba(255, 255, 255, .78);
    }
    .chat {
      padding: 18px;
      background: rgba(255, 255, 255, .055);
      border-color: rgba(255, 255, 255, .13);
    }
    .msg {
      background: rgba(255, 255, 255, .09);
      color: var(--text);
      border-color: rgba(255, 255, 255, .13);
      box-shadow: 0 10px 28px rgba(0, 0, 0, .1);
      font-size: 14.5px;
    }
    .you {
      background: rgba(98, 168, 255, .18);
      border-color: rgba(98, 168, 255, .28);
    }
    .nova {
      background: rgba(255, 255, 255, .085);
    }
    .msg pre {
      background: #0f172a;
      border: 1px solid rgba(255, 255, 255, .12);
      color: #edf5ff;
    }
    input,
    textarea,
    select {
      background: rgba(255, 255, 255, .09);
      border-color: rgba(255, 255, 255, .14);
      color: var(--text);
      font-family: "Segoe UI", system-ui, sans-serif;
    }
    textarea {
      min-height: 62px;
      border-radius: 8px;
    }
    .reader {
      background: rgba(15, 23, 42, .96);
      border-color: rgba(255, 255, 255, .16);
      box-shadow: 0 18px 50px rgba(0, 0, 0, .22);
    }
    .media-tools {
      display: grid;
      grid-template-columns: minmax(180px, 1fr) minmax(260px, 2fr) 56px;
      gap: 8px;
      align-items: center;
      padding: 10px;
      border: 1px solid rgba(255, 255, 255, .12);
      border-radius: 8px;
      background: rgba(255, 255, 255, .055);
    }
    .media-tools input[type="file"] {
      max-width: 100%;
      color: var(--muted);
    }
    .media-tools input[type="text"] {
      min-width: 0;
      border: 1px solid rgba(255, 255, 255, .14);
      border-radius: 8px;
      padding: 10px 12px;
    }
    .modal {
      background: rgba(4, 8, 17, .62);
      backdrop-filter: blur(10px);
    }
    .modal-card,
    .lock-card {
      background: #151b2c;
      border-color: rgba(255, 255, 255, .14);
      color: var(--text);
      box-shadow: var(--shadow);
    }
    .lock-screen {
      background:
        radial-gradient(circle at 20% 20%, rgba(98, 168, 255, .24), transparent 34%),
        linear-gradient(145deg, #111827, #0f172a);
    }
    .lock-card input {
      background: rgba(255, 255, 255, .09);
      border-color: rgba(255, 255, 255, .14);
      color: var(--text);
    }
    @media (prefers-color-scheme: dark) {
      .app, .chat { background: linear-gradient(180deg, rgba(25, 31, 47, .96), rgba(17, 24, 39, .94)); }
      input, textarea, button.small, .chip, .msg, select, .history-item, .icon-btn, .nav button { background: rgba(255, 255, 255, .08); }
      .you { background: rgba(98, 168, 255, .18); }
    }
    @media (max-width: 760px) {
      header { flex-wrap: wrap; }
      .nav { order: 3; width: 100%; }
      .nav button { flex: 1; }
      .app { width: calc(100vw - 12px); height: calc(100vh - 12px); }
      body.chat-page main,
      body.voice-page main,
      main { grid-template-columns: 1fr; }
      .media-tools { grid-template-columns: 1fr; }
    }
    /* Bigger ChatGPT-style workspace with hidden History/Settings drawers. */
    .app { position: relative; }
    body.chat-page main {
      grid-template-columns: minmax(0, 1fr);
    }
    body.chat-page .work-panel {
      grid-column: 1 / -1;
      min-height: 0;
      grid-template-rows: auto minmax(0, 1fr) auto auto;
    }
    body.chat-page .chat {
      min-height: 0;
      max-height: none;
      padding: 22px;
    }
    body.chat-page .msg {
      max-width: min(920px, 94%);
      font-size: 15.5px;
    }
    .history-panel,
    .settings-panel {
      position: absolute;
      top: 82px;
      bottom: 104px;
      width: min(360px, calc(100vw - 34px));
      z-index: 12;
      display: none;
      grid-template-rows: auto 1fr;
      gap: 10px;
      padding: 14px;
      border: 1px solid rgba(255, 255, 255, .14);
      border-radius: 8px;
      background: rgba(20, 25, 38, .98);
      box-shadow: 0 24px 70px rgba(0, 0, 0, .35);
      backdrop-filter: blur(14px);
    }
    .history-panel { left: 18px; }
    .settings-panel { right: 18px; }
    body.history-open .history-panel,
    body.settings-open .settings-panel {
      display: grid;
    }
    .settings-list {
      min-height: 0;
      overflow: auto;
      display: grid;
      gap: 10px;
      align-content: start;
    }
    .setting-item {
      display: grid;
      grid-template-columns: 1fr auto;
      gap: 10px;
      align-items: center;
      padding: 12px;
      border: 1px solid rgba(255, 255, 255, .12);
      border-radius: 8px;
      background: rgba(255, 255, 255, .07);
    }
    .setting-item strong { display: block; font-size: 14px; }
    .setting-item small { display: block; color: var(--muted); line-height: 1.35; margin-top: 4px; }
    .switch {
      position: relative;
      width: 48px;
      height: 28px;
      display: inline-block;
    }
    .switch input { opacity: 0; width: 0; height: 0; }
    .slider {
      position: absolute;
      cursor: pointer;
      inset: 0;
      border-radius: 999px;
      background: rgba(255, 255, 255, .18);
      border: 1px solid rgba(255, 255, 255, .18);
      transition: .18s ease;
    }
    .slider::before {
      content: "";
      position: absolute;
      height: 22px;
      width: 22px;
      left: 2px;
      top: 2px;
      border-radius: 50%;
      background: #fff;
      transition: .18s ease;
    }
    .switch input:checked + .slider {
      background: linear-gradient(135deg, #62a8ff, #6ee7b7);
    }
    .switch input:checked + .slider::before {
      transform: translateX(20px);
    }
    .modal-card {
      position: relative;
    }
    .modal-close {
      position: absolute;
      top: 10px;
      right: 10px;
      width: 32px;
      height: 32px;
      border-radius: 999px;
      border: 1px solid rgba(255, 255, 255, .14);
      background: rgba(255, 255, 255, .08);
      color: var(--text);
      cursor: pointer;
    }
    .controls {
      grid-template-columns: auto minmax(0, 1fr) auto auto auto;
      position: relative;
    }
    .controls textarea {
      min-height: 76px;
      max-height: 190px;
      font-size: 16px;
    }
    .composer-menu-btn {
      width: 44px;
      height: 44px;
      border-radius: 999px;
      font-size: 24px;
      line-height: 1;
      padding: 0;
    }
    .composer-toolbox {
      position: absolute;
      left: 0;
      bottom: calc(100% + 10px);
      width: min(320px, calc(100vw - 42px));
      display: none;
      z-index: 22;
      padding: 10px;
      border: 1px solid rgba(255, 255, 255, .15);
      border-radius: 8px;
      background: rgba(32, 33, 36, .98);
      box-shadow: 0 22px 58px rgba(0, 0, 0, .38);
      backdrop-filter: blur(14px);
    }
    .composer-toolbox.active { display: grid; gap: 6px; }
    .tool-row {
      width: 100%;
      min-height: 42px;
      display: grid;
      grid-template-columns: 28px 1fr auto;
      align-items: center;
      gap: 8px;
      padding: 8px 10px;
      border: 0;
      border-radius: 8px;
      background: transparent;
      color: var(--text);
      font: inherit;
      text-align: left;
      cursor: pointer;
    }
    .tool-row:hover { background: rgba(255, 255, 255, .08); }
    .tool-row .tool-icon {
      color: var(--muted);
      font-size: 17px;
      text-align: center;
    }
    .tool-row small {
      display: block;
      margin-top: 2px;
      color: var(--muted);
      font-size: 11px;
      line-height: 1.2;
    }
    .tool-divider {
      height: 1px;
      background: rgba(255, 255, 255, .12);
      margin: 5px 0;
    }
    .mini-switch {
      position: relative;
      width: 42px;
      height: 24px;
      display: inline-block;
    }
    .mini-switch input { opacity: 0; width: 0; height: 0; }
    .mini-switch span {
      position: absolute;
      cursor: pointer;
      inset: 0;
      border-radius: 999px;
      background: rgba(255, 255, 255, .18);
      border: 1px solid rgba(255, 255, 255, .16);
      transition: .16s ease;
    }
    .mini-switch span::before {
      content: "";
      position: absolute;
      width: 18px;
      height: 18px;
      left: 2px;
      top: 2px;
      border-radius: 50%;
      background: #fff;
      transition: .16s ease;
    }
    .mini-switch input:checked + span {
      background: linear-gradient(135deg, #62a8ff, #6ee7b7);
    }
    .mini-switch input:checked + span::before {
      transform: translateX(18px);
    }
    @media (max-width: 760px) {
      .history-panel,
      .settings-panel {
        top: 122px;
        bottom: 150px;
        left: 10px;
        right: 10px;
        width: auto;
      }
      .controls { grid-template-columns: auto 1fr auto auto; }
      .controls textarea { grid-column: 1 / -1; }
      .composer-toolbox { left: 0; bottom: calc(100% + 8px); }
    }
  </style>
</head>
<body class="chat-page">
  <div class="app">
    <header>
      <div class="brand"><span class="dots"><span class="b"></span><span class="r"></span><span class="y"></span><span class="g"></span></span> Nova Assistant</div>
      <nav class="nav" aria-label="Assistant mode">
        <button id="chatPageBtn" class="active" type="button">Chat Assistant</button>
        <button id="voicePageBtn" type="button">Voice Assistant</button>
      </nav>
      <button id="powerToggle" class="power-toggle" type="button" aria-pressed="true" title="Turn Nova assistant on or off">ON</button>
      <div id="status" class="status">Checking API</div>
    </header>
    <main>
      <section class="history-panel">
        <div class="panel-title">
          <span>History</span>
          <button id="clearHistory" class="icon-btn" title="Refresh history" type="button">↻</button>
        </div>
        <div id="historyList" class="history-list"></div>
      </section>
      <section id="settingsPanel" class="settings-panel">
        <div class="panel-title">
          <span>Settings</span>
          <button id="settingsClose" class="icon-btn" title="Close settings" type="button">&#10005;</button>
        </div>
        <div id="featureSettingsList" class="settings-list"></div>
      </section>
      <section class="assistant-panel">
        <button id="mic" class="orb" aria-label="Start listening" title="Start listening"></button>
        <div>
          <h1 id="prompt" class="prompt">Hi Adarsh.</h1>
          <p id="sub" class="sub">Click Start Voice once to turn mic ON. Click again to turn OFF.</p>
        </div>
        <button id="voiceButton" class="small primary" type="button">Start Voice</button>
        <div class="voice-controls">
          <select id="voiceSelect" aria-label="Assistant voice"></select>
          <button id="testVoice" class="small" type="button">Test Voice</button>
        </div>
      </section>
      <section class="work-panel">
        <div class="panel-title">
          <span>Conversation</span>
          <div class="panel-actions" aria-label="Chat tools">
            <button id="historyToggle" class="icon-btn" title="Open or close history" type="button">&#9776;</button>
            <button id="settingsToggle" class="icon-btn" title="Open or close settings" type="button">&#9881;</button>
            <button id="newChatButton" class="icon-btn" title="New chat, keep old history saved" type="button">&#43;</button>
            <button id="exportChatButton" class="icon-btn" title="Export chat history" type="button">&#8681;</button>
            <button id="chatHelpButton" class="icon-btn" title="ChatGPT-style tools" type="button">?</button>
          </div>
        </div>
        <div id="chat" class="chat" aria-live="polite"></div>
        <div class="chips" hidden></div>
        <div class="media-tools" aria-label="File and Media Studio">
          <input id="mediaFile" type="file" accept="image/*,video/*,.txt,.md,.py,.js,.ts,.html,.css,.json,.csv,.docx,.xlsx,.pptx,.pdf,.log,.bat,.cmd,.ps1,.xml,.yaml,.yml,.ini,.env,.rtf" title="Upload file, photo, or video" multiple>
          <input id="mediaPrompt" type="text" autocomplete="off" placeholder="Select files, then say what to do: analyze, crop 9:16, cinematic portrait, merge, create video...">
          <button id="mediaRun" class="small primary" type="button" title="Run selected file/media task">&#9654;</button>
        </div>
      </section>
    </main>
    <form id="form" class="controls">
      <button id="composerMenuButton" class="small composer-menu-btn" type="button" title="Open tools and permissions" aria-expanded="false">&#43;</button>
      <div id="composerToolbox" class="composer-toolbox" aria-label="Composer tools">
        <button id="attachFilesButton" class="tool-row" type="button">
          <span class="tool-icon">&#8853;</span>
          <span>Add photos &amp; files<small>Attach selected files to Media/File Studio</small></span>
          <span></span>
        </button>
        <label class="tool-row" for="planModeToggle">
          <span class="tool-icon">&#8756;</span>
          <span>Plan mode<small>Plan first; do not execute</small></span>
          <span class="mini-switch"><input id="planModeToggle" type="checkbox"><span></span></span>
        </label>
        <label class="tool-row" for="pursueGoalToggle">
          <span class="tool-icon">&#9673;</span>
          <span>Pursue goal<small>Use safe autopilot for clear steps</small></span>
          <span class="mini-switch"><input id="pursueGoalToggle" type="checkbox"><span></span></span>
        </label>
        <div class="tool-divider"></div>
        <button id="pluginsButton" class="tool-row" type="button">
          <span class="tool-icon">&#8759;</span>
          <span>Plugins<small>Open Nova feature tools</small></span>
          <span>&#8250;</span>
        </button>
        <button id="permissionsButton" class="tool-row" type="button">
          <span class="tool-icon">&#43;</span>
          <span>Default permissions<small>Open local feature permissions</small></span>
          <span>&#8964;</span>
        </button>
      </div>
      <textarea id="text" autocomplete="off" rows="2" placeholder="Optional: type here if mic is blocked"></textarea>
      <button id="micSmall" class="small" type="button">Mic</button>
      <button id="desktopMic" class="small" type="button" title="Same toggle as Start Voice: desktop microphone ON/OFF">Desktop Mic</button>
      <button id="sendButton" class="small primary" type="submit">Send</button>
    </form>
    <div id="reader" class="reader" aria-label="Read answer controls">
      <button id="readerShuffle" type="button" title="Restart answer">&#8634;</button>
      <button id="readerPrev" type="button" title="Previous line">&#9664;</button>
      <button id="readerPlay" class="play" type="button" title="Play or pause">&#9654;</button>
      <button id="readerNext" type="button" title="Next line">&#9654;</button>
      <button id="readerStop" type="button" title="Stop">&#9632;</button>
      <input id="readerProgress" type="range" min="0" max="0" value="0" aria-label="Read progress">
      <button id="readerClose" type="button" title="Close reader">&#10005;</button>
    </div>
  </div>
  <div id="permissionModal" class="modal" role="dialog" aria-modal="true" aria-labelledby="permissionTitle">
    <div class="modal-card">
      <button id="permissionClose" class="modal-close" type="button" title="Close permission popup">&#10005;</button>
      <h2 id="permissionTitle">Permission needed</h2>
      <p id="permissionText">Nova needs confirmation before doing this.</p>
      <div class="modal-actions">
        <button id="permissionCancel" class="small" type="button">Cancel</button>
        <button id="permissionAllow" class="small primary" type="button">Allow</button>
      </div>
    </div>
  </div>
  <div id="lockScreen" class="lock-screen" role="dialog" aria-modal="true" aria-labelledby="lockTitle">
    <div class="lock-card">
      <h2 id="lockTitle">Nova Locked</h2>
      <p>Enter your local PIN. This lock stays on this laptop only.</p>
      <input id="pinInput" type="password" inputmode="numeric" autocomplete="current-password" placeholder="PIN">
      <button id="unlockButton" class="small primary" type="button">Unlock</button>
      <p id="lockMessage"></p>
    </div>
  </div>
  <script>
    const mic = document.getElementById('mic');
    const voiceButton = document.getElementById('voiceButton');
    const micSmall = document.getElementById('micSmall');
    const desktopMic = document.getElementById('desktopMic');
    const chatPageBtn = document.getElementById('chatPageBtn');
    const voicePageBtn = document.getElementById('voicePageBtn');
    const voiceSelect = document.getElementById('voiceSelect');
    const testVoice = document.getElementById('testVoice');
    const powerToggle = document.getElementById('powerToggle');
    const statusEl = document.getElementById('status');
    const promptEl = document.getElementById('prompt');
    const subEl = document.getElementById('sub');
    const chat = document.getElementById('chat');
    const historyList = document.getElementById('historyList');
    const clearHistory = document.getElementById('clearHistory');
    const historyToggle = document.getElementById('historyToggle');
    const settingsToggle = document.getElementById('settingsToggle');
    const settingsPanel = document.getElementById('settingsPanel');
    const settingsClose = document.getElementById('settingsClose');
    const featureSettingsList = document.getElementById('featureSettingsList');
    const newChatButton = document.getElementById('newChatButton');
    const exportChatButton = document.getElementById('exportChatButton');
    const chatHelpButton = document.getElementById('chatHelpButton');
    const composerMenuButton = document.getElementById('composerMenuButton');
    const composerToolbox = document.getElementById('composerToolbox');
    const attachFilesButton = document.getElementById('attachFilesButton');
    const planModeToggle = document.getElementById('planModeToggle');
    const pursueGoalToggle = document.getElementById('pursueGoalToggle');
    const pluginsButton = document.getElementById('pluginsButton');
    const permissionsButton = document.getElementById('permissionsButton');
    const mediaFile = document.getElementById('mediaFile');
    const mediaPrompt = document.getElementById('mediaPrompt');
    const mediaRun = document.getElementById('mediaRun');
    const form = document.getElementById('form');
    const text = document.getElementById('text');
    const sendButton = document.getElementById('sendButton');
    const reader = document.getElementById('reader');
    const readerShuffle = document.getElementById('readerShuffle');
    const readerPrev = document.getElementById('readerPrev');
    const readerPlay = document.getElementById('readerPlay');
    const readerNext = document.getElementById('readerNext');
    const readerStop = document.getElementById('readerStop');
    const readerClose = document.getElementById('readerClose');
    const readerProgress = document.getElementById('readerProgress');
    const permissionModal = document.getElementById('permissionModal');
    const permissionText = document.getElementById('permissionText');
    const permissionAllow = document.getElementById('permissionAllow');
    const permissionCancel = document.getElementById('permissionCancel');
    const permissionClose = document.getElementById('permissionClose');
    const lockScreen = document.getElementById('lockScreen');
    const pinInput = document.getElementById('pinInput');
    const unlockButton = document.getElementById('unlockButton');
    const lockMessage = document.getElementById('lockMessage');
    const SpeechRecognition = window.SpeechRecognition || window.webkitSpeechRecognition;
    let recognition = null;
    let listening = false;
    let micWanted = false;
    let desktopMicWanted = false;
    let desktopMicAbort = null;
    let desktopMicErrorCount = 0;
    let voiceBusy = false;
    let readerText = '';
    let readerLines = [];
    let readerIndex = 0;
    let readerPaused = false;
    let readerActive = false;
    let readerUtterance = null;
    let readerManualStop = false;
    let readerRunId = 0;
    let readerTimer = null;
    let testVoiceBusy = false;
    let historyTimer = null;
    let speechNetworkErrors = 0;
    let pendingConfirmCommand = '';
    let pendingPermissionAction = null;
    let setupPopupShown = false;
    let shownNeedIds = new Set();
    let sendBusy = false;
    let recognitionRestartTimer = null;
    let accessToken = '';
    let assistantPowerOn = true;

    function setPage(page) {
      document.body.classList.toggle('voice-page', page === 'voice');
      document.body.classList.toggle('chat-page', page !== 'voice');
      voicePageBtn.classList.toggle('active', page === 'voice');
      chatPageBtn.classList.toggle('active', page !== 'voice');
      if (page !== 'voice' && micWanted) toggleVoice();
    }

    function setComposerMenu(open) {
      composerToolbox.classList.toggle('active', Boolean(open));
      composerMenuButton.setAttribute('aria-expanded', open ? 'true' : 'false');
      composerMenuButton.textContent = open ? '\u00D7' : '+';
    }

    function syncComposerModesFromSettings(items = []) {
      const plan = items.find(item => item.key === 'plan_mode_enabled');
      const pursue = items.find(item => item.key === 'pursue_goal_mode_enabled');
      if (plan) planModeToggle.checked = Boolean(plan.enabled);
      if (pursue) pursueGoalToggle.checked = Boolean(pursue.enabled);
    }

    function escapeHtml(value) {
      return String(value).replace(/[&<>"']/g, ch => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[ch]));
    }
    function authHeaders(extra = {}) {
      const headers = Object.assign({}, extra);
      if (accessToken) headers['X-Nova-Token'] = accessToken;
      return headers;
    }
    async function copyToClipboard(value) {
      try {
        await navigator.clipboard.writeText(value);
        subEl.textContent = 'Copied.';
      } catch (e) {
        text.value = value;
        text.focus();
        text.select();
        subEl.textContent = 'Clipboard blocked, so I placed it in the chat box.';
      }
    }
    function renderMessage(value) {
      const parts = String(value).split(/```([\s\S]*?)```/g);
      return parts.map((part, index) => {
        if (index % 2 === 1) {
          const code = part.replace(/^[a-zA-Z0-9_+-]+\n/, '');
          return `<pre><code>${escapeHtml(code.trim())}</code></pre>`;
        }
        return escapeHtml(part);
      }).join('');
    }
    function makeActionButton(label, title, fn) {
      const btn = document.createElement('button');
      btn.className = 'icon-btn';
      btn.type = 'button';
      btn.title = title;
      btn.textContent = label;
      btn.addEventListener('click', fn);
      return btn;
    }
    function trimChat() {
      while (chat.children.length > 80) {
        chat.removeChild(chat.firstElementChild);
      }
    }
    function add(kind, value, userText = '') {
      const div = document.createElement('div');
      div.className = `msg ${kind}`;
      div.innerHTML = renderMessage(value);
      if (kind === 'you') {
        const actions = document.createElement('div');
        actions.className = 'msg-actions';
        actions.appendChild(makeActionButton('\u29C9', 'Copy prompt', () => copyToClipboard(value)));
        actions.appendChild(makeActionButton('\u270E', 'Edit prompt', () => { text.value = value; text.focus(); }));
        div.appendChild(actions);
      }
      if (kind === 'nova') {
        const actions = document.createElement('div');
        actions.className = 'msg-actions';
        actions.appendChild(makeActionButton('\u29C9', 'Copy answer', () => copyToClipboard(value)));
        actions.appendChild(makeActionButton('\uD83D\uDD0A', 'Read / stop this answer', () => openReader(value)));
        actions.appendChild(makeActionButton('\uD83D\uDC4D', 'Like', () => feedback('like', userText, value)));
        actions.appendChild(makeActionButton('\uD83D\uDC4E', 'Dislike', () => feedback('dislike', userText, value)));
        actions.appendChild(makeActionButton('\u2935', 'Fork / continue from here', () => { text.value = `Continue from this answer:\n${value}\n\n`; text.focus(); }));
        if (userText) actions.appendChild(makeActionButton('\u27F3', 'Regenerate answer', () => send(userText)));
        div.appendChild(actions);
      }
      chat.appendChild(div);
      trimChat();
      requestAnimationFrame(() => { chat.scrollTop = chat.scrollHeight; });
    }
    async function feedback(kind, user, assistant) {
      await fetch('/api/feedback', {
        method: 'POST',
        headers: authHeaders({'Content-Type': 'application/json'}),
        body: JSON.stringify({kind, user, assistant})
      });
      subEl.textContent = kind === 'like' ? 'Saved as helpful.' : 'Saved feedback.';
      loadHistory();
    }
    function readFileAsDataUrl(file) {
      return new Promise((resolve, reject) => {
        const reader = new FileReader();
        reader.onload = () => resolve(reader.result);
        reader.onerror = () => reject(reader.error || new Error('File read failed'));
        reader.readAsDataURL(file);
      });
    }
    async function uploadSelectedMedia(analyze = false) {
      const files = Array.from(mediaFile.files || []);
      if (!files.length) {
        add('nova', 'Select a file, photo, or video first.');
        return;
      }
      mediaRun.disabled = true;
      subEl.textContent = `Uploading ${files.length} file(s) to local Nova folder...`;
      try {
        for (const file of files) {
          const dataUrl = await readFileAsDataUrl(file);
          const res = await fetch('/api/media/upload', {
            method: 'POST',
            headers: authHeaders({'Content-Type': 'application/json'}),
            body: JSON.stringify({filename: file.name, data: dataUrl, analyze, prompt: mediaPrompt.value || ''})
          });
          const data = await res.json();
          if (res.status === 403) {
            setLocked(true);
            throw new Error(data.reply || 'Nova is locked.');
          }
          if (!res.ok || data.ok === false) throw new Error(data.reply || data.error || 'Upload failed');
          add('nova', data.reply || 'File uploaded.');
        }
        clearTimeout(historyTimer);
        historyTimer = setTimeout(loadHistory, 450);
      } catch (e) {
        add('nova', 'File upload failed: ' + e.message);
      } finally {
        mediaRun.disabled = false;
      }
    }
    async function runSmartMediaTask() {
      const files = Array.from(mediaFile.files || []);
      const prompt = (mediaPrompt.value || '').trim();
      if (!files.length && !prompt) {
        add('nova', 'Select a photo/video/file or type a prompt first.');
        return;
      }
      mediaRun.disabled = true;
      mediaRun.textContent = '...';
      subEl.textContent = files.length ? `Working on ${files.length} selected file(s)...` : 'Creating from prompt...';
      try {
        const payloadFiles = [];
        for (const file of files) {
          const dataUrl = await readFileAsDataUrl(file);
          payloadFiles.push({filename: file.name, type: file.type || '', data: dataUrl});
        }
        const res = await fetch('/api/media/run', {
          method: 'POST',
          headers: authHeaders({'Content-Type': 'application/json'}),
          body: JSON.stringify({prompt, files: payloadFiles})
        });
        const data = await res.json();
        if (res.status === 403) {
          setLocked(true);
          throw new Error(data.reply || 'Nova is locked.');
        }
        if (!res.ok || data.ok === false) throw new Error(data.reply || data.error || 'Media task failed');
        add('you', prompt || `Selected ${files.length} file(s)`);
        add('nova', data.reply || 'Done.', prompt);
        subEl.textContent = 'Media task complete.';
        clearTimeout(historyTimer);
        historyTimer = setTimeout(loadHistory, 450);
      } catch (e) {
        add('nova', 'Media task failed: ' + e.message);
        subEl.textContent = 'Media task failed.';
      } finally {
        mediaRun.disabled = false;
        mediaRun.textContent = '\u25B6';
      }
    }
    function splitReaderText(value) {
      return String(value)
        .replace(/```[\s\S]*?```/g, block => block.replace(/```/g, ''))
        .split(/(?<=[.!?])\s+|\n+/)
        .map(x => x.trim())
        .filter(Boolean);
    }
    function updateReader() {
      readerProgress.max = Math.max(0, readerLines.length - 1);
      readerProgress.value = Math.min(readerIndex, Math.max(0, readerLines.length - 1));
      readerPlay.textContent = readerActive && !readerPaused ? '\u275A\u275A' : '\u25B6';
      const hasLines = readerLines.length > 0;
      [readerShuffle, readerPrev, readerPlay, readerNext, readerStop, readerProgress].forEach(el => el.disabled = !hasLines);
    }
    function readerDelayForLine(line) {
      const words = String(line || '').trim().split(/\s+/).filter(Boolean).length;
      return Math.max(1400, Math.min(14000, words * 430 + 900));
    }
    async function stopReader(keepOpen = true) {
      // Use Windows desktop voice backend, not browser SpeechSynthesis.
      // This fixes the black browser voice-control bar that appears but gives no sound.
      readerManualStop = true;
      readerRunId += 1;
      clearTimeout(readerTimer);
      readerTimer = null;
      try { await fetch('/api/stop_speech', {method: 'POST', headers: authHeaders()}); } catch (e) {}
      if ('speechSynthesis' in window) {
        try { speechSynthesis.cancel(); } catch (e) {}
      }
      readerUtterance = null;
      readerActive = false;
      readerPaused = false;
      updateReader();
      if (!keepOpen) reader.classList.remove('active');
      subEl.textContent = keepOpen ? 'Voice stopped.' : 'Voice control bar closed.';
    }
    async function speakReaderLine() {
      if (!readerLines.length) {
        subEl.textContent = 'Nothing to read in this answer.';
        updateReader();
        return;
      }
      const runId = ++readerRunId;
      readerManualStop = false;
      readerActive = true;
      readerPaused = false;
      clearTimeout(readerTimer);
      readerTimer = null;
      updateReader();
      subEl.textContent = 'Reading with Windows desktop voice...';
      try {
        await fetch('/api/stop_speech', {method: 'POST', headers: authHeaders()});
        const line = readerLines[readerIndex];
        const res = await fetch('/api/speak', {
          method: 'POST',
          headers: authHeaders({'Content-Type': 'application/json'}),
          body: JSON.stringify({text: line})
        });
        const data = await res.json().catch(() => ({ok: res.ok}));
        if (!res.ok || data.ok === false) throw new Error(data.error || 'Windows voice backend failed');
        readerTimer = setTimeout(async () => {
          if (runId !== readerRunId || readerManualStop || !readerActive || readerPaused) return;
          if (readerIndex < readerLines.length - 1) {
            readerIndex += 1;
            updateReader();
            await speakReaderLine();
          } else {
            readerActive = false;
            readerPaused = false;
            readerUtterance = null;
            subEl.textContent = 'Read complete.';
            updateReader();
          }
        }, readerDelayForLine(line));
      } catch (e) {
        readerActive = false;
        readerPaused = false;
        readerUtterance = null;
        updateReader();
        subEl.textContent = 'No sound from Windows voice. Run INSTALL_AUDIO_FIX.cmd and check Windows output device.';
        add('nova', 'Voice output failed: ' + e.message + '. Try: open sound settings, unmute speaker, choose correct output device, then type test voice.');
      }
    }
    function openReader(value) {
      if (readerText === value && (readerActive || readerPaused)) {
        stopReader(true);
        return;
      }
      readerText = value;
      readerLines = splitReaderText(value);
      readerIndex = 0;
      reader.classList.add('active');
      speakReaderLine();
    }
    async function pauseResumeReader() {
      if (!readerLines.length) return;
      if (readerActive && !readerPaused) {
        readerPaused = true;
        readerActive = false;
        readerManualStop = true;
        readerRunId += 1;
        clearTimeout(readerTimer);
        readerTimer = null;
        try { await fetch('/api/stop_speech', {method: 'POST', headers: authHeaders()}); } catch (e) {}
        subEl.textContent = 'Reading paused.';
        updateReader();
      } else {
        await speakReaderLine();
      }
    }
    async function moveReader(delta) {
      if (!readerLines.length) return;
      readerIndex = Math.max(0, Math.min(readerLines.length - 1, readerIndex + delta));
      if (readerActive || readerPaused) {
        await speakReaderLine();
      } else {
        updateReader();
      }
    }
    function showPermissionPopup(message, confirmCommand, onAllow = null) {
      pendingConfirmCommand = confirmCommand || '';
      pendingPermissionAction = typeof onAllow === 'function' ? onAllow : null;
      permissionText.textContent = message || 'Nova needs confirmation before doing this.';
      permissionAllow.textContent = (pendingConfirmCommand || pendingPermissionAction) ? 'Allow' : 'OK';
      permissionModal.classList.add('active');
      permissionAllow.focus();
    }
    function hidePermissionPopup() {
      permissionModal.classList.remove('active');
      pendingConfirmCommand = '';
      pendingPermissionAction = null;
      permissionAllow.textContent = 'Allow';
    }
    function setLocked(locked) {
      lockScreen.classList.toggle('active', locked);
      if (locked) {
        pinInput.focus();
        setPowerState(false);
      }
    }
    async function unlockNova() {
      const pin = pinInput.value.trim();
      if (!pin) {
        lockMessage.textContent = 'Enter your PIN.';
        return;
      }
      unlockButton.disabled = true;
      try {
        const res = await fetch('/api/unlock', {
          method: 'POST',
          headers: {'Content-Type': 'application/json'},
          body: JSON.stringify({pin})
        });
        const data = await res.json();
        if (!res.ok || !data.ok) {
          lockMessage.textContent = data.message || 'Wrong PIN.';
          return;
        }
        accessToken = data.token || '';
        pinInput.value = '';
        lockMessage.textContent = '';
        setLocked(false);
        refreshApiStatus();
        loadHistory();
      } catch (e) {
        lockMessage.textContent = 'Unlock failed. Keep Nova terminal open.';
      } finally {
        unlockButton.disabled = false;
      }
    }
    function setVoiceButtons(active) {
      voiceButton.textContent = active ? 'Stop Voice' : 'Start Voice';
      micSmall.textContent = active ? 'Stop' : 'Mic';
      desktopMic.textContent = active ? 'Stop Mic' : 'Desktop Mic';
      statusEl.textContent = active ? 'Listening' : 'Ready';
      mic.classList.toggle('listening', active);
    }
    function stopBrowserRecognition() {
      micWanted = false;
      clearTimeout(recognitionRestartTimer);
      try { if (recognition) recognition.abort(); } catch (e) {}
      try { if (recognition) recognition.stop(); } catch (e) {}
    }
    function stopDesktopMicLoop() {
      desktopMicWanted = false;
      if (desktopMicAbort) {
        try { desktopMicAbort.abort(); } catch (e) {}
        desktopMicAbort = null;
      }
      desktopMic.disabled = false;
    }
    function stopAllMic() {
      voiceBusy = false;
      speechNetworkErrors = 0;
      desktopMicErrorCount = 0;
      stopBrowserRecognition();
      stopDesktopMicLoop();
      if ('speechSynthesis' in window) speechSynthesis.cancel();
      setVoiceButtons(false);
    }
    function setPowerState(enabled) {
      assistantPowerOn = Boolean(enabled);
      powerToggle.textContent = assistantPowerOn ? 'ON' : 'OFF';
      powerToggle.classList.toggle('off', !assistantPowerOn);
      powerToggle.setAttribute('aria-pressed', assistantPowerOn ? 'true' : 'false');
      powerToggle.title = assistantPowerOn ? 'Turn Nova assistant off' : 'Turn Nova assistant on';
      if (!assistantPowerOn && (micWanted || desktopMicWanted)) {
        stopAllMic();
        promptEl.textContent = 'Assistant off.';
      }
    }
    function scheduleRecognitionStart(delay = 120) {
      clearTimeout(recognitionRestartTimer);
      if (!recognition || !micWanted || voiceBusy) return;
      recognitionRestartTimer = setTimeout(() => {
        if (!recognition || !micWanted || voiceBusy || listening) return;
        try {
          recognition.start();
        } catch (e) {
          if (!String(e && e.name || e).includes('InvalidStateError')) {
            subEl.textContent = 'Mic could not start yet. I will keep trying until you press Stop Voice.';
            showPermissionPopup('Nova could not start the browser microphone yet. Allow microphone permission in the browser/Windows popup. I will keep retrying until you press Stop Voice.', '');
            scheduleRecognitionStart(1800);
          }
        }
      }, delay);
    }
    async function backendSpeak(value) {
      try {
        await fetch('/api/stop_speech', {method: 'POST', headers: authHeaders()});
        const res = await fetch('/api/speak_sync', {
          method: 'POST',
          headers: authHeaders({'Content-Type': 'application/json'}),
          body: JSON.stringify({text: value})
        });
        const data = await res.json().catch(() => ({ok: res.ok}));
        if (!res.ok || data.ok === false) throw new Error(data.error || 'voice failed');
        return;
      } catch (e) {
        subEl.textContent = 'Windows voice could not speak. Check terminal for error.';
      }
    }
    async function loadVoices() {
      try {
        const res = await fetch('/api/voices');
        const data = await res.json();
        voiceSelect.innerHTML = '';
        if (!data.voices || !data.voices.length) {
          const opt = document.createElement('option');
          opt.textContent = 'Default Windows voice';
          opt.value = '';
          voiceSelect.appendChild(opt);
          return;
        }
        data.voices.forEach(v => {
          const opt = document.createElement('option');
          const detail = (data.details || []).find(item => item.name === v) || {};
          opt.textContent = detail.culture ? `${v} (${detail.culture})` : v;
          opt.value = v;
          if (v === data.selected) opt.selected = true;
          voiceSelect.appendChild(opt);
        });
        if (!data.hindi_voice_available) {
          const opt = document.createElement('option');
          opt.textContent = 'Hindi voice not installed - say setup Hindi voice';
          opt.value = '';
          voiceSelect.appendChild(opt);
        }
      } catch (e) {
        voiceSelect.innerHTML = '<option>Voice list unavailable</option>';
      }
    }
    async function loadHistory() {
      if (loadHistory.loading) return;
      loadHistory.loading = true;
      try {
        const res = await fetch('/api/history', {headers: authHeaders()});
        const data = await res.json();
        historyList.innerHTML = '';
        if (!data.items || !data.items.length) {
          historyList.innerHTML = '<div class="history-item">No history yet<small>Your searches and chats will appear here.</small></div>';
          return;
        }
        data.items.slice(-10).reverse().forEach(item => {
          const btn = document.createElement('button');
          btn.className = 'history-item';
          btn.type = 'button';
          btn.innerHTML = `${escapeHtml(item.user).slice(0, 80)}<small>${escapeHtml(item.assistant).slice(0, 120)}</small>`;
          btn.addEventListener('click', () => {
            text.value = item.user;
            add('you', item.user);
            add('nova', item.assistant, item.user);
          });
          historyList.appendChild(btn);
        });
      } catch (e) {
        historyList.innerHTML = '<div class="history-item">History unavailable</div>';
      } finally {
        loadHistory.loading = false;
      }
    }
    function applyFeatureSideEffects(key, enabled) {
      if (key === 'voice_enabled' && !enabled) stopAllMic();
      if (key === 'assistant_enabled') setPowerState(enabled);
      if (key === 'control_mode' && !enabled) subEl.textContent = 'Laptop control is OFF.';
    }
    async function updateFeatureSetting(key, enabled, checkbox) {
      try {
        const res = await fetch('/api/feature_settings', {
          method: 'POST',
          headers: authHeaders({'Content-Type': 'application/json'}),
          body: JSON.stringify({key, enabled})
        });
        const data = await res.json();
        if (res.status === 403) {
          setLocked(true);
          throw new Error(data.reply || 'Nova is locked.');
        }
        if (!res.ok || data.ok === false) throw new Error(data.reply || 'Setting failed');
        applyFeatureSideEffects(key, enabled);
        if (key === 'plan_mode_enabled') planModeToggle.checked = Boolean(enabled);
        if (key === 'pursue_goal_mode_enabled') pursueGoalToggle.checked = Boolean(enabled);
        subEl.textContent = data.reply || 'Setting updated.';
        if (checkbox) checkbox.checked = Boolean(enabled);
        if (data.settings) renderFeatureSettings(data.settings.items || []);
        refreshApiStatus();
      } catch (e) {
        if (checkbox) checkbox.checked = !enabled;
        add('nova', 'Setting update failed: ' + e.message);
      }
    }
    function renderFeatureSettings(items) {
      syncComposerModesFromSettings(items);
      featureSettingsList.innerHTML = '';
      if (!items.length) {
        featureSettingsList.innerHTML = '<div class="history-item">No settings available</div>';
        return;
      }
      items.forEach(item => {
        const row = document.createElement('div');
        row.className = 'setting-item';
        const info = document.createElement('div');
        info.innerHTML = `<strong>${escapeHtml(item.label)}</strong><small>${escapeHtml(item.description)}</small>`;
        const label = document.createElement('label');
        label.className = 'switch';
        const input = document.createElement('input');
        input.type = 'checkbox';
        input.checked = Boolean(item.enabled);
        const slider = document.createElement('span');
        slider.className = 'slider';
        label.appendChild(input);
        label.appendChild(slider);
        input.addEventListener('change', () => {
          const wantsOn = input.checked;
          if (!wantsOn) {
            updateFeatureSetting(item.key, false, input);
            return;
          }
          input.checked = false;
          showPermissionPopup(
            `${item.label}: ${item.description}\n\nDo you really want to turn this ON?`,
            '',
            async () => updateFeatureSetting(item.key, true, input)
          );
        });
        row.appendChild(info);
        row.appendChild(label);
        featureSettingsList.appendChild(row);
      });
    }
    async function loadFeatureSettings() {
      try {
        const res = await fetch('/api/feature_settings', {headers: authHeaders()});
        const data = await res.json();
        if (res.status === 403) {
          setLocked(true);
          return;
        }
        syncComposerModesFromSettings(data.items || []);
        renderFeatureSettings(data.items || []);
      } catch (e) {
        featureSettingsList.innerHTML = '<div class="history-item">Settings unavailable</div>';
      }
    }
    function toggleHistoryPanel() {
      document.body.classList.toggle('history-open');
      if (document.body.classList.contains('history-open')) {
        document.body.classList.remove('settings-open');
        loadHistory();
      }
    }
    function toggleSettingsPanel() {
      document.body.classList.toggle('settings-open');
      if (document.body.classList.contains('settings-open')) {
        document.body.classList.remove('history-open');
        loadFeatureSettings();
      }
    }
    function openSettingsPanel(message = '') {
      document.body.classList.add('settings-open');
      document.body.classList.remove('history-open');
      loadFeatureSettings();
      if (message) subEl.textContent = message;
    }
    function toggleComposerFeature(key, checkbox, label, description, exclusiveKey = '', exclusiveCheckbox = null) {
      const wantsOn = checkbox.checked;
      if (!wantsOn) {
        updateFeatureSetting(key, false, checkbox);
        return;
      }
      checkbox.checked = false;
      showPermissionPopup(
        `${label}: ${description}\n\nDo you really want to turn this ON?`,
        '',
        async () => {
          if (exclusiveKey && exclusiveCheckbox && exclusiveCheckbox.checked) {
            await updateFeatureSetting(exclusiveKey, false, exclusiveCheckbox);
          }
          await updateFeatureSetting(key, true, checkbox);
        }
      );
    }
    async function refreshApiStatus() {
      try {
        const [statusRes, networkRes] = await Promise.all([fetch('/api/status'), fetch('/api/network')]);
        const data = await statusRes.json();
        const network = await networkRes.json();
        const laptop = data.laptop_connected ? 'Laptop connected. ' : 'Laptop not connected. ';
        const local = data.local_only_mode ? 'Local-only. ' : '';
        const mobile = data.mobile_mode && data.mobile_url ? ` Mobile: ${data.mobile_url}` : '';
        if (data.access_lock_enabled && !accessToken) setLocked(true);
        setPowerState(Boolean(data.assistant_enabled));
        if (data.online) {
          const first = data.order && data.order.length ? data.order[0] : 'API';
          statusEl.textContent = `${network.online ? 'Online' : 'Net?'}: ${first}`;
          subEl.textContent = `${local}${laptop}Brain connected: ${data.order.join(', ')}.${mobile}`;
        } else {
          statusEl.textContent = network.online ? 'API offline' : 'Offline';
          subEl.textContent = local + laptop + (data.message || 'Add an API key in .env for online thinking.') + mobile;
        }
        const needs = Array.isArray(data.needs) ? data.needs : [];
        const actionableNeed = needs.find(item => item.confirmCommand && !shownNeedIds.has(item.id));
        const infoNeed = needs.find(item => !item.confirmCommand && !shownNeedIds.has(item.id));
        const need = actionableNeed || infoNeed;
        if (need) {
          shownNeedIds.add(need.id);
          showPermissionPopup(`${need.title}: ${need.message} This stays on this laptop only.`, need.confirmCommand || '');
        }
      } catch (e) {
        statusEl.textContent = 'Status error';
      }
    }
    async function checkNetworkMessage() {
      if (!navigator.onLine) return 'Browser is offline. Check Wi-Fi.';
      try {
        const res = await fetch('/api/network');
        const data = await res.json();
        return data.message || 'Network checked.';
      } catch (e) {
        return 'Browser could not reach Nova network checker.';
      }
    }
    async function send(command, fromVoice = false) {
      command = (command || '').trim();
      if (!command) return;
      if (fromVoice && /^(stop voice|stop mic|stop listening|mic off|voice off|band mic|mic band|voice band|sunna band|listening band)$/i.test(command)) {
        stopAllMic();
        add('you', command);
        add('nova', 'Voice stopped. Click Start Voice again when you want me to listen.');
        promptEl.textContent = 'Voice stopped.';
        subEl.textContent = 'Mic is OFF by your command.';
        return;
      }
      if (sendBusy) {
        subEl.textContent = 'Still working on the previous command.';
        return;
      }
      sendBusy = true;
      sendButton.disabled = true;
      if (fromVoice) {
        voiceBusy = true;
        try { recognition.stop(); } catch (e) {}
        subEl.textContent = `Heard: ${command}`;
      }
      add('you', command);
      promptEl.textContent = 'Working on it...';
      statusEl.textContent = 'Thinking';
      if (/\b(wifi|wi-fi|wi fi|bluetooth|blue tooth)\b/i.test(command) && /\b(on|off|enable|disable|turn on|turn off|switch on|switch off)\b/i.test(command)) {
        subEl.textContent = 'Windows permission popup may appear. Press Yes to allow Nova, then I will verify the result.';
      }
      try {
        const res = await fetch('/api/command', {
          method: 'POST',
          headers: authHeaders({'Content-Type': 'application/json'}),
          body: JSON.stringify({
            command,
            planMode: Boolean(planModeToggle.checked),
            pursueGoal: Boolean(pursueGoalToggle.checked)
          })
        });
        const data = await res.json();
        if (res.status === 403) {
          setLocked(true);
          throw new Error(data.reply || 'Locked');
        }
        if (!res.ok) throw new Error(`HTTP ${res.status}`);
        const reply = data.reply || 'Nova did not return a reply.';
        add('nova', reply, command);
        promptEl.textContent = reply.length > 62 ? 'Done.' : reply;
        if (data.needsConfirm) {
          subEl.textContent = 'Permission needed on screen.';
          showPermissionPopup(reply, data.confirmCommand);
        } else if (data.needsPermission) {
          subEl.textContent = 'Permission or setting needed.';
          showPermissionPopup(reply, '');
        } else {
          subEl.textContent = 'Say another command.';
        }
        if (/Assistant is ON/i.test(reply)) setPowerState(true);
        if (/Assistant is OFF/i.test(reply)) setPowerState(false);
        statusEl.textContent = 'Ready';
        const shouldSpeakReply = fromVoice || document.body.classList.contains('voice-page');
        if (shouldSpeakReply) {
          if (reply.length < 360 && !reply.includes('```')) await backendSpeak(reply);
          else await backendSpeak('I wrote the full answer in chat. Use the read button if you want me to read it.');
        }
        clearTimeout(historyTimer);
        historyTimer = setTimeout(loadHistory, 450);
        if (data.exit) setTimeout(() => window.close(), 500);
      } catch (e) {
        const message = 'Nova server did not answer. Keep the terminal open, then try again.';
        add('nova', message, command);
        promptEl.textContent = 'Connection issue.';
        subEl.textContent = message;
        statusEl.textContent = 'Error';
      } finally {
        sendBusy = false;
        sendButton.disabled = false;
        if (fromVoice) {
          setTimeout(() => {
            voiceBusy = false;
            scheduleRecognitionStart(120);
          }, 900);
        }
      }
    }
    function setupRecognition() {
      if (!SpeechRecognition) {
        subEl.textContent = 'Browser voice needs Chrome or Edge. Use Desktop Mic instead.';
        return;
      }
      recognition = new SpeechRecognition();
      recognition.lang = 'en-IN';
      recognition.continuous = false;
      recognition.interimResults = false;
      recognition.maxAlternatives = 1;
      recognition.onstart = () => {
        listening = true;
        setVoiceButtons(true);
        promptEl.textContent = 'Browser listening...';
      };
      recognition.onend = () => {
        listening = false;
        if (!micWanted) setVoiceButtons(false);
      };
      recognition.onerror = (event) => {
        listening = false;
        if (!micWanted) return;
        if (event.error === 'no-speech') {
          subEl.textContent = 'No speech heard. Mic is still ON. Click Stop Voice to turn it OFF, or speak again.';
          scheduleRecognitionStart(450);
          return;
        }
        if (event.error === 'network') {
          speechNetworkErrors += 1;
          subEl.textContent = 'Browser speech network issue. Mic mode is still ON; I will retry. Desktop Mic is usually better.';
          scheduleRecognitionStart(Math.min(4500, 900 + speechNetworkErrors * 500));
          return;
        }
        if (event.error === 'audio-capture') {
          subEl.textContent = 'No microphone found. Check Windows input device.';
          showPermissionPopup('No microphone is available to the browser. Check Windows input device or allow microphone permission. I will keep retrying until you press Stop Voice.', '');
          scheduleRecognitionStart(2200);
          return;
        }
        if (event.error === 'not-allowed') {
          subEl.textContent = 'Microphone permission blocked. Allow mic in browser address bar, or use Desktop Mic.';
          showPermissionPopup('Microphone permission is blocked. Use the browser address bar mic icon and choose Allow. I will keep the mic mode ON and retry until you press Stop Voice.', '');
          scheduleRecognitionStart(2200);
          return;
        }
        subEl.textContent = `Mic issue: ${event.error}. Click Stop Voice to turn OFF.`;
      };
      recognition.onresult = (event) => {
        const last = event.results[event.results.length - 1];
        const command = last[0].transcript;
        speechNetworkErrors = 0;
        subEl.textContent = `Heard: ${command}`;
        try { recognition.stop(); } catch (e) {}
        send(command, true);
      };
    }

    async function setPowerFromButton() {
      const nextEnabled = !assistantPowerOn;
      powerToggle.disabled = true;
      const oldLabel = powerToggle.textContent;
      powerToggle.textContent = nextEnabled ? 'ON...' : 'OFF...';
      try {
        const res = await fetch('/api/power', {
          method: 'POST',
          headers: authHeaders({'Content-Type': 'application/json'}),
          body: JSON.stringify({enabled: nextEnabled})
        });
        const data = await res.json();
        if (res.status === 403) {
          setLocked(true);
          throw new Error(data.reply || 'Nova is locked.');
        }
        if (!res.ok || !data.ok) throw new Error(data.reply || 'Power change failed');
        setPowerState(Boolean(data.enabled));
        const reply = data.reply || (data.enabled ? 'Assistant is ON.' : 'Assistant is OFF.');
        add('nova', reply);
        promptEl.textContent = data.enabled ? 'Ready.' : 'Assistant off.';
        subEl.textContent = data.enabled ? 'ON/OFF button fixed. Say or type a command.' : 'Assistant control is paused. Press ON to enable.';
        statusEl.textContent = data.enabled ? 'Ready' : 'Paused';
      } catch (e) {
        powerToggle.textContent = oldLabel;
        subEl.textContent = 'ON/OFF button failed: ' + e.message;
        await send(nextEnabled ? 'assistant on' : 'assistant off');
      } finally {
        powerToggle.disabled = false;
        refreshApiStatus();
      }
    }

    async function desktopListen() {
      if (!assistantPowerOn || powerToggle.classList.contains('off')) {
        subEl.textContent = 'Assistant is OFF. Press ON first.';
        return;
      }
      if (desktopMicWanted) {
        stopAllMic();
        promptEl.textContent = 'Voice stopped.';
        subEl.textContent = 'Mic is OFF. Click Start Voice once to turn ON again.';
        return;
      }
      stopBrowserRecognition();
      desktopMicWanted = true;
      desktopMicErrorCount = 0;
      setVoiceButtons(true);
      promptEl.textContent = 'Desktop mic ON';
      subEl.textContent = 'Desktop mic is ON. Speak a command. Click again to turn OFF.';
      while (desktopMicWanted && assistantPowerOn) {
        if (sendBusy || voiceBusy) {
          await new Promise(r => setTimeout(r, 350));
          continue;
        }
        desktopMicAbort = new AbortController();
        try {
          const res = await fetch('/api/listen_once', {method: 'POST', headers: authHeaders(), signal: desktopMicAbort.signal});
          const data = await res.json();
          if (!desktopMicWanted) break;
          if (!res.ok || !data.ok) throw new Error(data.error || data.text || 'Desktop mic failed');
          if (!data.text) {
            subEl.textContent = 'Listening... no speech heard. Mic is still ON. Click Stop Voice to turn OFF.';
            await new Promise(r => setTimeout(r, 250));
            continue;
          }
          subEl.textContent = `Desktop heard: ${data.text}`;
          await send(data.text, true);
        } catch (e) {
          if (!desktopMicWanted || String(e.name || e).includes('Abort')) break;
          desktopMicErrorCount += 1;
          subEl.textContent = 'Desktop mic had an issue, but it is still ON. I will retry. Click Stop Voice to turn OFF.';
          if (desktopMicErrorCount === 1 || desktopMicErrorCount % 5 === 0) {
            add('nova', 'Desktop mic retrying. If Windows shows a microphone permission popup, press Allow. Error: ' + e.message);
          }
          await new Promise(r => setTimeout(r, 900));
          continue;
        } finally {
          desktopMicAbort = null;
        }
      }
      stopAllMic();
    }
    function toggleVoice() {
      // Main Start/Stop button now uses the desktop microphone loop.
      // This avoids the browser SpeechRecognition network hiccup that kept retrying.
      if (desktopMicWanted || micWanted) {
        stopAllMic();
        promptEl.textContent = 'Voice stopped.';
        subEl.textContent = 'Mic is OFF. Click Start Voice once to turn ON again.';
        return;
      }
      desktopListen();
    }

    mic.addEventListener('click', toggleVoice);
    voiceButton.addEventListener('click', toggleVoice);
    micSmall.addEventListener('click', toggleVoice);
    desktopMic.addEventListener('click', desktopListen);
    composerMenuButton.addEventListener('click', (event) => {
      event.preventDefault();
      setComposerMenu(!composerToolbox.classList.contains('active'));
    });
    attachFilesButton.addEventListener('click', (event) => {
      event.preventDefault();
      setComposerMenu(false);
      mediaFile.click();
    });
    planModeToggle.addEventListener('change', () => {
      toggleComposerFeature(
        'plan_mode_enabled',
        planModeToggle,
        'Plan mode',
        'Nova will plan first and will not execute the command.',
        'pursue_goal_mode_enabled',
        pursueGoalToggle
      );
    });
    pursueGoalToggle.addEventListener('change', () => {
      toggleComposerFeature(
        'pursue_goal_mode_enabled',
        pursueGoalToggle,
        'Pursue goal',
        'Nova will use safe autopilot for clear multi-step goals.',
        'plan_mode_enabled',
        planModeToggle
      );
    });
    pluginsButton.addEventListener('click', (event) => {
      event.preventDefault();
      setComposerMenu(false);
      openSettingsPanel('Plugins and Nova tools are in Settings.');
    });
    permissionsButton.addEventListener('click', (event) => {
      event.preventDefault();
      setComposerMenu(false);
      openSettingsPanel('Default permissions opened. Turn features ON only after approving the popup.');
    });
    document.addEventListener('click', (event) => {
      if (!composerToolbox.classList.contains('active')) return;
      if (form.contains(event.target)) return;
      setComposerMenu(false);
    });
    document.addEventListener('keydown', (event) => {
      if (event.key === 'Escape') setComposerMenu(false);
    });
    powerToggle.addEventListener('click', setPowerFromButton);
    chatPageBtn.addEventListener('click', () => setPage('chat'));
    voicePageBtn.addEventListener('click', () => setPage('voice'));
    form.addEventListener('submit', (e) => {
      e.preventDefault();
      const value = text.value;
      send(value);
      if (value.trim()) text.value = '';
    });
    text.addEventListener('keydown', (event) => {
      if (event.key === 'Enter' && !event.shiftKey) {
        event.preventDefault();
        form.requestSubmit();
      }
    });
    text.addEventListener('input', () => {
      text.style.height = 'auto';
      text.style.height = Math.min(text.scrollHeight, 190) + 'px';
    });
    document.querySelectorAll('.chip').forEach(btn => btn.addEventListener('click', () => send(btn.textContent)));
    historyToggle.addEventListener('click', toggleHistoryPanel);
    settingsToggle.addEventListener('click', toggleSettingsPanel);
    settingsClose.addEventListener('click', () => document.body.classList.remove('settings-open'));
    clearHistory.addEventListener('click', loadHistory);
    newChatButton.addEventListener('click', () => {
      chat.innerHTML = '';
      send('new chat');
    });
    exportChatButton.addEventListener('click', () => send('export chat'));
    chatHelpButton.addEventListener('click', () => send('chatgpt help'));
    mediaRun.addEventListener('click', runSmartMediaTask);
    mediaPrompt.addEventListener('keydown', (event) => {
      if (event.key === 'Enter') {
        event.preventDefault();
        runSmartMediaTask();
      }
    });
    mediaFile.addEventListener('change', () => {
      const files = Array.from(mediaFile.files || []);
      if (files.length === 1) subEl.textContent = `Selected media: ${files[0].name}`;
      else if (files.length > 1) subEl.textContent = `Selected ${files.length} media files.`;
    });
    permissionAllow.addEventListener('click', () => {
      const command = pendingConfirmCommand;
      const action = pendingPermissionAction;
      hidePermissionPopup();
      if (action) action();
      else if (command) send(command);
    });
    permissionCancel.addEventListener('click', () => {
      hidePermissionPopup();
      add('nova', 'Cancelled.');
      subEl.textContent = 'Permission cancelled.';
    });
    permissionClose.addEventListener('click', () => {
      hidePermissionPopup();
      subEl.textContent = 'Permission closed.';
    });
    permissionModal.addEventListener('click', (event) => {
      if (event.target === permissionModal) {
        hidePermissionPopup();
        subEl.textContent = 'Permission cancelled.';
      }
    });
    unlockButton.addEventListener('click', unlockNova);
    pinInput.addEventListener('keydown', (event) => {
      if (event.key === 'Enter') unlockNova();
    });
    voiceSelect.addEventListener('change', async () => {
      await fetch('/api/voice', {
        method: 'POST',
        headers: authHeaders({'Content-Type': 'application/json'}),
        body: JSON.stringify({voice: voiceSelect.value})
      });
      backendSpeak(`Voice changed to ${voiceSelect.value || 'default'}.`);
    });
    testVoice.addEventListener('click', async () => {
      if (testVoiceBusy) return;
      testVoiceBusy = true;
      testVoice.disabled = true;
      await backendSpeak('Hello Adarsh. My voice is now in slow and clear English mode.');
      setTimeout(() => {
        testVoiceBusy = false;
        testVoice.disabled = false;
      }, 2200);
    });
    readerPlay.addEventListener('click', pauseResumeReader);
    readerPrev.addEventListener('click', () => moveReader(-1));
    readerNext.addEventListener('click', () => moveReader(1));
    readerStop.addEventListener('click', () => stopReader(true));
    readerClose.addEventListener('click', () => stopReader(false));
    readerShuffle.addEventListener('click', () => { readerIndex = 0; speakReaderLine(); });
    readerProgress.addEventListener('input', () => { readerIndex = Number(readerProgress.value || 0); speakReaderLine(); });
    setupRecognition();
    refreshApiStatus();
    loadVoices();
    loadHistory();
    loadFeatureSettings();
  </script>
</body>
</html>
"""


def web_gui_main(port=8765, host="127.0.0.1", mobile=False, mobile_pin=""):
    import http.server
    import socketserver

    global SERVER_MOBILE_MODE, SERVER_MOBILE_URL, SERVER_LOCAL_URL
    prepare_assistant_runtime()
    SERVER_MOBILE_MODE = bool(mobile)

    class ReusableTCPServer(socketserver.ThreadingMixIn, socketserver.TCPServer):
        allow_reuse_address = True
        daemon_threads = True

    class Handler(http.server.BaseHTTPRequestHandler):
        def log_message(self, fmt, *args):
            log("Web GUI: " + (fmt % args))

        def send_text(self, status, body, content_type="text/plain; charset=utf-8"):
            data = body.encode("utf-8")
            self.send_response(status)
            self.send_header("Content-Type", content_type)
            self.send_header("Content-Length", str(len(data)))
            self.send_header("Cache-Control", "no-store")
            self.send_header("X-Content-Type-Options", "nosniff")
            self.send_header("Referrer-Policy", "no-referrer")
            self.send_header("Permissions-Policy", "camera=(), geolocation=(), payment=()")
            self.end_headers()
            self.wfile.write(data)

        def authorized(self):
            if not access_lock_enabled():
                return True
            return self.headers.get("X-Nova-Token", "") in ACCESS_TOKENS

        def require_auth_json(self):
            if self.authorized():
                return True
            self.send_text(403, json.dumps({"ok": False, "reply": "Nova is locked. Enter your local PIN."}), "application/json; charset=utf-8")
            return False

        def do_GET(self):
            if self.path in ["/", "/index.html"]:
                self.send_text(200, WEB_GUI_HTML, "text/html; charset=utf-8")
            elif self.path == "/api/status":
                status = api_status()
                if not status["requests_installed"]:
                    status["message"] = "Python requests package is missing."
                elif not status["order"]:
                    status["message"] = "No online API key found in .env."
                elif not status.get("python_http", False):
                    status["message"] = status.get("network_message") or "Python/API HTTP is blocked or filtered."
                else:
                    status["message"] = "Online brain ready."
                status["needs"] = setup_needs()
                status["mobile_mode"] = SERVER_MOBILE_MODE
                status["mobile_url"] = SERVER_MOBILE_URL
                status["local_url"] = SERVER_LOCAL_URL
                status["media"] = media_tool_status()
                self.send_text(200, json.dumps(status), "application/json; charset=utf-8")
            elif self.path == "/api/voices":
                details = windows_voice_details()
                payload = {
                    "voices": [v["name"] for v in details],
                    "details": details,
                    "selected": settings.get("voice_name", ""),
                    "mode": settings.get("voice_language_mode", "auto"),
                    "hindi_voice_available": bool(available_hindi_voice()),
                    "enabled": settings.get("voice_enabled", True),
                }
                self.send_text(200, json.dumps(payload), "application/json; charset=utf-8")
            elif self.path == "/api/history":
                if not self.require_auth_json():
                    return
                self.send_text(200, json.dumps({"items": recent_history_pairs(10)}), "application/json; charset=utf-8")
            elif self.path == "/api/feature_settings":
                if not self.require_auth_json():
                    return
                self.send_text(200, json.dumps(feature_settings_payload()), "application/json; charset=utf-8")
            elif self.path == "/api/network":
                self.send_text(200, json.dumps(network_status()), "application/json; charset=utf-8")
            elif self.path == "/api/audio_status":
                self.send_text(200, json.dumps(audio_status()), "application/json; charset=utf-8")
            elif self.path == "/api/diagnostics":
                if not self.require_auth_json():
                    return
                self.send_text(200, json.dumps({"text": assistant_diagnostics(repair=False)}), "application/json; charset=utf-8")
            else:
                self.send_text(404, "Not found")

        def do_POST(self):
            if self.path != "/api/unlock" and not self.require_auth_json():
                return
            if self.path == "/api/feedback":
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    data = json.loads(raw)
                    save_feedback(data.get("kind", ""), data.get("user", ""), data.get("assistant", ""))
                    self.send_text(200, '{"ok":true}', "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "error": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/unlock":
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    pin = json.loads(raw).get("pin", "")
                    if verify_access_pin(pin):
                        self.send_text(200, json.dumps({"ok": True, "token": issue_access_token()}), "application/json; charset=utf-8")
                    else:
                        self.send_text(403, json.dumps({"ok": False, "message": "Wrong PIN."}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "message": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/media/upload":
                length = int(self.headers.get("Content-Length", "0"))
                max_len = int(settings.get("max_media_upload_mb", 80)) * 1024 * 1024 * 2
                if length > max_len:
                    self.send_text(413, json.dumps({"ok": False, "reply": f"Upload is too large. Limit is {settings.get('max_media_upload_mb', 80)} MB."}), "application/json; charset=utf-8")
                    return
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    data = json.loads(raw or "{}")
                    upload_ext = Path(str(data.get("filename", ""))).suffix.lower()
                    if upload_ext in IMAGE_EXTS | VIDEO_EXTS and not settings.get("media_studio_enabled", True):
                        self.send_text(403, json.dumps({"ok": False, "reply": "Media Studio is OFF in Nova Settings."}), "application/json; charset=utf-8")
                        return
                    if upload_ext and upload_ext not in IMAGE_EXTS | VIDEO_EXTS and not settings.get("file_studio_enabled", True):
                        self.send_text(403, json.dumps({"ok": False, "reply": "File Studio is OFF in Nova Settings."}), "application/json; charset=utf-8")
                        return
                    ok, msg, path = save_uploaded_media(data.get("filename", ""), data.get("data", ""))
                    if not ok:
                        self.send_text(400, json.dumps({"ok": False, "reply": msg}), "application/json; charset=utf-8")
                        return
                    reply = msg
                    if data.get("analyze") and path:
                        if path.suffix.lower() in IMAGE_EXTS:
                            _, reply = analyze_image_file(path, data.get("prompt", ""))
                        elif path.suffix.lower() in VIDEO_EXTS:
                            _, reply = analyze_video_file(path)
                        else:
                            _, reply = analyze_general_file(path)
                    record_interaction("uploaded file: " + (data.get("filename", "") or "file"), reply)
                    self.send_text(200, json.dumps({"ok": True, "reply": reply, "path": str(path) if path else ""}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "reply": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/media/run":
                length = int(self.headers.get("Content-Length", "0"))
                max_len = int(settings.get("max_media_upload_mb", 80)) * 1024 * 1024 * 4
                if length > max_len:
                    self.send_text(413, json.dumps({"ok": False, "reply": f"Selected upload is too large. Limit is about {settings.get('max_media_upload_mb', 80)} MB per task."}), "application/json; charset=utf-8")
                    return
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    data = json.loads(raw or "{}")
                    prompt = str(data.get("prompt", "") or "")
                    files_payload = data.get("files", []) or []
                    selected_paths = []
                    for item in files_payload[:20]:
                        filename = item.get("filename", "")
                        upload_ext = Path(str(filename)).suffix.lower()
                        if upload_ext in IMAGE_EXTS | VIDEO_EXTS and not settings.get("media_studio_enabled", True):
                            self.send_text(403, json.dumps({"ok": False, "reply": "Media Studio is OFF in Nova Settings."}), "application/json; charset=utf-8")
                            return
                        if upload_ext and upload_ext not in IMAGE_EXTS | VIDEO_EXTS and not settings.get("file_studio_enabled", True):
                            self.send_text(403, json.dumps({"ok": False, "reply": "File Studio is OFF in Nova Settings."}), "application/json; charset=utf-8")
                            return
                        ok, msg, path = save_uploaded_media(filename, item.get("data", ""))
                        if not ok:
                            self.send_text(400, json.dumps({"ok": False, "reply": msg}), "application/json; charset=utf-8")
                            return
                        if path:
                            selected_paths.append(path)
                    c_prompt = normalize_words(prompt)
                    if not selected_paths:
                        if any(x in c_prompt for x in ["image", "photo", "picture", "video", "media", "wallpaper", "poster", "thumbnail"]) and not settings.get("media_studio_enabled", True):
                            self.send_text(403, json.dumps({"ok": False, "reply": "Media Studio is OFF in Nova Settings."}), "application/json; charset=utf-8")
                            return
                        if any(x in c_prompt for x in ["file", "document", "docx", "xlsx", "pptx", "pdf", "csv", "json"]) and not settings.get("file_studio_enabled", True):
                            self.send_text(403, json.dumps({"ok": False, "reply": "File Studio is OFF in Nova Settings."}), "application/json; charset=utf-8")
                            return
                    ok, reply = smart_selected_media_task(prompt, selected_paths)
                    record_interaction("media: " + (prompt.strip() or ", ".join(p.name for p in selected_paths)), reply)
                    status = 200 if ok else 400
                    self.send_text(status, json.dumps({"ok": bool(ok), "reply": reply, "paths": [str(p) for p in selected_paths]}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "reply": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/power":
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    token = self.headers.get("X-Nova-Token", "")
                    if access_lock_enabled() and token not in ACCESS_TOKENS:
                        self.send_text(403, json.dumps({"ok": False, "reply": "Nova is locked. Enter your local PIN."}), "application/json; charset=utf-8")
                        return
                    data = json.loads(raw or "{}")
                    enabled = bool(data.get("enabled", True))
                    reply = set_assistant_power(enabled)
                    self.send_text(200, json.dumps({"ok": True, "enabled": enabled, "reply": reply}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "reply": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/feature_settings":
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    data = json.loads(raw or "{}")
                    ok, reply = set_feature_setting(data.get("key", ""), bool(data.get("enabled", False)))
                    status = 200 if ok else 400
                    self.send_text(status, json.dumps({"ok": ok, "reply": reply, "settings": feature_settings_payload()}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "reply": str(e)}), "application/json; charset=utf-8")
                return
            if self.path in ["/api/speak", "/api/speak_sync"]:
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    text = json.loads(raw or "{}").get("text", "")
                    if self.path == "/api/speak_sync":
                        ok = speak(text, wait=True)
                    else:
                        threading.Thread(target=speak, args=(text, False), daemon=True).start()
                        ok = True
                    self.send_text(200, json.dumps({"ok": bool(ok)}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "error": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/stop_speech":
                stop_speech()
                self.send_text(200, '{"ok":true}', "application/json; charset=utf-8")
                return
            if self.path == "/api/listen_once":
                try:
                    text = listen_once(timeout=7, phrase_time_limit=8)
                    self.send_text(200, json.dumps({"ok": True, "text": text}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "error": str(e)}), "application/json; charset=utf-8")
                return
            if self.path == "/api/voice":
                length = int(self.headers.get("Content-Length", "0"))
                raw = self.rfile.read(length).decode("utf-8")
                try:
                    voice = json.loads(raw).get("voice", "")
                    settings["voice_name"] = voice
                    settings["voice_enabled"] = True
                    save_json(SETTINGS_FILE, settings)
                    self.send_text(200, json.dumps({"ok": True, "voice": voice}), "application/json; charset=utf-8")
                except Exception as e:
                    self.send_text(500, json.dumps({"ok": False, "error": str(e)}), "application/json; charset=utf-8")
                return
            if self.path != "/api/command":
                self.send_text(404, '{"reply":"Not found"}', "application/json; charset=utf-8")
                return
            length = int(self.headers.get("Content-Length", "0"))
            raw = self.rfile.read(length).decode("utf-8")
            try:
                token = self.headers.get("X-Nova-Token", "")
                if access_lock_enabled() and token not in ACCESS_TOKENS:
                    self.send_text(403, json.dumps({"reply": "Nova is locked. Enter your local PIN.", "exit": False}), "application/json; charset=utf-8")
                    return
                command_payload = json.loads(raw or "{}")
                command = command_payload.get("command", "")
                if len(str(command)) > int(settings.get("max_command_length", 1200)):
                    self.send_text(413, json.dumps({"reply": "Command is too long. For safety, keep one command under 1200 characters.", "exit": False}), "application/json; charset=utf-8")
                    return
                plan_mode = bool(command_payload.get("planMode", False) or settings.get("plan_mode_enabled", False))
                pursue_goal = bool(command_payload.get("pursueGoal", False) or settings.get("pursue_goal_mode_enabled", False))
                if plan_mode:
                    reply = plan_mode_response(command)
                elif pursue_goal and looks_like_autopilot(normalize_words(command)):
                    reply = autopilot_command(command, ask_confirm=None)
                else:
                    reply = handle_command(command, ask_confirm=None)
                reply = force_english_reply(reply)
                if reply != "EXIT":
                    record_interaction(command, reply)
                needs_confirm = isinstance(reply, str) and reply.startswith("For safety I need a clear confirmation. Say: confirm ")
                confirm_command = ""
                if needs_confirm:
                    confirm_command = reply.split("Say: ", 1)[1].strip()
                needs_permission = isinstance(reply, str) and any(
                    marker in reply.lower()
                    for marker in [
                        "off in nova settings",
                        "permission is off",
                        "permission needed",
                        "windows permission popup",
                        "turn it on and allow",
                        "allow the popup",
                    ]
                )
                payload = {
                    "reply": "Closing Nova." if reply == "EXIT" else reply,
                    "exit": reply == "EXIT",
                    "needsConfirm": needs_confirm,
                    "needsPermission": bool(needs_permission and not needs_confirm),
                    "confirmCommand": confirm_command,
                }
            except Exception as e:
                log(f"Web command error: {e}")
                payload = {"reply": f"I tried, but got an error: {e}", "exit": False}
            self.send_text(200, json.dumps(payload), "application/json; charset=utf-8")

    def find_port(start):
        for p in range(start, start + 20):
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                if s.connect_ex(("127.0.0.1", p)) != 0:
                    return p
        return start

    port = find_port(port)
    server_host = host or ("0.0.0.0" if mobile else "127.0.0.1")
    SERVER_LOCAL_URL = f"http://127.0.0.1:{port}"
    SERVER_MOBILE_URL = ""
    mobile_pin_note = ""
    if mobile:
        pin_to_show = ""
        if mobile_pin:
            ok, msg = set_access_pin(mobile_pin)
            if ok:
                pin_to_show = re.sub(r"\D", "", str(mobile_pin))
            else:
                generated = generate_pin()
                set_access_pin(generated)
                pin_to_show = generated
                mobile_pin_note = msg + " Generated a safe PIN instead."
        elif not access_lock_enabled():
            generated = generate_pin()
            set_access_pin(generated)
            pin_to_show = generated
        else:
            mobile_pin_note = "Use your existing Nova PIN."
        lan_ip = get_lan_ip()
        SERVER_MOBILE_URL = f"http://{lan_ip}:{port}"
        lines = [
            "Nova Mobile Mode",
            f"Laptop URL: {SERVER_LOCAL_URL}",
            f"Mobile URL: {SERVER_MOBILE_URL}",
            "Use this only on your own trusted Wi-Fi/LAN.",
            "Phone and laptop must be on the same Wi-Fi network.",
            "If Windows Firewall asks, allow Private networks.",
            f"PIN: {pin_to_show}" if pin_to_show else mobile_pin_note,
            "Mobile can control this laptop through Nova commands; it does not control phone apps.",
        ]
        MOBILE_ACCESS_FILE.write_text("\n".join(x for x in lines if x), encoding="utf-8")
    url = SERVER_LOCAL_URL
    log(f"Starting Google Assistant style GUI at {url}")
    speak("Nova Assistant GUI is ready.")
    webbrowser.open(url)
    with ReusableTCPServer((server_host, port), Handler) as httpd:
        safe_print(f"Nova Google Assistant style GUI running at {url}")
        if mobile:
            safe_print(f"Nova mobile URL: {SERVER_MOBILE_URL}")
            safe_print(f"Mobile access details saved to: {MOBILE_ACCESS_FILE}")
        safe_print("Keep this terminal open. Close it to stop Nova.")
        httpd.serve_forever()


class VoiceWorker:
    def __init__(self, app):
        self.app = app
        self.running = False
        self.thread = None

    def start(self):
        if self.running:
            return
        self.running = True
        self.thread = threading.Thread(target=self.loop, daemon=True)
        self.thread.start()
        self.app.add_bot("Listening started.")
        speak("Listening started.")

    def stop(self):
        self.running = False
        self.app.add_bot("Listening stopped.")
        speak("Listening stopped.")

    def loop(self):
        while self.running:
            text = listen_once(timeout=8, phrase_time_limit=8)
            if not self.running:
                break
            if text:
                self.app.process_user(text)
            time.sleep(0.1)


class NovaGUI:
    def __init__(self, root):
        self.root = root
        self.voice = VoiceWorker(self)
        self.root.title(APP_NAME)
        self.root.geometry("820x620")
        self.root.minsize(760, 560)
        self.build()
        speak(f"{APP_NAME} opened. I am ready.")

    def build(self):
        style = ttk.Style()
        try:
            style.theme_use("clam")
        except Exception:
            pass

        # Keyboard-fixed layout:
        # The command box is at the TOP so it is always visible even on small screens.
        main = ttk.Frame(self.root, padding=12)
        main.pack(fill="both", expand=True)

        header = ttk.Label(main, text="🤖 Nova AI Ultimate", font=("Segoe UI", 20, "bold"))
        header.pack(anchor="w")
        sub = ttk.Label(main, text="Type command here or use voice. Online AI uses .env API keys; laptop actions work offline.", font=("Segoe UI", 10))
        sub.pack(anchor="w", pady=(0, 8))

        # INPUT ROW FIRST - this fixes the issue where keyboard box was hidden at bottom.
        row = ttk.Frame(main)
        row.pack(fill="x", pady=(0, 8))
        self.entry = ttk.Entry(row, font=("Segoe UI", 13))
        self.entry.pack(side="left", fill="x", expand=True, ipady=6)
        self.entry.bind("<Return>", lambda e: self.send())
        ttk.Button(row, text="Send", command=self.send).pack(side="left", padx=(6, 0))
        ttk.Button(row, text="Start Voice", command=self.voice.start).pack(side="left", padx=(6, 0))
        ttk.Button(row, text="Stop Voice", command=self.voice.stop).pack(side="left", padx=(6, 0))

        # Quick suggestion buttons are hidden; all commands still work through text or voice.

        self.chat = scrolledtext.ScrolledText(main, wrap="word", font=("Segoe UI", 11), height=15)
        self.chat.pack(fill="both", expand=True)
        self.chat.configure(state="disabled")

        self.add_bot("Ready. First type: learn laptop. I will scan installed Start Menu apps. I will not Google unless you say search/google.")
        self.root.after(500, self.focus_input)

    def focus_input(self):
        try:
            self.root.lift()
            self.entry.focus_set()
            self.entry.focus_force()
        except Exception:
            pass

    def add_line(self, who, text):
        self.chat.configure(state="normal")
        self.chat.insert("end", f"{who}: {text}\n\n")
        self.chat.see("end")
        self.chat.configure(state="disabled")

    def add_bot(self, text):
        self.add_line("Nova", text)

    def ask_confirm(self, text):
        return messagebox.askyesno(APP_NAME, text)

    def process_user(self, text):
        self.add_line("You", text)
        def work():
            reply = handle_command(text, ask_confirm=self.ask_confirm)
            reply = force_english_reply(reply)
            if reply == "EXIT":
                self.root.after(0, self.root.destroy)
                return
            self.root.after(0, lambda: self.add_bot(reply))
            speak(reply)
        threading.Thread(target=work, daemon=True).start()

    def send(self):
        text = self.entry.get().strip()
        if not text:
            return
        self.entry.delete(0, "end")
        self.process_user(text)

    def quick(self, text):
        self.entry.delete(0, "end")
        self.entry.insert(0, text)
        self.send()


def console_main():
    speak(f"{APP_NAME} voice terminal mode started. Say a command after the listening message.")
    if sr:
        print("Voice mode is ON. Say commands like: hello, learn laptop, open settings, play jhol song in spotify.")
        print("Say 'exit' to close Nova.")
        while True:
            cmd = listen_once(timeout=10, phrase_time_limit=8)
            if not cmd:
                continue
            print("You:", cmd)
            reply = handle_command(cmd, ask_confirm=lambda q: input(q + " (yes/no): ").lower().startswith("y"))
            reply = force_english_reply(reply)
            if reply == "EXIT":
                speak("Closing Nova.")
                break
            print("Nova:", reply)
            speak(reply)
        return

    speak("Voice recognition package is not installed, so keyboard mode is open.")
    while True:
        try:
            cmd = input("You: ").strip()
        except KeyboardInterrupt:
            break
        if not cmd:
            continue
        reply = handle_command(cmd, ask_confirm=lambda q: input(q + " (yes/no): ").lower().startswith("y"))
        reply = force_english_reply(reply)
        if reply == "EXIT":
            break
        print("Nova:", reply)
        speak(reply)


def preflight_checks():
    """Fast non-launching setup check for START/RUN scripts and manual testing."""
    required = {
        "python-dotenv": "dotenv",
        "requests": "requests",
        "pyttsx3": "pyttsx3",
        "SpeechRecognition": "speech_recognition",
        "PyAudio": "pyaudio",
        "pyautogui": "pyautogui",
        "pyperclip": "pyperclip",
        "pygetwindow": "pygetwindow",
        "python-docx": "docx",
        "openpyxl": "openpyxl",
        "python-pptx": "pptx",
        "Pillow": "PIL",
    }
    lines = [
        "Nova preflight check",
        f"- Python: {sys.version.split()[0]} at {sys.executable}",
        f"- Workspace: {BASE_DIR}",
    ]
    missing = []
    for label, module in required.items():
        ok = importlib.util.find_spec(module) is not None
        lines.append(f"- {label}: {'OK' if ok else 'MISSING'}")
        if not ok:
            missing.append(label)
    status = api_status()
    lines.append(f"- Assistant switch: {'ON' if settings.get('assistant_enabled', True) else 'OFF'}")
    lines.append(f"- Control mode: {'ON' if settings.get('control_mode', True) else 'OFF'}")
    lines.append(f"- Laptop profile: {'OK' if LAPTOP_PROFILE_FILE.exists() else 'MISSING'}")
    lines.append(f"- App index: {'OK' if APP_INDEX_FILE.exists() else 'MISSING'}")
    lines.append(f"- Online brain: {'READY: ' + ', '.join(status['order']) if status['online'] else 'OFFLINE/OPTIONAL'}")
    media = media_tool_status()
    lines.append(f"- Media images: {'OK' if media['pillow'] else 'MISSING'}")
    lines.append(f"- Media videos: {'OK' if media['ffmpeg'] else 'OPTIONAL FFMPEG MISSING'}")
    try:
        voice_count = len(windows_voice_names())
        lines.append(f"- Windows voices: {voice_count if voice_count else 'not detected'}")
    except Exception as e:
        lines.append(f"- Windows voices: check failed ({e})")
    if missing:
        lines.append("Result: needs package repair. Run INSTALL_AND_RUN.bat.")
    else:
        lines.append("Result: setup looks ready. Nova was not launched.")
    return "\n".join(lines), not missing


def main():
    parser = argparse.ArgumentParser(add_help=False)
    parser.add_argument("--console", action="store_true")
    parser.add_argument("--tk", action="store_true")
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--mobile", action="store_true")
    parser.add_argument("--host", default="")
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--pin", default="")
    args, _ = parser.parse_known_args()
    if args.check:
        text, ok = preflight_checks()
        print(text)
        return 0 if ok else 1
    if args.self_test:
        print(full_function_audit(repair=False))
        return 0
    if args.console:
        console_main()
        return
    if not args.tk:
        host = args.host or ("0.0.0.0" if args.mobile else "127.0.0.1")
        web_gui_main(port=args.port, host=host, mobile=args.mobile, mobile_pin=args.pin)
        return
    if tk:
        try:
            root = tk.Tk()
            NovaGUI(root)
            root.mainloop()
            return
        except Exception as e:
            log(f"GUI startup failed: {e}")
            print("Nova GUI could not start because Python Tkinter/Tcl is not installed correctly.")
            print("Starting console mode instead. Type commands here, or type exit to close.")
            console_main()
    else:
        console_main()


if __name__ == "__main__":
    sys.exit(main() or 0)
